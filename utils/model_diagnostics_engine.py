from __future__ import annotations

from dataclasses import dataclass
from fractions import Fraction
from io import BytesIO
from itertools import product
import math
from typing import Iterable

import numpy as np

try:
    from scipy.optimize import least_squares
except ImportError:
    least_squares = None

try:
    from openpyxl import load_workbook
except ImportError:
    load_workbook = None

from utils.kinetic_fitting_engine import KineticFittingProblem, parse_fitting_workbook, residual_vector, species_scales


class ModelDiagnosticsValidationError(ValueError):
    pass


@dataclass(frozen=True)
class CompletedFitInfo:
    problem: KineticFittingProblem
    summary: dict[str, object]
    fitted_parameters: dict[str, float]
    measured_species: tuple[str, ...]


@dataclass(frozen=True)
class InvariantResult:
    coefficients: tuple[int, ...]
    participating_species: tuple[str, ...]
    fully_observed: bool
    missing_species: tuple[str, ...]


@dataclass(frozen=True)
class BalancePoint:
    dataset_id: str
    time_s: float
    balance_value: float
    baseline_value: float
    normalized_deviation: float | None


@dataclass(frozen=True)
class BalanceUncertaintyPoint:
    dataset_id: str
    time_s: float
    balance_value: float
    baseline_value: float
    normalized_deviation: float | None
    balance_sigma: float
    baseline_sigma: float
    normalized_deviation_sigma: float | None


@dataclass(frozen=True)
class ReconciledObservation:
    dataset_id: str
    time_s: float
    species_id: str
    measured_value: float
    reconciled_value: float
    correction: float
    sigma: float
    standardized_correction: float | None
    balance_before: float
    balance_after: float
    target_balance: float
    weighted_inconsistency: float | None




@dataclass(frozen=True)
class LocalParameterDiagnostic:
    parameter_id: str
    fitted_value: float
    lower_bound: float
    upper_bound: float
    finite_difference_step: float
    finite_difference_method: str
    local_parameter_scale: float
    relative_sensitivity_rms: float
    one_percent_effect_rms: float
    approximate_standard_error: float | None
    relative_standard_error_percent: float | None
    bound_position_percent: float
    nearest_bound_distance_percent: float
    strongest_correlation_parameter: str | None
    strongest_correlation: float | None


@dataclass(frozen=True)
class LocalParameterDiagnosticsResult:
    weighting: str
    parameter_ids: tuple[str, ...]
    fitted_values: tuple[float, ...]
    weighted_residuals: tuple[float, ...]
    jacobian: np.ndarray
    scaled_jacobian: np.ndarray
    parameter_scales: tuple[float, ...]
    singular_values: tuple[float, ...]
    condition_number: float
    numerical_rank: int
    degrees_of_freedom: int
    objective_sse: float
    residual_variance_estimate: float | None
    covariance: np.ndarray
    correlation: np.ndarray
    weakest_direction: tuple[float, ...]
    parameter_rows: tuple[LocalParameterDiagnostic, ...]

@dataclass(frozen=True)
class ResidualDiagnosticPoint:
    dataset_id: str
    time_s: float
    species_id: str
    measured_value: float
    fitted_value: float
    residual_model_minus_measured: float
    supplied_sigma: float | None
    effective_sigma: float
    uncertainty_standardized_residual: float | None
    species_scale: float
    fit_scaled_residual: float | None


def _header_map(ws) -> dict[str, int]:
    return {
        str(cell.value).strip(): idx
        for idx, cell in enumerate(ws[1], start=1)
        if cell.value not in (None, "")
    }


def parse_completed_fit_workbook(workbook_bytes: bytes) -> CompletedFitInfo:
    """Read a KineticFittingPot completed kinetic-fitting workbook.

    The original model/data sheets are parsed by the existing fitting parser so
    post-fit diagnostics use exactly the same model definition and observation
    set as the fit. Fit_Parameters and Fit_Summary are then read as result
    metadata; they are not used to change the stoichiometric checks.
    """
    problem = parse_fitting_workbook(workbook_bytes)
    if load_workbook is None:
        raise RuntimeError("Post-fit diagnostics requires openpyxl.")
    try:
        wb = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    except Exception as exc:
        raise ModelDiagnosticsValidationError(f"KineticFittingPot could not read this completed fitting workbook: {exc}") from exc

    required = {"Fit_Parameters", "Fit_Predictions", "Fit_Residuals", "Fit_Summary"}
    missing = sorted(required.difference(wb.sheetnames))
    if missing:
        raise ModelDiagnosticsValidationError(
            "This does not appear to be a completed KineticFittingPot kinetic-fitting workbook. "
            "Missing result sheet(s): " + ", ".join(missing)
        )

    summary: dict[str, object] = {}
    ws = wb["Fit_Summary"]
    for row in range(2, ws.max_row + 1):
        key = ws.cell(row=row, column=1).value
        if key in (None, ""):
            continue
        summary[str(key).strip()] = ws.cell(row=row, column=2).value

    fitted: dict[str, float] = {}
    ws = wb["Fit_Parameters"]
    headers = _header_map(ws)
    for required_header in ("Parameter_ID", "Fitted_Value"):
        if required_header not in headers:
            raise ModelDiagnosticsValidationError(
                f"Fit_Parameters sheet is missing '{required_header}'."
            )
    for row in range(2, ws.max_row + 1):
        pid = ws.cell(row=row, column=headers["Parameter_ID"]).value
        value = ws.cell(row=row, column=headers["Fitted_Value"]).value
        if pid in (None, "") or value in (None, ""):
            continue
        try:
            fitted[str(pid).strip()] = float(value)
        except (TypeError, ValueError) as exc:
            raise ModelDiagnosticsValidationError(
                f"Fit_Parameters contains a non-numeric fitted value for '{pid}'."
            ) from exc

    observed_set = {obs.species_id for obs in problem.observations}
    measured = tuple(sid for sid in problem.model.species_ids if sid in observed_set)
    return CompletedFitInfo(problem=problem, summary=summary, fitted_parameters=fitted, measured_species=measured)


def _as_fraction(value: float, max_denominator: int = 10000) -> Fraction:
    x = float(value)
    if abs(x) < 1e-12:
        return Fraction(0)
    return Fraction(str(x)).limit_denominator(max_denominator)


def _rref_fraction(matrix: np.ndarray) -> tuple[list[list[Fraction]], tuple[int, ...]]:
    a = [[_as_fraction(v) for v in row] for row in np.asarray(matrix, dtype=float)]
    if not a:
        return [], tuple()
    m, n = len(a), len(a[0])
    pivots: list[int] = []
    r = 0
    for c in range(n):
        pivot = next((i for i in range(r, m) if a[i][c] != 0), None)
        if pivot is None:
            continue
        a[r], a[pivot] = a[pivot], a[r]
        p = a[r][c]
        a[r] = [v / p for v in a[r]]
        for i in range(m):
            if i == r or a[i][c] == 0:
                continue
            f = a[i][c]
            a[i] = [a[i][j] - f * a[r][j] for j in range(n)]
        pivots.append(c)
        r += 1
        if r == m:
            break
    return a, tuple(pivots)


def exact_rank(matrix: np.ndarray) -> int:
    _rref, pivots = _rref_fraction(np.asarray(matrix, dtype=float))
    return len(pivots)


def _integerize(vector: Iterable[Fraction | int | float]) -> tuple[int, ...]:
    vals = [v if isinstance(v, Fraction) else _as_fraction(float(v)) for v in vector]
    lcm = 1
    for v in vals:
        lcm = abs(lcm * v.denominator) // math.gcd(lcm, v.denominator)
    integers = [int(v * lcm) for v in vals]
    gcd = 0
    for v in integers:
        gcd = math.gcd(gcd, abs(v))
    if gcd:
        integers = [v // gcd for v in integers]
    for v in integers:
        if v != 0:
            if v < 0:
                integers = [-x for x in integers]
            break
    return tuple(integers)


def _raw_integer_nullspace(matrix: np.ndarray) -> list[tuple[int, ...]]:
    """Return an exact integer basis for null(matrix)."""
    a, pivots = _rref_fraction(np.asarray(matrix, dtype=float))
    if not a:
        n = int(np.asarray(matrix).shape[1]) if np.asarray(matrix).ndim == 2 else 0
        return [tuple(1 if i == j else 0 for i in range(n)) for j in range(n)]
    n = len(a[0])
    free = [j for j in range(n) if j not in pivots]
    basis: list[tuple[int, ...]] = []
    for free_col in free:
        v = [Fraction(0) for _ in range(n)]
        v[free_col] = Fraction(1)
        for row_i, pivot_col in enumerate(pivots):
            v[pivot_col] = -a[row_i][free_col]
        basis.append(_integerize(v))
    return basis


def _normalise_integer_vector(vector: np.ndarray | tuple[int, ...]) -> tuple[int, ...]:
    vals = [int(v) for v in vector]
    gcd = 0
    for v in vals:
        gcd = math.gcd(gcd, abs(v))
    if gcd:
        vals = [v // gcd for v in vals]
    for v in vals:
        if v != 0:
            if v < 0:
                vals = [-x for x in vals]
            break
    return tuple(vals)


def _simple_integer_basis(raw_basis: list[tuple[int, ...]]) -> list[tuple[int, ...]]:
    """Prefer sparse, same-sign, small-integer combinations of an exact basis.

    Null-space bases are not unique. For teaching and diagnostics a sparse
    non-negative conservation relation is usually much easier to read than a
    dense algebraic basis. For small nullities we search low-order integer
    combinations and retain an independent set. The represented invariant
    space is unchanged.
    """
    q = len(raw_basis)
    if q <= 1:
        return raw_basis
    b = np.asarray(raw_basis, dtype=int)
    if q > 6:
        return raw_basis

    candidates: set[tuple[int, ...]] = set()
    for coeff in product(range(-2, 3), repeat=q):
        if not any(coeff):
            continue
        v = np.asarray(coeff, dtype=int) @ b
        if np.all(v == 0):
            continue
        candidates.add(_normalise_integer_vector(v))

    def score(v: tuple[int, ...]) -> tuple:
        nz = [x for x in v if x != 0]
        mixed_sign = int(any(x < 0 for x in nz) and any(x > 0 for x in nz))
        return (mixed_sign, len(nz), sum(abs(x) for x in v), max(abs(x) for x in v), v)

    selected: list[tuple[int, ...]] = []
    for candidate in sorted(candidates, key=score):
        if len(selected) >= q:
            break
        trial = np.asarray(selected + [candidate], dtype=float)
        if np.linalg.matrix_rank(trial, tol=1e-10) > len(selected):
            selected.append(candidate)
    return selected if len(selected) == q else raw_basis


def invariant_basis(stoichiometric_matrix: np.ndarray) -> tuple[tuple[int, ...], ...]:
    """Return a readable independent basis of l such that l^T S = 0."""
    s = np.asarray(stoichiometric_matrix, dtype=float)
    if s.ndim != 2:
        raise ValueError("Stoichiometric matrix must be two-dimensional.")
    raw = _raw_integer_nullspace(s.T)
    return tuple(_simple_integer_basis(raw))


def measured_only_invariant_basis(
    stoichiometric_matrix: np.ndarray,
    species_ids: tuple[str, ...],
    measured_species: tuple[str, ...],
) -> tuple[tuple[int, ...], ...]:
    """Return all independent invariants that can be evaluated using measured species only."""
    s = np.asarray(stoichiometric_matrix, dtype=float)
    measured_set = set(measured_species)
    measured_indices = [i for i, sid in enumerate(species_ids) if sid in measured_set]
    if not measured_indices:
        return tuple()
    reduced = s[measured_indices, :]
    reduced_basis = invariant_basis(reduced)
    full_basis: list[tuple[int, ...]] = []
    for v in reduced_basis:
        full = [0 for _ in species_ids]
        for idx, coeff in zip(measured_indices, v):
            full[idx] = int(coeff)
        full_basis.append(tuple(full))
    return tuple(full_basis)


def simple_nonnegative_integer_invariants(
    stoichiometric_matrix: np.ndarray,
    max_coefficient: int = 4,
    max_candidates: int = 500_000,
) -> tuple[tuple[int, ...], ...]:
    """Enumerate primitive non-negative integer invariants within a small coefficient box.

    These relationships do not add independent information beyond the left-null
    space. They are generated because small same-sign integer totals are often
    much easier to interpret chemically and experimentally than an arbitrary
    algebraic basis. Scaled duplicates are removed by retaining primitive
    vectors only (gcd of non-zero coefficients = 1).
    """
    s = np.asarray(stoichiometric_matrix, dtype=float)
    if s.ndim != 2:
        raise ValueError("Stoichiometric matrix must be two-dimensional.")
    max_coefficient = int(max_coefficient)
    if max_coefficient < 1:
        raise ValueError("max_coefficient must be at least 1.")
    n_species = int(s.shape[0])
    search_size = (max_coefficient + 1) ** n_species
    if search_size > int(max_candidates):
        raise ModelDiagnosticsValidationError(
            f"Simple-integer invariant enumeration would require {search_size:,} candidate vectors. "
            f"Reduce the coefficient cap or use the independent invariant basis for this {n_species}-species model."
        )

    found: list[tuple[int, ...]] = []
    for vector in product(range(max_coefficient + 1), repeat=n_species):
        if not any(vector):
            continue
        gcd = 0
        for value in vector:
            gcd = math.gcd(gcd, abs(int(value)))
        if gcd != 1:
            continue
        if np.allclose(np.asarray(vector, dtype=float) @ s, 0.0, rtol=0.0, atol=1e-10):
            found.append(tuple(int(v) for v in vector))

    def score(v: tuple[int, ...]) -> tuple:
        nonzero = sum(1 for x in v if x != 0)
        return (nonzero, sum(v), max(v), v)

    return tuple(sorted(found, key=score))


def _aggregate_observations(
    info: CompletedFitInfo,
    required_species: tuple[str, ...],
) -> dict[tuple[str, float, str], tuple[float, float | None]]:
    """Average duplicate observations and combine supplied sigma for the mean."""
    raw: dict[tuple[str, float, str], list[tuple[float, float | None]]] = {}
    required = set(required_species)
    for obs in info.problem.observations:
        if obs.species_id not in required:
            continue
        raw.setdefault((obs.dataset_id, float(obs.time_s), obs.species_id), []).append(
            (float(obs.measured_value), None if obs.sigma is None else float(obs.sigma))
        )
    aggregated: dict[tuple[str, float, str], tuple[float, float | None]] = {}
    for key, rows in raw.items():
        values = np.asarray([x[0] for x in rows], dtype=float)
        supplied = [x[1] for x in rows]
        sigma: float | None = None
        if supplied and all(x is not None and np.isfinite(x) and x > 0 for x in supplied):
            sigma = float(np.sqrt(np.sum(np.square(np.asarray(supplied, dtype=float)))) / len(supplied))
        aggregated[key] = (float(np.mean(values)), sigma)
    return aggregated


def effective_measurement_sigma(
    measured_value: float,
    supplied_sigma: float | None,
    assumed_relative_fraction: float,
    absolute_floor: float,
) -> float:
    """Return the 1-sigma measurement uncertainty used by diagnostics.

    A positive workbook sigma takes precedence. Otherwise KineticFittingPot combines
    an assumed relative term with an absolute floor in quadrature.
    """
    if supplied_sigma is not None and np.isfinite(supplied_sigma) and supplied_sigma > 0:
        return float(supplied_sigma)
    rel = max(0.0, float(assumed_relative_fraction))
    floor = max(0.0, float(absolute_floor))
    return float(np.sqrt((rel * abs(float(measured_value))) ** 2 + floor ** 2))


def balance_points_with_uncertainty(
    info: CompletedFitInfo,
    coefficients: tuple[int, ...],
    assumed_relative_fraction: float = 0.03,
    absolute_floor: float = 0.0,
) -> tuple[BalanceUncertaintyPoint, ...]:
    """Evaluate a fully observed invariant and propagate independent measurement uncertainty.

    The uncertainty on B(t)-B(t0) is propagated as sqrt(sigma_B(t)^2 +
    sigma_B(t0)^2). The observed baseline magnitude is then used only as the
    normalising scale. At t0 the difference is identically zero because the same
    observation defines both terms.
    """
    species_ids = info.problem.model.species_ids
    required = tuple(sid for sid, c in zip(species_ids, coefficients) if c != 0)
    if not required:
        return tuple()
    measured = set(info.measured_species)
    if any(sid not in measured for sid in required):
        raise ModelDiagnosticsValidationError(
            "This invariant contains unmeasured species and cannot be evaluated directly from raw observations."
        )

    dataset_map = {d.dataset_id: d for d in info.problem.datasets}
    values = _aggregate_observations(info, required)
    result: list[BalanceUncertaintyPoint] = []
    for dataset_id, dataset in dataset_map.items():
        if dataset.volume_model != "constant":
            continue
        times = sorted({time for ds, time, _sid in values if ds == dataset_id})
        complete_times = [
            t for t in times if all((dataset_id, t, sid) in values for sid in required)
        ]
        if not complete_times:
            continue
        rows: list[tuple[float, float, float]] = []
        for t in complete_times:
            total = 0.0
            variance = 0.0
            for sid, coeff in zip(species_ids, coefficients):
                if coeff == 0:
                    continue
                concentration, supplied_sigma = values[(dataset_id, t, sid)]
                sigma_c = effective_measurement_sigma(
                    concentration, supplied_sigma, assumed_relative_fraction, absolute_floor
                )
                scale = float(coeff) * float(dataset.volume_m3)
                total += scale * concentration
                variance += (scale * sigma_c) ** 2
            rows.append((float(t), float(total), float(np.sqrt(max(0.0, variance)))))
        baseline = rows[0][1]
        baseline_sigma = rows[0][2]
        denom = abs(baseline)
        for index, (time_s, value, sigma_b) in enumerate(rows):
            deviation = None if denom <= 1e-15 else (value - baseline) / denom
            if denom <= 1e-15:
                deviation_sigma = None
            elif index == 0:
                deviation_sigma = 0.0
            else:
                deviation_sigma = float(np.sqrt(sigma_b ** 2 + baseline_sigma ** 2) / denom)
            result.append(BalanceUncertaintyPoint(
                dataset_id=dataset_id,
                time_s=time_s,
                balance_value=value,
                baseline_value=baseline,
                normalized_deviation=deviation,
                balance_sigma=sigma_b,
                baseline_sigma=baseline_sigma,
                normalized_deviation_sigma=deviation_sigma,
            ))
    return tuple(result)


def reconcile_observations_for_invariant(
    info: CompletedFitInfo,
    coefficients: tuple[int, ...],
    assumed_relative_fraction: float = 0.03,
    absolute_floor: float = 0.0,
) -> tuple[ReconciledObservation, ...]:
    """Reconcile raw observations to one fully observed invariant.

    For each constant-volume dataset the earliest complete measured balance is
    treated as the fixed target. Later measurements are adjusted by the minimum
    diagonal-weighted least-squares correction subject to l^T x = target. Raw
    observations are never modified.
    """
    species_ids = info.problem.model.species_ids
    required = tuple(sid for sid, c in zip(species_ids, coefficients) if c != 0)
    coeff_map = {sid: float(c) for sid, c in zip(species_ids, coefficients) if c != 0}
    if not required:
        return tuple()
    measured = set(info.measured_species)
    if any(sid not in measured for sid in required):
        raise ModelDiagnosticsValidationError(
            "Data reconciliation requires a fully observed invariant."
        )

    dataset_map = {d.dataset_id: d for d in info.problem.datasets}
    values = _aggregate_observations(info, required)
    output: list[ReconciledObservation] = []
    for dataset_id, dataset in dataset_map.items():
        if dataset.volume_model != "constant":
            continue
        times = sorted({time for ds, time, _sid in values if ds == dataset_id})
        complete_times = [t for t in times if all((dataset_id, t, sid) in values for sid in required)]
        if not complete_times:
            continue

        first_t = complete_times[0]
        y0 = np.asarray([values[(dataset_id, first_t, sid)][0] for sid in required], dtype=float)
        lvec = np.asarray([coeff_map[sid] for sid in required], dtype=float)
        target_concentration_balance = float(lvec @ y0)
        target_mole_balance = target_concentration_balance * float(dataset.volume_m3)

        for t in complete_times:
            y = np.asarray([values[(dataset_id, t, sid)][0] for sid in required], dtype=float)
            sigmas = np.asarray([
                effective_measurement_sigma(
                    values[(dataset_id, t, sid)][0],
                    values[(dataset_id, t, sid)][1],
                    assumed_relative_fraction,
                    absolute_floor,
                )
                for sid in required
            ], dtype=float)
            positive = sigmas > 0
            if not np.any(positive):
                sigmas = np.ones_like(sigmas)
            else:
                # A zero uncertainty would otherwise make a channel infinitely
                # precise. Use a very small positive scale only for channels
                # explicitly assigned zero while other channels have uncertainty.
                min_positive = float(np.min(sigmas[positive]))
                sigmas = np.where(sigmas > 0, sigmas, max(1e-12, min_positive * 1e-6))

            variance = np.square(sigmas)
            denom = float(np.sum(np.square(lvec) * variance))
            balance_before_c = float(lvec @ y)
            if denom <= 0 or not np.isfinite(denom):
                x = y.copy()
            else:
                multiplier = (balance_before_c - target_concentration_balance) / denom
                x = y - variance * lvec * multiplier
            correction = x - y
            standardized = np.divide(
                correction, sigmas, out=np.full_like(correction, np.nan), where=sigmas > 0
            )
            weighted = float(np.nansum(np.square(standardized))) if np.any(np.isfinite(standardized)) else None
            balance_after_c = float(lvec @ x)
            for idx, sid in enumerate(required):
                output.append(ReconciledObservation(
                    dataset_id=dataset_id,
                    time_s=float(t),
                    species_id=sid,
                    measured_value=float(y[idx]),
                    reconciled_value=float(x[idx]),
                    correction=float(correction[idx]),
                    sigma=float(sigmas[idx]),
                    standardized_correction=None if not np.isfinite(standardized[idx]) else float(standardized[idx]),
                    balance_before=balance_before_c * float(dataset.volume_m3),
                    balance_after=balance_after_c * float(dataset.volume_m3),
                    target_balance=target_mole_balance,
                    weighted_inconsistency=weighted,
                ))
    return tuple(output)


def residual_diagnostic_points(
    workbook_bytes: bytes,
    info: CompletedFitInfo,
    assumed_relative_fraction: float = 0.03,
    absolute_floor: float = 0.002,
) -> tuple[ResidualDiagnosticPoint, ...]:
    """Read completed-fit predictions and construct residual diagnostic records.

    Residual sign is kept identical to KineticFittingPot kinetic fitting:
    fitted/model value minus measured value. Positive residual therefore means
    that the fitted model lies above the observation.

    Two distinct normalisations are retained:
    - ``fit_scaled_residual`` reproduces the species scale used by the
      ``species_scaled`` fitting objective (maximum absolute measured value by
      species); it is a numerical fit scale, not a statistical standard error.
    - ``uncertainty_standardized_residual`` divides by supplied measurement
      sigma where available, otherwise by the explicit relative-plus-floor
      diagnostic uncertainty model used elsewhere on the page.
    """
    if load_workbook is None:
        raise RuntimeError("Post-fit residual diagnostics requires openpyxl.")
    try:
        wb = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    except Exception as exc:
        raise ModelDiagnosticsValidationError(
            f"KineticFittingPot could not read residual diagnostics from this workbook: {exc}"
        ) from exc
    if "Fit_Predictions" not in wb.sheetnames:
        raise ModelDiagnosticsValidationError("Completed workbook is missing Fit_Predictions.")

    ws = wb["Fit_Predictions"]
    headers = _header_map(ws)
    required_headers = (
        "Dataset_ID",
        "Time_s",
        "Species_ID",
        "Measured_Value_kmol_m3",
        "Fitted_Value_kmol_m3",
    )
    missing = [h for h in required_headers if h not in headers]
    if missing:
        raise ModelDiagnosticsValidationError(
            "Fit_Predictions is missing required column(s): " + ", ".join(missing)
        )

    aggregate = _aggregate_observations(info, tuple(info.measured_species))
    scales = species_scales(info.problem)
    output: list[ResidualDiagnosticPoint] = []
    for row in range(2, ws.max_row + 1):
        ds = ws.cell(row=row, column=headers["Dataset_ID"]).value
        time_value = ws.cell(row=row, column=headers["Time_s"]).value
        sid = ws.cell(row=row, column=headers["Species_ID"]).value
        measured_value = ws.cell(row=row, column=headers["Measured_Value_kmol_m3"]).value
        fitted_value = ws.cell(row=row, column=headers["Fitted_Value_kmol_m3"]).value
        if ds in (None, "") or sid in (None, ""):
            continue
        try:
            time_s = float(time_value)
            measured = float(measured_value)
            fitted = float(fitted_value)
        except (TypeError, ValueError) as exc:
            raise ModelDiagnosticsValidationError(
                f"Fit_Predictions row {row} contains a non-numeric time/measured/fitted value."
            ) from exc
        dataset_id = str(ds).strip()
        species_id = str(sid).strip()
        residual = fitted - measured
        supplied_sigma = None
        key = (dataset_id, time_s, species_id)
        if key in aggregate:
            supplied_sigma = aggregate[key][1]
        sigma = effective_measurement_sigma(
            measured, supplied_sigma, assumed_relative_fraction, absolute_floor
        )
        z = None if sigma <= 0 or not np.isfinite(sigma) else float(residual / sigma)
        scale = float(scales.get(species_id, max(abs(measured), 1e-6)))
        fit_scaled = None if scale <= 0 or not np.isfinite(scale) else float(residual / scale)
        output.append(
            ResidualDiagnosticPoint(
                dataset_id=dataset_id,
                time_s=time_s,
                species_id=species_id,
                measured_value=measured,
                fitted_value=fitted,
                residual_model_minus_measured=float(residual),
                supplied_sigma=supplied_sigma,
                effective_sigma=float(sigma),
                uncertainty_standardized_residual=z,
                species_scale=scale,
                fit_scaled_residual=fit_scaled,
            )
        )
    return tuple(output)



def local_parameter_diagnostics(
    info: CompletedFitInfo,
    *,
    relative_step: float = 1e-4,
    span_step_floor: float = 1e-7,
) -> LocalParameterDiagnosticsResult:
    """Calculate local Jacobian and parameter-correlation diagnostics at the fitted optimum.

    The Jacobian is formed from the *same weighted residual vector* used by the
    completed kinetic fit. Central finite differences are used whenever the
    fitted point is sufficiently far from both parameter bounds; one-sided
    differences are used at/near a bound.

    Singular values and the reported condition number require a parameter
    scale. KineticFittingPot uses a local relative scale ``abs(theta_hat)`` whenever
    it is safely non-zero; parameters at/near zero use 1% of their admissible
    bound span instead. This makes the geometry interpretable as sensitivity to
    fractional/local parameter movements without pretending that condition
    numbers are invariant to parameter units.

    Approximate covariance/standard errors use the usual local least-squares
    linearisation of the selected fitting objective. They are objective-based
    approximations, not profile-likelihood intervals and not formal statistical
    confidence intervals unless the residual weighting correctly represents the
    experimental error model.
    """
    problem = info.problem
    parameter_ids = tuple(problem.fit_parameter_ids)
    if not parameter_ids:
        raise ModelDiagnosticsValidationError("No fitted parameters are available for local diagnostics.")

    try:
        values = np.asarray([float(info.fitted_parameters[pid]) for pid in parameter_ids], dtype=float)
    except KeyError as exc:
        raise ModelDiagnosticsValidationError(
            f"Completed workbook is missing fitted value for parameter '{exc.args[0]}'."
        ) from exc
    if not np.all(np.isfinite(values)):
        raise ModelDiagnosticsValidationError("Fitted parameter values must be finite for local diagnostics.")

    weighting = str(info.summary.get("Weighting", "species_scaled") or "species_scaled").strip().lower()
    if weighting not in {"unweighted", "species_scaled"}:
        raise ModelDiagnosticsValidationError(
            f"Unsupported completed-fit weighting '{weighting}' for local diagnostics."
        )

    lower = np.asarray(problem.lower_bounds, dtype=float)
    upper = np.asarray(problem.upper_bounds, dtype=float)
    if values.shape != lower.shape or values.shape != upper.shape:
        raise ModelDiagnosticsValidationError("Fitted-parameter dimensions do not match the original fitting problem.")
    spans = upper - lower
    if np.any(~np.isfinite(spans)) or np.any(spans <= 0):
        raise ModelDiagnosticsValidationError("Local diagnostics require finite, ordered parameter bounds.")

    try:
        base_residual = np.asarray(residual_vector(problem, values, weighting), dtype=float)
    except Exception as exc:
        raise ModelDiagnosticsValidationError(
            f"KineticFittingPot could not evaluate the residual vector at the fitted optimum: {exc}"
        ) from exc
    if not np.all(np.isfinite(base_residual)):
        raise ModelDiagnosticsValidationError("Residual vector at the fitted optimum contains non-finite values.")

    n_obs = int(base_residual.size)
    n_par = int(values.size)
    jac = np.empty((n_obs, n_par), dtype=float)
    steps: list[float] = []
    methods: list[str] = []

    for j in range(n_par):
        theta = float(values[j])
        span = float(spans[j])
        # Stable local perturbation: mostly relative to the fitted value, with a
        # very small fraction of the admissible range as a floor for values at 0.
        target = max(float(relative_step) * abs(theta), float(span_step_floor) * span, 1e-14)
        room_minus = max(0.0, theta - float(lower[j]))
        room_plus = max(0.0, float(upper[j]) - theta)
        central_h = min(target, 0.49 * room_minus, 0.49 * room_plus)

        try:
            if central_h > max(1e-14, 1e-12 * span):
                plus = values.copy(); minus = values.copy()
                plus[j] += central_h; minus[j] -= central_h
                r_plus = np.asarray(residual_vector(problem, plus, weighting), dtype=float)
                r_minus = np.asarray(residual_vector(problem, minus, weighting), dtype=float)
                derivative = (r_plus - r_minus) / (2.0 * central_h)
                h_used = central_h
                method = "central"
            elif room_plus > max(1e-14, 1e-12 * span):
                h = min(target, 0.49 * room_plus)
                plus = values.copy(); plus[j] += h
                r_plus = np.asarray(residual_vector(problem, plus, weighting), dtype=float)
                derivative = (r_plus - base_residual) / h
                h_used = h
                method = "forward"
            elif room_minus > max(1e-14, 1e-12 * span):
                h = min(target, 0.49 * room_minus)
                minus = values.copy(); minus[j] -= h
                r_minus = np.asarray(residual_vector(problem, minus, weighting), dtype=float)
                derivative = (base_residual - r_minus) / h
                h_used = h
                method = "backward"
            else:
                raise ModelDiagnosticsValidationError(
                    f"Parameter '{parameter_ids[j]}' has no usable finite-difference room inside its bounds."
                )
        except ModelDiagnosticsValidationError:
            raise
        except Exception as exc:
            raise ModelDiagnosticsValidationError(
                f"Finite-difference sensitivity failed for parameter '{parameter_ids[j]}': {exc}"
            ) from exc

        if derivative.shape != base_residual.shape or not np.all(np.isfinite(derivative)):
            raise ModelDiagnosticsValidationError(
                f"Finite-difference sensitivity for parameter '{parameter_ids[j]}' contains non-finite values."
            )
        jac[:, j] = derivative
        steps.append(float(h_used))
        methods.append(method)

    # Local parameter scaling for singular values/condition number. A pure raw
    # Jacobian condition number is dependent on parameter units, so use fitted
    # magnitude except when that magnitude is effectively zero.
    parameter_scales = np.empty(n_par, dtype=float)
    for j, theta in enumerate(values):
        span = float(spans[j])
        near_zero_threshold = max(1e-12, 1e-10 * span)
        parameter_scales[j] = abs(float(theta)) if abs(float(theta)) > near_zero_threshold else 0.01 * span
        parameter_scales[j] = max(parameter_scales[j], 1e-14)
    scaled_jac = jac * parameter_scales[np.newaxis, :]

    try:
        _u, singular, vt = np.linalg.svd(scaled_jac, full_matrices=False)
    except np.linalg.LinAlgError as exc:
        raise ModelDiagnosticsValidationError(f"SVD of the local Jacobian failed: {exc}") from exc
    singular = np.asarray(singular, dtype=float)
    if singular.size:
        tol = max(scaled_jac.shape) * np.finfo(float).eps * float(singular[0])
        numerical_rank = int(np.sum(singular > tol))
        condition = float(singular[0] / singular[-1]) if singular[-1] > tol else float("inf")
        weakest = np.asarray(vt[-1], dtype=float)
        if np.max(np.abs(weakest)) > 0:
            weakest = weakest / np.max(np.abs(weakest))
    else:
        numerical_rank = 0
        condition = float("inf")
        weakest = np.full(n_par, np.nan)

    sse = float(np.dot(base_residual, base_residual))
    dof = int(n_obs - n_par)
    residual_variance = float(sse / dof) if dof > 0 else None

    normal = jac.T @ jac
    try:
        normal_pinv = np.linalg.pinv(normal, rcond=1e-12)
    except np.linalg.LinAlgError as exc:
        raise ModelDiagnosticsValidationError(f"Local covariance calculation failed: {exc}") from exc
    if residual_variance is None:
        covariance = np.full((n_par, n_par), np.nan)
    else:
        covariance = float(residual_variance) * normal_pinv

    diag_cov = np.diag(covariance).astype(float)
    se = np.sqrt(np.maximum(diag_cov, 0.0)) if np.all(np.isfinite(diag_cov)) else np.full(n_par, np.nan)
    denom = np.sqrt(np.outer(diag_cov, diag_cov))
    correlation = np.divide(
        covariance,
        denom,
        out=np.full_like(covariance, np.nan, dtype=float),
        where=np.isfinite(denom) & (denom > 0),
    )
    for j in range(n_par):
        if np.isfinite(correlation[j, j]):
            correlation[j, j] = 1.0

    relative_sensitivity = np.sqrt(np.mean(np.square(scaled_jac), axis=0))
    rows: list[LocalParameterDiagnostic] = []
    for j, pid in enumerate(parameter_ids):
        theta = float(values[j])
        span = float(spans[j])
        position = 100.0 * (theta - float(lower[j])) / span
        nearest_distance = min(theta - float(lower[j]), float(upper[j]) - theta)
        nearest = 100.0 * nearest_distance / float(parameter_scales[j])
        if n_par > 1 and np.any(np.isfinite(correlation[j, :])):
            candidates = [(k, float(correlation[j, k])) for k in range(n_par) if k != j and np.isfinite(correlation[j, k])]
            if candidates:
                k_best, corr_best = max(candidates, key=lambda kv: abs(kv[1]))
                partner = parameter_ids[k_best]
            else:
                corr_best = None; partner = None
        else:
            corr_best = None; partner = None
        se_j = float(se[j]) if j < len(se) and np.isfinite(se[j]) else None
        rel_se = None
        if se_j is not None and abs(theta) > max(1e-14, 1e-12 * span):
            rel_se = 100.0 * se_j / abs(theta)
        rows.append(LocalParameterDiagnostic(
            parameter_id=pid,
            fitted_value=theta,
            lower_bound=float(lower[j]),
            upper_bound=float(upper[j]),
            finite_difference_step=float(steps[j]),
            finite_difference_method=methods[j],
            local_parameter_scale=float(parameter_scales[j]),
            relative_sensitivity_rms=float(relative_sensitivity[j]),
            one_percent_effect_rms=float(0.01 * relative_sensitivity[j]),
            approximate_standard_error=se_j,
            relative_standard_error_percent=rel_se,
            bound_position_percent=float(position),
            nearest_bound_distance_percent=float(nearest),
            strongest_correlation_parameter=partner,
            strongest_correlation=corr_best,
        ))

    return LocalParameterDiagnosticsResult(
        weighting=weighting,
        parameter_ids=parameter_ids,
        fitted_values=tuple(float(v) for v in values),
        weighted_residuals=tuple(float(v) for v in base_residual),
        jacobian=jac,
        scaled_jacobian=scaled_jac,
        parameter_scales=tuple(float(v) for v in parameter_scales),
        singular_values=tuple(float(v) for v in singular),
        condition_number=condition,
        numerical_rank=numerical_rank,
        degrees_of_freedom=dof,
        objective_sse=sse,
        residual_variance_estimate=residual_variance,
        covariance=covariance,
        correlation=correlation,
        weakest_direction=tuple(float(v) for v in weakest),
        parameter_rows=tuple(rows),
    )

def time_grid_regularity(
    times: Iterable[float],
    tolerance_fraction: float = 0.05,
) -> dict[str, object]:
    """Assess whether a time grid is sufficiently regular for a conventional ACF.

    The criterion is intentionally transparent: at least six unique time points
    are required and every interval must lie within ``tolerance_fraction`` of
    the median positive interval. KineticFittingPot does not interpolate an irregular
    residual sequence simply to make an ACF available.
    """
    unique = np.asarray(sorted({float(t) for t in times if np.isfinite(float(t))}), dtype=float)
    if len(unique) < 6:
        return {
            "regular": False,
            "reason": "At least six unique observation times are required for a useful residual ACF.",
            "times": unique,
            "intervals": np.diff(unique),
            "median_interval": None,
            "maximum_relative_interval_deviation": None,
        }
    intervals = np.diff(unique)
    positive = intervals[intervals > 0]
    if len(positive) != len(intervals) or len(positive) == 0:
        return {
            "regular": False,
            "reason": "Observation times must increase strictly for residual ACF analysis.",
            "times": unique,
            "intervals": intervals,
            "median_interval": None,
            "maximum_relative_interval_deviation": None,
        }
    median = float(np.median(positive))
    if median <= 0:
        regular = False
        max_rel = None
    else:
        max_rel = float(np.max(np.abs(positive - median) / median))
        regular = bool(max_rel <= float(tolerance_fraction) + 1e-12)
    reason = (
        f"Observation intervals are sufficiently regular (maximum deviation from the median interval {100.0 * max_rel:.1f}%)."
        if regular and max_rel is not None
        else f"Observation intervals are irregular (maximum deviation from the median interval {100.0 * max_rel:.1f}%)."
        if max_rel is not None
        else "Observation intervals are irregular."
    )
    return {
        "regular": regular,
        "reason": reason,
        "times": unique,
        "intervals": intervals,
        "median_interval": median,
        "maximum_relative_interval_deviation": max_rel,
    }


def residual_autocorrelation(values: Iterable[float], max_lag: int | None = None) -> tuple[tuple[int, float], ...]:
    """Return conventional mean-centred sample autocorrelation values.

    This routine assumes the caller has already verified a sufficiently regular
    time grid. Lag zero is included and equals one when variance is non-zero.
    """
    x = np.asarray([float(v) for v in values if np.isfinite(float(v))], dtype=float)
    n = len(x)
    if n < 2:
        return tuple()
    centred = x - float(np.mean(x))
    denom = float(np.dot(centred, centred))
    if denom <= 0:
        return tuple((lag, 1.0 if lag == 0 else 0.0) for lag in range(0, min(n - 1, max_lag or n - 1) + 1))
    if max_lag is None:
        max_lag = min(10, n // 2)
    max_lag = max(0, min(int(max_lag), n - 1))
    values_out: list[tuple[int, float]] = []
    for lag in range(max_lag + 1):
        numerator = float(np.dot(centred[: n - lag], centred[lag:])) if lag else denom
        values_out.append((lag, numerator / denom))
    return tuple(values_out)


def classify_invariants(
    basis: tuple[tuple[int, ...], ...],
    species_ids: tuple[str, ...],
    measured_species: tuple[str, ...],
) -> tuple[InvariantResult, ...]:
    measured = set(measured_species)
    results: list[InvariantResult] = []
    for coeffs in basis:
        participating = tuple(sid for sid, c in zip(species_ids, coeffs) if c != 0)
        missing = tuple(sid for sid in participating if sid not in measured)
        results.append(
            InvariantResult(
                coefficients=tuple(int(c) for c in coeffs),
                participating_species=participating,
                fully_observed=not missing,
                missing_species=missing,
            )
        )
    return tuple(results)


def format_invariant(coefficients: tuple[int, ...], species_ids: tuple[str, ...]) -> str:
    terms: list[str] = []
    for sid, coeff in zip(species_ids, coefficients):
        c = int(coeff)
        if c == 0:
            continue
        mag = abs(c)
        core = sid if mag == 1 else f"{mag}{sid}"
        if not terms:
            terms.append(core if c > 0 else f"-{core}")
        else:
            terms.append((" + " if c > 0 else " - ") + core)
    return "".join(terms) + " = constant" if terms else "0 = constant"


def stoichiometric_diagnostics(info: CompletedFitInfo) -> dict:
    model = info.problem.model
    s = np.asarray(model.stoichiometric_matrix, dtype=float)
    rank = exact_rank(s)
    full_basis = invariant_basis(s)
    measured_basis = measured_only_invariant_basis(s, model.species_ids, info.measured_species)
    classified = classify_invariants(full_basis, model.species_ids, info.measured_species)
    # Put directly measured invariants first without changing the represented basis.
    classified = tuple(sorted(classified, key=lambda x: (not x.fully_observed, len(x.participating_species), x.coefficients)))
    return {
        "rank": rank,
        "reaction_columns": int(s.shape[1]),
        "species_count": int(s.shape[0]),
        "invariant_dimension": int(s.shape[0] - rank),
        "full_basis": full_basis,
        "classified_full_basis": classified,
        "measured_basis": measured_basis,
        "measured_invariant_dimension": len(measured_basis),
    }


def balance_points_for_invariant(
    info: CompletedFitInfo,
    coefficients: tuple[int, ...],
) -> tuple[BalancePoint, ...]:
    """Evaluate a fully observed invariant directly from raw measurements.

    This first implementation deliberately supports direct concentration-to-mole
    monitoring for constant-volume datasets only. Open/variable-volume systems
    require additional process-boundary information and are not silently treated
    as closed concentration balances.
    """
    species_ids = info.problem.model.species_ids
    required = tuple(sid for sid, c in zip(species_ids, coefficients) if c != 0)
    if not required:
        return tuple()
    measured = set(info.measured_species)
    if any(sid not in measured for sid in required):
        raise ModelDiagnosticsValidationError(
            "This invariant contains unmeasured species and cannot be evaluated directly from raw observations."
        )

    dataset_map = {d.dataset_id: d for d in info.problem.datasets}
    # Multiple rows at the same dataset/time/species are averaged explicitly;
    # the UI reports that behaviour if it occurs.
    values: dict[tuple[str, float, str], list[float]] = {}
    for obs in info.problem.observations:
        if obs.species_id not in required:
            continue
        values.setdefault((obs.dataset_id, float(obs.time_s), obs.species_id), []).append(float(obs.measured_value))

    result: list[BalancePoint] = []
    for dataset_id, dataset in dataset_map.items():
        if dataset.volume_model != "constant":
            continue
        times = sorted({time for ds, time, _sid in values if ds == dataset_id})
        complete_times = [
            t for t in times
            if all((dataset_id, t, sid) in values for sid in required)
        ]
        if not complete_times:
            continue
        balances: list[tuple[float, float]] = []
        for t in complete_times:
            total = 0.0
            for sid, coeff in zip(species_ids, coefficients):
                if coeff == 0:
                    continue
                obs_values = values[(dataset_id, t, sid)]
                concentration = float(np.mean(obs_values))
                total += float(coeff) * concentration * float(dataset.volume_m3)
            balances.append((float(t), float(total)))
        baseline = balances[0][1]
        denom = abs(baseline)
        for time_s, value in balances:
            deviation = None if denom <= 1e-15 else (value - baseline) / denom
            result.append(BalancePoint(dataset_id, time_s, value, baseline, deviation))
    return tuple(result)


@dataclass(frozen=True)
class ProfilePoint:
    parameter_id: str
    fixed_value: float
    objective_sse: float
    delta_sse: float
    relative_increase_percent: float
    weighted_rmse: float
    success: bool
    nfev: int
    optimized_parameter_values: tuple[float, ...]
    is_zero_boundary_test: bool = False


@dataclass(frozen=True)
class ParameterProfileResult:
    parameter_id: str
    fitted_value: float
    lower_bound: float
    upper_bound: float
    grid_mode: str
    points: tuple[ProfilePoint, ...]
    zero_boundary_point: ProfilePoint | None


@dataclass(frozen=True)
class ProfileDiagnosticsResult:
    weighting: str
    base_objective_sse: float
    base_weighted_rmse: float
    parameter_ids: tuple[str, ...]
    profiles: tuple[ParameterProfileResult, ...]
    improved_solution_found: bool


def _profile_to_unit(values: np.ndarray, lower: np.ndarray, upper: np.ndarray) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    lower = np.asarray(lower, dtype=float)
    upper = np.asarray(upper, dtype=float)
    out = np.empty_like(values)
    positive = (lower > 0) & (upper > 0)
    if np.any(positive):
        clipped = np.clip(values[positive], lower[positive], upper[positive])
        out[positive] = (
            np.log(clipped) - np.log(lower[positive])
        ) / (
            np.log(upper[positive]) - np.log(lower[positive])
        )
    linear = ~positive
    if np.any(linear):
        out[linear] = (values[linear] - lower[linear]) / (upper[linear] - lower[linear])
    return np.clip(out, 0.0, 1.0)


def _profile_from_unit(z: np.ndarray, lower: np.ndarray, upper: np.ndarray) -> np.ndarray:
    z = np.asarray(z, dtype=float)
    lower = np.asarray(lower, dtype=float)
    upper = np.asarray(upper, dtype=float)
    out = np.empty_like(z)
    positive = (lower > 0) & (upper > 0)
    if np.any(positive):
        out[positive] = np.exp(
            np.log(lower[positive])
            + z[positive] * (np.log(upper[positive]) - np.log(lower[positive]))
        )
    linear = ~positive
    if np.any(linear):
        out[linear] = lower[linear] + z[linear] * (upper[linear] - lower[linear])
    return out


def _profile_grid_values(
    fitted_value: float,
    lower_bound: float,
    upper_bound: float,
    *,
    points: int,
    span_factor: float,
) -> tuple[np.ndarray, str]:
    """Construct a readable profile grid around the fitted value.

    Positive parameters are profiled multiplicatively because kinetic constants
    commonly span orders of magnitude. Parameters that cannot use a positive
    log scale are profiled linearly over a local fraction of their admissible
    interval. The exact fitted value is always included.
    """
    n = int(points)
    if n < 5:
        raise ModelDiagnosticsValidationError("Profile grid requires at least 5 points per parameter.")
    if n % 2 == 0:
        n += 1
    factor = float(span_factor)
    if not np.isfinite(factor) or factor <= 1.0:
        raise ModelDiagnosticsValidationError("Profile span factor must be greater than 1.")

    theta = float(fitted_value)
    lo = float(lower_bound)
    hi = float(upper_bound)
    if not (np.isfinite(theta) and np.isfinite(lo) and np.isfinite(hi) and lo < hi):
        raise ModelDiagnosticsValidationError("Profile grid requires finite fitted values and ordered finite bounds.")

    side = (n - 1) // 2
    if theta > 0 and lo >= 0:
        # Do not use zero in the ordinary log grid. A physically meaningful
        # zero-rate boundary is handled explicitly as a model-structure point.
        low_target = max(lo if lo > 0 else theta / factor, theta / factor, theta * 1e-8, 1e-15)
        high_target = min(hi, theta * factor)
        lower_values = np.asarray([], dtype=float)
        upper_values = np.asarray([], dtype=float)
        if low_target < theta * (1.0 - 1e-12):
            lower_values = np.geomspace(low_target, theta, side + 1)[:-1]
        if high_target > theta * (1.0 + 1e-12):
            upper_values = np.geomspace(theta, high_target, side + 1)[1:]
        values = np.concatenate([lower_values, np.asarray([theta]), upper_values])
        mode = "log-relative"
    else:
        span = hi - lo
        half = min(0.25 * span, max(abs(theta), 0.05 * span) * (factor - 1.0))
        low_target = max(lo, theta - half)
        high_target = min(hi, theta + half)
        values = np.linspace(low_target, high_target, n)
        values = np.append(values, theta)
        mode = "linear"

    values = np.asarray(sorted(set(float(v) for v in values if lo - 1e-14 <= v <= hi + 1e-14)), dtype=float)
    if not np.any(np.isclose(values, theta, rtol=1e-12, atol=max(1e-15, abs(theta) * 1e-12))):
        values = np.sort(np.append(values, theta))
    return values, mode


def _optimise_with_fixed_parameter(
    info: CompletedFitInfo,
    base_values: np.ndarray,
    *,
    fixed_index: int,
    fixed_value: float,
    start_values: np.ndarray,
    weighting: str,
    max_nfev: int,
) -> tuple[float, np.ndarray, int, bool]:
    if least_squares is None:
        raise RuntimeError("SciPy is required for profiled-objective diagnostics.")

    problem = info.problem
    n_par = len(problem.fit_parameter_ids)
    free_idx = np.asarray([i for i in range(n_par) if i != int(fixed_index)], dtype=int)
    full = np.asarray(base_values, dtype=float).copy()
    full[int(fixed_index)] = float(fixed_value)

    if free_idx.size == 0:
        res = np.asarray(residual_vector(problem, full, weighting), dtype=float)
        return float(np.sum(res ** 2)), full, 0, True

    lower = np.asarray(problem.lower_bounds, dtype=float)[free_idx]
    upper = np.asarray(problem.upper_bounds, dtype=float)[free_idx]
    start = np.asarray(start_values, dtype=float)[free_idx]
    start = np.clip(start, lower, upper)
    z0 = _profile_to_unit(start, lower, upper)

    def objective(z: np.ndarray) -> np.ndarray:
        values = np.asarray(base_values, dtype=float).copy()
        values[int(fixed_index)] = float(fixed_value)
        values[free_idx] = _profile_from_unit(np.asarray(z, dtype=float), lower, upper)
        try:
            residual = np.asarray(residual_vector(problem, values, weighting), dtype=float)
            if not np.all(np.isfinite(residual)):
                raise ValueError("non-finite residual")
            return residual
        except Exception:
            return np.full(len(problem.observations), 1e6, dtype=float)

    opt = least_squares(
        objective,
        z0,
        bounds=(np.zeros_like(z0), np.ones_like(z0)),
        method="trf",
        max_nfev=int(max_nfev),
        xtol=1e-8,
        ftol=1e-8,
        gtol=1e-8,
    )
    full[free_idx] = _profile_from_unit(opt.x, lower, upper)
    residual = objective(opt.x)
    sse = float(np.sum(residual ** 2))
    success = bool(opt.success) and np.isfinite(sse) and sse < 1e11
    return sse, full, int(opt.nfev), success


def profiled_objective_diagnostics(
    info: CompletedFitInfo,
    *,
    parameter_ids: Iterable[str] | None = None,
    grid_points: int = 11,
    span_factor: float = 20.0,
    max_nfev: int = 160,
    zero_boundary_parameter_ids: Iterable[str] = (),
) -> ProfileDiagnosticsResult:
    """Profile the fitted least-squares objective one parameter at a time.

    For every fixed parameter value, all remaining fitted parameters are
    re-optimised. The ordinary profile grid respects the parameter bounds saved
    in the original fitting workbook. Optional zero-boundary tests deliberately
    permit a selected non-negative kinetic constant to be fixed at exactly zero
    even when the fitting workbook used a small positive numerical lower bound;
    the remaining parameters retain their original bounds.

    The quantity returned is a *profiled objective*. It has a formal likelihood
    interpretation only when the residual weighting corresponds to a defensible
    probabilistic measurement-error model.
    """
    problem = info.problem
    all_ids = tuple(problem.fit_parameter_ids)
    if not all_ids:
        raise ModelDiagnosticsValidationError("No fitted parameters are available for profile diagnostics.")

    selected = tuple(parameter_ids) if parameter_ids is not None else all_ids
    unknown = [pid for pid in selected if pid not in all_ids]
    if unknown:
        raise ModelDiagnosticsValidationError("Unknown fitted parameter(s) requested for profiling: " + ", ".join(unknown))
    if not selected:
        raise ModelDiagnosticsValidationError("Select at least one parameter to profile.")

    zero_selected = set(str(pid) for pid in zero_boundary_parameter_ids)
    unknown_zero = sorted(zero_selected.difference(all_ids))
    if unknown_zero:
        raise ModelDiagnosticsValidationError("Unknown zero-boundary parameter(s): " + ", ".join(unknown_zero))

    try:
        base_values = np.asarray([float(info.fitted_parameters[pid]) for pid in all_ids], dtype=float)
    except KeyError as exc:
        raise ModelDiagnosticsValidationError(
            f"Completed workbook is missing fitted value for parameter '{exc.args[0]}'."
        ) from exc
    if not np.all(np.isfinite(base_values)):
        raise ModelDiagnosticsValidationError("Fitted parameter values must be finite for profile diagnostics.")

    weighting = str(info.summary.get("Weighting", "species_scaled") or "species_scaled").strip().lower()
    if weighting not in {"unweighted", "species_scaled"}:
        raise ModelDiagnosticsValidationError(f"Unsupported completed-fit weighting '{weighting}'.")

    base_residual = np.asarray(residual_vector(problem, base_values, weighting), dtype=float)
    base_sse = float(np.sum(base_residual ** 2))
    base_rmse = float(np.sqrt(np.mean(base_residual ** 2)))
    lower = np.asarray(problem.lower_bounds, dtype=float)
    upper = np.asarray(problem.upper_bounds, dtype=float)

    profiles: list[ParameterProfileResult] = []
    improved = False

    def make_point(pid: str, fixed: float, sse: float, values: np.ndarray, nfev: int, success: bool, zero: bool = False) -> ProfilePoint:
        nonlocal improved
        if np.isfinite(sse) and sse < base_sse * (1.0 - 1e-5):
            improved = True
        delta = float(sse - base_sse)
        rel = float(100.0 * delta / base_sse) if base_sse > 0 else float("nan")
        wrmse = float(np.sqrt(sse / len(problem.observations))) if np.isfinite(sse) and sse >= 0 else float("inf")
        return ProfilePoint(
            parameter_id=pid,
            fixed_value=float(fixed),
            objective_sse=float(sse),
            delta_sse=delta,
            relative_increase_percent=rel,
            weighted_rmse=wrmse,
            success=bool(success),
            nfev=int(nfev),
            optimized_parameter_values=tuple(float(v) for v in values),
            is_zero_boundary_test=bool(zero),
        )

    for pid in selected:
        j = all_ids.index(pid)
        theta = float(base_values[j])
        grid, mode = _profile_grid_values(
            theta,
            float(lower[j]),
            float(upper[j]),
            points=int(grid_points),
            span_factor=float(span_factor),
        )

        # Profile outward from the fitted optimum so each constrained solve can
        # warm-start from the nearest already-optimised point.
        atol = max(1e-15, abs(theta) * 1e-12)
        lower_grid = sorted([float(v) for v in grid if v < theta and not np.isclose(v, theta, rtol=1e-12, atol=atol)], reverse=True)
        upper_grid = sorted([float(v) for v in grid if v > theta and not np.isclose(v, theta, rtol=1e-12, atol=atol)])
        point_map: dict[float, ProfilePoint] = {
            theta: make_point(pid, theta, base_sse, base_values.copy(), 0, True)
        }

        warm = base_values.copy()
        for fixed in lower_grid:
            try:
                sse, values, nfev, success = _optimise_with_fixed_parameter(
                    info,
                    base_values,
                    fixed_index=j,
                    fixed_value=fixed,
                    start_values=warm,
                    weighting=weighting,
                    max_nfev=int(max_nfev),
                )
            except Exception:
                sse, values, nfev, success = float("inf"), warm.copy(), 0, False
            point_map[fixed] = make_point(pid, fixed, sse, values, nfev, success)
            if success:
                warm = values.copy()

        warm = base_values.copy()
        for fixed in upper_grid:
            try:
                sse, values, nfev, success = _optimise_with_fixed_parameter(
                    info,
                    base_values,
                    fixed_index=j,
                    fixed_value=fixed,
                    start_values=warm,
                    weighting=weighting,
                    max_nfev=int(max_nfev),
                )
            except Exception:
                sse, values, nfev, success = float("inf"), warm.copy(), 0, False
            point_map[fixed] = make_point(pid, fixed, sse, values, nfev, success)
            if success:
                warm = values.copy()

        zero_point = None
        if pid in zero_selected:
            if float(lower[j]) < 0:
                raise ModelDiagnosticsValidationError(
                    f"Parameter '{pid}' has a negative admissible lower bound; KineticFittingPot will not infer that zero is a reaction-removal boundary."
                )
            if np.isclose(theta, 0.0, atol=1e-15):
                zero_point = make_point(pid, 0.0, base_sse, base_values.copy(), 0, True, zero=True)
            else:
                try:
                    # Warm-start at the smallest successful ordinary profile point.
                    successful = [p for p in point_map.values() if p.success and p.fixed_value > 0]
                    start = min(successful, key=lambda p: p.fixed_value).optimized_parameter_values if successful else tuple(base_values)
                    sse, values, nfev, success = _optimise_with_fixed_parameter(
                        info,
                        base_values,
                        fixed_index=j,
                        fixed_value=0.0,
                        start_values=np.asarray(start, dtype=float),
                        weighting=weighting,
                        max_nfev=int(max_nfev),
                    )
                except Exception:
                    sse, values, nfev, success = float("inf"), base_values.copy(), 0, False
                zero_point = make_point(pid, 0.0, sse, values, nfev, success, zero=True)

        ordered = tuple(point_map[k] for k in sorted(point_map))
        profiles.append(
            ParameterProfileResult(
                parameter_id=pid,
                fitted_value=theta,
                lower_bound=float(lower[j]),
                upper_bound=float(upper[j]),
                grid_mode=mode,
                points=ordered,
                zero_boundary_point=zero_point,
            )
        )

    return ProfileDiagnosticsResult(
        weighting=weighting,
        base_objective_sse=base_sse,
        base_weighted_rmse=base_rmse,
        parameter_ids=all_ids,
        profiles=tuple(profiles),
        improved_solution_found=bool(improved),
    )
