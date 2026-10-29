from __future__ import annotations

from datetime import datetime, timezone
from io import BytesIO
import hashlib

import numpy as np
import pandas as pd

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    from openpyxl.utils.dataframe import dataframe_to_rows
except ImportError:  # pragma: no cover - guarded at runtime
    load_workbook = None
    Alignment = Font = PatternFill = None
    get_column_letter = dataframe_to_rows = None

from utils.model_diagnostics_engine import (
    format_invariant,
    residual_autocorrelation,
    time_grid_regularity,
)


DIAG_PREFIX = "Diag_"


def _clean_value(value):
    if value is None:
        return None
    if isinstance(value, (np.bool_, bool)):
        return bool(value)
    if isinstance(value, (np.integer,)):
        return int(value)
    if isinstance(value, (np.floating, float)):
        x = float(value)
        return x if np.isfinite(x) else None
    if isinstance(value, (tuple, list, set)):
        return ", ".join(str(v) for v in value)
    return value


def _clean_dataframe(df: pd.DataFrame | None) -> pd.DataFrame:
    if df is None:
        return pd.DataFrame()
    out = df.copy()
    for col in out.columns:
        out[col] = out[col].map(_clean_value)
    return out


def _style_sheet(ws, *, dark: bool = False) -> None:
    fill = PatternFill("solid", fgColor="111827" if dark else "0F766E")
    font = Font(bold=True, color="FFFFFF")
    if ws.max_row >= 1:
        for cell in ws[1]:
            cell.fill = fill
            cell.font = font
            cell.alignment = Alignment(horizontal="center", vertical="top", wrap_text=True)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
    for row in ws.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
    for idx, col in enumerate(ws.iter_cols(1, ws.max_column), start=1):
        width = max(12, min(42, max((len(str(c.value)) if c.value is not None else 0) for c in col) + 2))
        ws.column_dimensions[get_column_letter(idx)].width = width


def _add_dataframe(wb, title: str, df: pd.DataFrame | None, *, dark: bool = False) -> None:
    ws = wb.create_sheet(title)
    frame = _clean_dataframe(df)
    if frame.empty and len(frame.columns) == 0:
        ws.append(["Status"])
        ws.append(["No results available for this diagnostic in the saved run."])
    else:
        for row in dataframe_to_rows(frame, index=False, header=True):
            ws.append([_clean_value(v) for v in row])
        if ws.max_row == 1:
            ws.append([None for _ in range(max(1, ws.max_column))])
    _style_sheet(ws, dark=dark)


def _residual_extra_tables(residual_artifact: dict | None) -> tuple[pd.DataFrame, pd.DataFrame]:
    if not residual_artifact or residual_artifact.get("points") is None:
        return pd.DataFrame(), pd.DataFrame()
    df = residual_artifact["points"].copy()
    if df.empty or not {"Dataset", "Species", "Time_s", "Residual"}.issubset(df.columns):
        return pd.DataFrame(), pd.DataFrame()
    acf_rows: list[dict] = []
    corr_rows: list[dict] = []

    for (dataset, species), group in df.groupby(["Dataset", "Species"], sort=False):
        group = group.sort_values("Time_s")
        values = group["Residual"].dropna().to_numpy(dtype=float)
        regularity = time_grid_regularity(group["Time_s"].to_numpy(dtype=float), tolerance_fraction=0.05)
        if bool(regularity["regular"]):
            acf = residual_autocorrelation(values, max_lag=min(10, max(1, len(values) // 2)))
            for lag, rho in acf:
                acf_rows.append({
                    "Dataset": dataset,
                    "Species": species,
                    "Regular_grid": True,
                    "Median_interval_s": regularity.get("median_interval"),
                    "Lag": lag,
                    "Autocorrelation": rho,
                    "Note": regularity.get("reason"),
                })
        else:
            acf_rows.append({
                "Dataset": dataset,
                "Species": species,
                "Regular_grid": False,
                "Median_interval_s": regularity.get("median_interval"),
                "Lag": None,
                "Autocorrelation": None,
                "Note": regularity.get("reason"),
            })

    for dataset, ddf in df.groupby("Dataset", sort=False):
        pivot = ddf.pivot_table(index="Time_s", columns="Species", values="Residual", aggfunc="mean")
        species = list(pivot.columns)
        for i, a in enumerate(species):
            for b in species[i:]:
                if a == b:
                    series = pivot[a].dropna()
                    n_common = int(len(series))
                    corr = 1.0 if n_common >= 1 else np.nan
                else:
                    paired = pivot[[a, b]].dropna()
                    n_common = int(len(paired))
                    corr = paired[a].corr(paired[b]) if n_common >= 3 else np.nan
                corr_rows.append({
                    "Dataset": dataset,
                    "Species_1": a,
                    "Species_2": b,
                    "Common_time_points": n_common,
                    "Residual_correlation": corr,
                })
    return pd.DataFrame(acf_rows), pd.DataFrame(corr_rows)


def build_model_diagnostics_workbook(
    original_bytes: bytes,
    info,
    diag: dict,
    *,
    integer_artifact: dict | None = None,
    balance_artifact: dict | None = None,
    reconciliation_artifact: dict | None = None,
    residual_artifact: dict | None = None,
    local_artifact: dict | None = None,
    profile_artifact: dict | None = None,
    reactionpot_version: str = "KineticFittingPot v1.40",
) -> bytes:
    """Append a reproducible diagnostics record to a completed fitting workbook.

    The original model/data and Fit_* sheets are preserved. Existing Diag_*
    sheets are replaced so re-exporting the same completed fit is deterministic
    apart from the recorded export timestamp.
    """
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to generate diagnostics output.")
    wb = load_workbook(BytesIO(original_bytes), data_only=False, read_only=False)
    for name in list(wb.sheetnames):
        if name.startswith(DIAG_PREFIX):
            del wb[name]

    species_ids = tuple(info.problem.model.species_ids)
    reaction_ids = tuple(info.problem.model.reaction_ids)

    summary_rows = [
        ["KineticFittingPot model-diagnostics export", reactionpot_version, "Completed fitting workbook plus post-fit diagnostic record"],
        ["Exported UTC", datetime.now(timezone.utc).isoformat(timespec="seconds"), "Informational timestamp"],
        ["Source workbook SHA-256", hashlib.sha256(original_bytes).hexdigest(), "Hash of uploaded completed fitting workbook"],
        ["Model", info.summary.get("Model", info.problem.model.model_name), ""],
        ["Datasets", len(info.problem.datasets), ""],
        ["Measured species", len(info.measured_species), ", ".join(info.measured_species)],
        ["Species in model", diag.get("species_count"), ""],
        ["Rate/reaction columns", diag.get("reaction_columns"), ""],
        ["Stoichiometric rank", diag.get("rank"), ""],
        ["Independent full-model invariants", diag.get("invariant_dimension"), ""],
        ["Independent measured-only invariants", diag.get("measured_invariant_dimension"), ""],
        ["Fitted parameters", len(info.fitted_parameters), ", ".join(info.fitted_parameters)],
        ["Fit weighting", info.summary.get("Weighting"), "Inherited from completed fit"],
        ["Best weighted SSE", info.summary.get("Best weighted SSE"), "Inherited from completed fit"],
        ["Best raw RMSE", info.summary.get("Best raw RMSE"), "kmol m^-3; inherited from completed fit"],
        ["Profile diagnostics included", bool(profile_artifact and profile_artifact.get("result") is not None), "Run parameter profiles before export to include them"],
    ]
    _add_dataframe(wb, "Diag_Summary", pd.DataFrame(summary_rows, columns=["Item", "Value", "Note"]), dark=True)

    settings_rows = []
    if integer_artifact:
        settings_rows.append(["Simple integer invariant maximum coefficient", integer_artifact.get("max_coefficient"), "Stoichiometric catalogue setting"])
    if balance_artifact:
        settings_rows.extend([
            ["Measured invariant monitored", balance_artifact.get("label"), "Stoichiometric-consistency monitor"],
            ["Assumed relative measurement uncertainty (fraction, 1 sigma)", balance_artifact.get("relative_fraction"), "Used when workbook sigma is absent"],
            ["Absolute measurement uncertainty floor (kmol m^-3, 1 sigma)", balance_artifact.get("absolute_floor"), "Combined in quadrature with relative term"],
        ])
    elif residual_artifact:
        settings_rows.extend([
            ["Residual diagnostic relative uncertainty (fraction, 1 sigma)", residual_artifact.get("relative_fraction"), "Defaulted because no measured invariant supplied controls"],
            ["Residual diagnostic absolute uncertainty floor (kmol m^-3, 1 sigma)", residual_artifact.get("absolute_floor"), "Defaulted because no measured invariant supplied controls"],
        ])
    if profile_artifact:
        settings_rows.extend([
            ["Profile parameters", profile_artifact.get("selected_parameters"), "Stored profile run"],
            ["Zero-boundary test parameters", profile_artifact.get("zero_parameters"), "Stored profile run"],
            ["Profile grid points per parameter", profile_artifact.get("grid_points"), "Stored profile run"],
            ["Profile span factor", profile_artifact.get("span_factor"), "Stored profile run"],
            ["Profile max function evaluations per point", profile_artifact.get("max_nfev"), "Stored profile run"],
            ["Displayed controls match stored profile run", profile_artifact.get("controls_match"), "False means page controls changed after the saved result was calculated"],
        ])
    _add_dataframe(wb, "Diag_Settings", pd.DataFrame(settings_rows, columns=["Setting", "Value", "Note"]))

    s = np.asarray(info.problem.model.stoichiometric_matrix, dtype=float)
    stoich_df = pd.DataFrame(s, columns=reaction_ids)
    stoich_df.insert(0, "Species", species_ids)
    _add_dataframe(wb, "Diag_Stoich_Matrix", stoich_df)

    invariant_rows = []
    for i, item in enumerate(diag.get("classified_full_basis", ()), start=1):
        invariant_rows.append({
            "Invariant": f"B{i}",
            "Expression": format_invariant(item.coefficients, species_ids),
            "Fully_observed": item.fully_observed,
            "Missing_species": ", ".join(item.missing_species),
            **{sid: coeff for sid, coeff in zip(species_ids, item.coefficients)},
        })
    _add_dataframe(wb, "Diag_Invariants", pd.DataFrame(invariant_rows))

    measured_rows = []
    for i, coeffs in enumerate(diag.get("measured_basis", ()), start=1):
        measured_rows.append({
            "Invariant": f"M{i}",
            "Expression": format_invariant(coeffs, species_ids),
            **{sid: coeff for sid, coeff in zip(species_ids, coeffs)},
        })
    _add_dataframe(wb, "Diag_Measured_Invariants", pd.DataFrame(measured_rows))

    integer_rows = []
    if integer_artifact:
        for i, item in enumerate(integer_artifact.get("classified", ()), start=1):
            integer_rows.append({
                "Catalogue_item": i,
                "Expression": format_invariant(item.coefficients, species_ids),
                "Fully_observed": item.fully_observed,
                "Missing_species": ", ".join(item.missing_species),
                "Coefficient_sum": int(sum(item.coefficients)),
                **{sid: coeff for sid, coeff in zip(species_ids, item.coefficients)},
            })
    _add_dataframe(wb, "Diag_Integer_Invariants", pd.DataFrame(integer_rows))

    if balance_artifact:
        coeffs = tuple(balance_artifact.get("coefficients", ()))
        balance_meta = pd.DataFrame([{
            "Selected_invariant": balance_artifact.get("label"),
            "Expression": format_invariant(coeffs, species_ids) if coeffs else None,
            **{sid: coeff for sid, coeff in zip(species_ids, coeffs)},
        }])
        _add_dataframe(wb, "Diag_Balance_Definition", balance_meta)
        _add_dataframe(wb, "Diag_Balance_Points", balance_artifact.get("balance_dataframe"))
        _add_dataframe(wb, "Diag_Balance_Summary", balance_artifact.get("summary_dataframe"))

    if reconciliation_artifact:
        _add_dataframe(wb, "Diag_Recon_Observations", reconciliation_artifact.get("observations"))
        _add_dataframe(wb, "Diag_Recon_Species", reconciliation_artifact.get("species_summary"))
        _add_dataframe(wb, "Diag_Recon_Datasets", reconciliation_artifact.get("dataset_summary"))

    if residual_artifact:
        _add_dataframe(wb, "Diag_Residual_Points", residual_artifact.get("points"))
        _add_dataframe(wb, "Diag_Residual_Species", residual_artifact.get("species_summary"))
        _add_dataframe(wb, "Diag_Residual_Datasets", residual_artifact.get("dataset_summary"))
        _add_dataframe(wb, "Diag_Residual_By_DS_Sp", residual_artifact.get("dataset_species_summary"))
        acf_df, corr_df = _residual_extra_tables(residual_artifact)
        _add_dataframe(wb, "Diag_Residual_ACF", acf_df)
        _add_dataframe(wb, "Diag_Residual_Corr", corr_df)

    if local_artifact and local_artifact.get("result") is not None:
        result = local_artifact["result"]
        local_summary = pd.DataFrame([
            ["Weighting", result.weighting],
            ["Observations", len(result.weighted_residuals)],
            ["Fitted parameters", len(result.parameter_ids)],
            ["Numerical rank", result.numerical_rank],
            ["Degrees of freedom", result.degrees_of_freedom],
            ["Condition number", result.condition_number],
            ["Objective SSE", result.objective_sse],
            ["Residual variance estimate", result.residual_variance_estimate],
        ], columns=["Metric", "Value"])
        _add_dataframe(wb, "Diag_Local_Summary", local_summary)
        _add_dataframe(wb, "Diag_Local_Parameters", local_artifact.get("parameter_table"))
        _add_dataframe(wb, "Diag_Local_Pairs", local_artifact.get("pair_table"))
        _add_dataframe(wb, "Diag_Singular_Values", pd.DataFrame({
            "Direction": np.arange(1, len(result.singular_values) + 1),
            "Singular_value": list(result.singular_values),
            "Normalised_to_largest": [v / result.singular_values[0] if result.singular_values and result.singular_values[0] else np.nan for v in result.singular_values],
        }))
        corr_df = pd.DataFrame(np.asarray(result.correlation, dtype=float), columns=result.parameter_ids)
        corr_df.insert(0, "Parameter", result.parameter_ids)
        _add_dataframe(wb, "Diag_Parameter_Corr", corr_df)
        cov_df = pd.DataFrame(np.asarray(result.covariance, dtype=float), columns=result.parameter_ids)
        cov_df.insert(0, "Parameter", result.parameter_ids)
        _add_dataframe(wb, "Diag_Parameter_Cov", cov_df)
        weak_df = pd.DataFrame({"Parameter": result.parameter_ids, "Weakest_direction_coefficient": result.weakest_direction})
        _add_dataframe(wb, "Diag_Weakest_Direction", weak_df)
        jac_df = pd.DataFrame(np.asarray(result.jacobian, dtype=float), columns=result.parameter_ids)
        jac_df.insert(0, "Residual_row", np.arange(1, len(jac_df) + 1))
        _add_dataframe(wb, "Diag_Jacobian", jac_df)
        sjac_df = pd.DataFrame(np.asarray(result.scaled_jacobian, dtype=float), columns=result.parameter_ids)
        sjac_df.insert(0, "Residual_row", np.arange(1, len(sjac_df) + 1))
        _add_dataframe(wb, "Diag_Scaled_Jacobian", sjac_df)
        wr_df = pd.DataFrame({"Residual_row": np.arange(1, len(result.weighted_residuals) + 1), "Weighted_residual": result.weighted_residuals})
        _add_dataframe(wb, "Diag_Weighted_Residuals", wr_df)

    if profile_artifact and profile_artifact.get("result") is not None:
        result = profile_artifact["result"]
        profile_summary = pd.DataFrame([
            ["Weighting", result.weighting],
            ["Baseline objective SSE", result.base_objective_sse],
            ["Baseline weighted RMSE", result.base_weighted_rmse],
            ["Parameters profiled", len(result.profiles)],
            ["Improved solution found", result.improved_solution_found],
        ], columns=["Metric", "Value"])
        _add_dataframe(wb, "Diag_Profile_Summary", profile_summary)

        point_rows = []
        zero_rows = []
        all_pids = tuple(result.parameter_ids)
        for profile in result.profiles:
            for point in profile.points:
                row = {
                    "Profile_parameter": profile.parameter_id,
                    "Fitted_value": profile.fitted_value,
                    "Lower_bound": profile.lower_bound,
                    "Upper_bound": profile.upper_bound,
                    "Grid_mode": profile.grid_mode,
                    "Fixed_value": point.fixed_value,
                    "Fixed_over_fitted": point.fixed_value / profile.fitted_value if profile.fitted_value != 0 else np.nan,
                    "Objective_SSE": point.objective_sse,
                    "Delta_SSE": point.delta_sse,
                    "Relative_increase_percent": point.relative_increase_percent,
                    "Weighted_RMSE": point.weighted_rmse,
                    "Optimiser_success": point.success,
                    "nfev": point.nfev,
                }
                row.update({f"Optimised_{pid}": value for pid, value in zip(all_pids, point.optimized_parameter_values)})
                point_rows.append(row)
            zp = profile.zero_boundary_point
            if zp is not None:
                row = {
                    "Profile_parameter": profile.parameter_id,
                    "Fitted_value": profile.fitted_value,
                    "Fixed_value": 0.0,
                    "Objective_SSE": zp.objective_sse,
                    "Delta_SSE": zp.delta_sse,
                    "Relative_increase_percent": zp.relative_increase_percent,
                    "Weighted_RMSE": zp.weighted_rmse,
                    "Optimiser_success": zp.success,
                    "nfev": zp.nfev,
                }
                row.update({f"Optimised_{pid}": value for pid, value in zip(all_pids, zp.optimized_parameter_values)})
                zero_rows.append(row)
        _add_dataframe(wb, "Diag_Profile_Points", pd.DataFrame(point_rows))
        _add_dataframe(wb, "Diag_Zero_Tests", pd.DataFrame(zero_rows))
    else:
        _add_dataframe(wb, "Diag_Profile_Summary", pd.DataFrame([
            {"Status": "Profile diagnostics were not included in this export. Run parameter profiles first, then download the diagnostics workbook."}
        ]))

    out = BytesIO()
    wb.save(out)
    return out.getvalue()
