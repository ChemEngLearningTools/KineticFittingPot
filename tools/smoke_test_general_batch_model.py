from __future__ import annotations

from pathlib import Path
from io import BytesIO
import sys

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from utils.batch_model import initial_moles_from_component_volumes, reference_parameters, simulate_van_de_vusse
from utils.general_batch_model import (
    ModelValidationError,
    build_results_workbook,
    parse_model_workbook,
    simulate_general_isothermal_batch,
    species_peak,
)


def workbook_with_volume_mode(raw: bytes, mode: str) -> bytes:
    wb = load_workbook(BytesIO(raw), data_only=False)
    ws = wb["Model"]
    for row in range(2, ws.max_row + 1):
        if ws.cell(row=row, column=1).value == "Volume_Model":
            ws.cell(row=row, column=2).value = mode
            break
    else:
        raise AssertionError("Template Model sheet does not contain Volume_Model.")
    out = BytesIO()
    wb.save(out)
    return out.getvalue()


def check_against_builtin(raw: bytes, mode: str) -> tuple[dict, dict, dict]:
    model = parse_model_workbook(workbook_with_volume_mode(raw, mode))
    result = simulate_general_isothermal_batch(model)
    peak = species_peak(model, result, "x3")

    ref = reference_parameters()
    n0 = initial_moles_from_component_volumes(
        np.array([0.5, 0.5, 0.0, 0.0, 0.0]),
        ref.molecular_weight_kg_per_kmol,
        ref.density_kg_per_m3,
    )
    builtin_mode = "composition-dependent" if mode == "composition_dependent" else "constant"
    built_in = simulate_van_de_vusse(
        initial_moles_kmol=n0,
        temperature_c=60.0,
        batch_time_s=7200.0,
        output_points=1200,
        volume_model=builtin_mode,
        pre_exponential=ref.pre_exponential,
        activation_energy_kj_per_kmol=ref.activation_energy_kj_per_kmol,
        molecular_weight_kg_per_kmol=ref.molecular_weight_kg_per_kmol,
        density_kg_per_m3=ref.density_kg_per_m3,
    )

    dt = abs(peak["time_s"] - built_in["optimum"]["time_s"])
    dn = abs(peak["amount_kmol"] - float(built_in["optimum"]["n_kmol"][3]))
    dv = abs(result["final_volume_m3"] - float(built_in["volume_m3"][-1]))

    # The dedicated simulator locates its displayed optimum on a dense regular
    # grid, whereas the general engine refines the local maximum continuously.
    # A sub-second time difference is therefore expected in the variable-volume case.
    if dt > 1.0:
        raise AssertionError(f"{mode}: Excel-defined x3 peak time differs from built-in reference by {dt:.6g} s")
    if dn > 2e-7:
        raise AssertionError(f"{mode}: Excel-defined x3 peak amount differs from built-in reference by {dn:.6g} kmol")
    if dv > 1e-10:
        raise AssertionError(f"{mode}: final liquid volume differs from built-in reference by {dv:.6g} m3")
    if result["min_amount_kmol"] < -1e-8:
        raise AssertionError(f"{mode}: unexpected negative amount: {result['min_amount_kmol']:.6g} kmol")

    return model, result, peak


def main() -> None:
    template = ROOT / "data" / "templates" / "KineticFittingPot_Isothermal_Batch_Model_Template.xlsx"
    raw = template.read_bytes()

    comp_model, comp_result, comp_peak = check_against_builtin(raw, "composition_dependent")
    const_model, const_result, const_peak = check_against_builtin(raw, "constant")

    completed = build_results_workbook(raw, comp_model, comp_result, "x3")
    if len(completed) <= len(raw):
        raise AssertionError("Completed workbook did not grow after adding simulation outputs.")

    completed_wb = load_workbook(BytesIO(completed), data_only=False, read_only=True)
    results_headers = [cell.value for cell in next(completed_wb["Results_TimeSeries"].iter_rows(min_row=1, max_row=1))]
    if "volume_m3" not in results_headers:
        raise AssertionError("Completed workbook does not include the volume_m3 trajectory.")

    # Security regression: arbitrary Python in a rate cell must be rejected.
    malicious_wb = load_workbook(BytesIO(raw), data_only=False)
    malicious_wb["Reactions"]["C2"] = "__import__('os').system('echo unsafe')"
    malicious_buffer = BytesIO()
    malicious_wb.save(malicious_buffer)
    try:
        parse_model_workbook(malicious_buffer.getvalue())
    except ModelValidationError:
        pass
    else:
        raise AssertionError("Unsafe arbitrary-Python rate expression was not rejected.")

    # Variable-volume consistency regression: an inconsistent initial charge
    # must be rejected rather than silently rescaled.
    bad_wb = load_workbook(BytesIO(raw), data_only=False)
    bad_wb["Species"]["E2"] = 500.0
    bad_buffer = BytesIO()
    bad_wb.save(bad_buffer)
    try:
        parse_model_workbook(bad_buffer.getvalue())
    except ModelValidationError:
        pass
    else:
        raise AssertionError("Inconsistent variable-volume initial charge was not rejected.")

    print("General Excel-defined isothermal batch smoke test: PASS")
    print(
        "Composition-dependent: "
        f"x3 peak={comp_peak['time_s']:.6f} s, {comp_peak['amount_kmol']:.9f} kmol; "
        f"Vfinal={comp_result['final_volume_m3']:.9f} m3"
    )
    print(
        "Constant: "
        f"x3 peak={const_peak['time_s']:.6f} s, {const_peak['amount_kmol']:.9f} kmol; "
        f"Vfinal={const_result['final_volume_m3']:.9f} m3"
    )


if __name__ == "__main__":
    main()
