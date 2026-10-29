from __future__ import annotations

from dataclasses import dataclass, replace
from io import BytesIO
from typing import Iterable
from time import monotonic

import numpy as np

try:
    from scipy.integrate import solve_ivp
    from scipy.optimize import least_squares
except ImportError:
    solve_ivp = None
    least_squares = None

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:
    load_workbook = None

from utils.general_batch_model import (
    GeneralBatchModel,
    ModelValidationError,
    evaluate_rates,
    model_volume_m3,
    parse_model_workbook,
)


@dataclass(frozen=True)
class FitDataset:
    dataset_id: str
    description: str
    volume_m3: float
    volume_model: str
    initial_concentration_kmol_m3: np.ndarray


@dataclass(frozen=True)
class Observation:
    dataset_id: str
    time_s: float
    species_id: str
    measured_value: float
    sigma: float | None


@dataclass(frozen=True)
class KineticFittingProblem:
    model: GeneralBatchModel
    datasets: tuple[FitDataset, ...]
    observations: tuple[Observation, ...]
    fit_parameter_indices: tuple[int, ...]
    lower_bounds: np.ndarray
    upper_bounds: np.ndarray
    starting_values: np.ndarray
    validation_datasets: tuple[FitDataset, ...] = ()
    validation_observations: tuple[Observation, ...] = ()

    @property
    def fit_parameter_ids(self) -> tuple[str, ...]:
        return tuple(self.model.parameter_ids[i] for i in self.fit_parameter_indices)


class FittingValidationError(ValueError):
    pass


class FittingBudgetExceeded(RuntimeError):
    """A public fitting job exhausted its cooperative computation budget."""


class FittingBudget:
    def __init__(self, *, max_integrations: int, max_rhs_evaluations: int, max_wall_time_s: float):
        self.max_integrations = int(max_integrations)
        self.max_rhs_evaluations = int(max_rhs_evaluations)
        self.max_wall_time_s = float(max_wall_time_s)
        self.started = monotonic()
        self.integrations = 0
        self.rhs_evaluations = 0
        self.failed_forward_evaluations = 0
        self.objective_calls = 0
        self.reason = None

    def check(self):
        if self.reason:
            raise FittingBudgetExceeded(self.reason)
        elapsed = monotonic() - self.started
        if elapsed >= self.max_wall_time_s:
            self.reason = f"Elapsed computation time exceeded {self.max_wall_time_s:g} s."
        elif self.integrations > self.max_integrations:
            self.reason = f"ODE integrations exceeded {self.max_integrations:,}."
        elif self.rhs_evaluations > self.max_rhs_evaluations:
            self.reason = f"ODE RHS evaluations exceeded {self.max_rhs_evaluations:,}."
        if self.reason:
            raise FittingBudgetExceeded(self.reason)

    def integration(self):
        self.integrations += 1
        self.check()

    def rhs(self):
        self.rhs_evaluations += 1
        self.check()

    def summary(self):
        return {"integrations": self.integrations, "rhs_evaluations": self.rhs_evaluations,
                "objective_calls": self.objective_calls,
                "failed_forward_evaluations": self.failed_forward_evaluations,
                "elapsed_s": monotonic() - self.started, "budget_reason": self.reason}


def _header_map(ws) -> dict[str, int]:
    return {
        str(cell.value).strip(): idx
        for idx, cell in enumerate(ws[1], start=1)
        if cell.value not in (None, "")
    }


def _iter_nonblank_rows(ws, start_row: int = 2) -> Iterable[int]:
    for row in range(start_row, ws.max_row + 1):
        vals = [ws.cell(row=row, column=c).value for c in range(1, ws.max_column + 1)]
        if any(v not in (None, "") for v in vals):
            yield row


def _float(value, label: str) -> float:
    try:
        x = float(value)
    except (TypeError, ValueError) as exc:
        raise FittingValidationError(f"{label} must be numeric.") from exc
    if not np.isfinite(x):
        raise FittingValidationError(f"{label} must be finite.")
    return x


def _yes(value) -> bool:
    text = str(value or "").strip().lower()
    return text in {"yes", "y", "true", "1"}


def parse_fitting_workbook(workbook_bytes: bytes) -> KineticFittingProblem:
    if load_workbook is None:
        raise RuntimeError("Kinetic fitting requires openpyxl.")
    model = parse_model_workbook(workbook_bytes)
    try:
        wb = load_workbook(BytesIO(workbook_bytes), data_only=False, read_only=False)
    except Exception as exc:
        raise FittingValidationError(f"KineticFittingPot could not read this fitting workbook: {exc}") from exc

    for name in ("Datasets", "Experimental_Data"):
        if name not in wb.sheetnames:
            raise FittingValidationError(f"Missing required worksheet '{name}'.")

    fit_idx = tuple(i for i, flag in enumerate(model.fit_later) if flag)
    if not fit_idx:
        raise FittingValidationError("No Parameters rows are marked Fit_Later=Yes.")
    if len(fit_idx) > 12:
        raise FittingValidationError("The first KineticFittingPot fitting engine supports at most 12 fitted parameters.")

    lower, upper, starts = [], [], []
    for i in fit_idx:
        pid = model.parameter_ids[i]
        lb = model.lower_bounds[i]
        ub = model.upper_bounds[i]
        start = float(model.parameter_values[i])
        if lb is None or ub is None or not np.isfinite(lb) or not np.isfinite(ub):
            raise FittingValidationError(f"Fitted parameter {pid} requires finite Lower_Bound and Upper_Bound values.")
        if not float(lb) < float(ub):
            raise FittingValidationError(f"Fitted parameter {pid}: lower bound must be below upper bound.")
        if not float(lb) <= start <= float(ub):
            raise FittingValidationError(f"Fitted parameter {pid}: starting Value must lie inside its bounds.")
        lower.append(float(lb)); upper.append(float(ub)); starts.append(start)

    # Datasets -------------------------------------------------------------
    ws = wb["Datasets"]
    headers = _header_map(ws)
    required = ["Dataset_ID", "Volume_m3", *[f"c0_{sid}" for sid in model.species_ids]]
    missing = [h for h in required if h not in headers]
    if missing:
        raise FittingValidationError("Datasets sheet is missing column(s): " + ", ".join(missing))

    datasets: list[FitDataset] = []
    validation_datasets: list[FitDataset] = []
    dataset_ids: set[str] = set()
    for row in _iter_nonblank_rows(ws):
        use_for_fit = "Use_for_Fit" not in headers or _yes(ws.cell(row=row, column=headers["Use_for_Fit"]).value)
        ds = str(ws.cell(row=row, column=headers["Dataset_ID"]).value or "").strip()
        if not ds:
            raise FittingValidationError(f"Datasets row {row}: Dataset_ID is required.")
        if ds in dataset_ids:
            raise FittingValidationError(f"Duplicate Dataset_ID '{ds}'.")
        dataset_ids.add(ds)
        description = str(ws.cell(row=row, column=headers.get("Description", 0)).value or ds).strip() if "Description" in headers else ds
        volume = _float(ws.cell(row=row, column=headers["Volume_m3"]).value, f"Dataset {ds} Volume_m3")
        if volume <= 0:
            raise FittingValidationError(f"Dataset {ds}: Volume_m3 must be positive.")
        volume_model = model.volume_model
        if "Volume_Model" in headers:
            raw = str(ws.cell(row=row, column=headers["Volume_Model"]).value or volume_model).strip().lower().replace("-", "_")
            if raw not in {"constant", "composition_dependent"}:
                raise FittingValidationError(f"Dataset {ds}: Volume_Model must be constant or composition_dependent.")
            volume_model = raw
        c0 = np.asarray([
            _float(ws.cell(row=row, column=headers[f"c0_{sid}"]).value, f"Dataset {ds} c0_{sid}")
            for sid in model.species_ids
        ], dtype=float)
        if np.any(c0 < 0):
            raise FittingValidationError(f"Dataset {ds}: initial concentrations cannot be negative.")
        if volume_model == "composition_dependent":
            n0 = c0 * volume
            implied = float(np.sum(n0 * model.molecular_weight_kg_per_kmol / model.density_kg_per_m3))
            mismatch = abs(implied - volume) / volume
            if mismatch > 1e-3:
                raise FittingValidationError(
                    f"Dataset {ds}: composition-dependent initial conditions imply {implied:.6g} m3, "
                    f"not the specified {volume:.6g} m3 ({100*mismatch:.3g}% mismatch)."
                )
        (datasets if use_for_fit else validation_datasets).append(FitDataset(ds, description, volume, volume_model, c0))

    if not datasets:
        raise FittingValidationError("No Datasets rows are selected for fitting.")
    if len(datasets) > 30:
        raise FittingValidationError("The first fitting engine supports at most 30 selected datasets.")

    # Observations ---------------------------------------------------------
    ws = wb["Experimental_Data"]
    headers = _header_map(ws)
    required = ["Dataset_ID", "Time_s", "Species_ID", "Measured_Value_kmol_m3"]
    missing = [h for h in required if h not in headers]
    if missing:
        raise FittingValidationError("Experimental_Data sheet is missing column(s): " + ", ".join(missing))

    observations: list[Observation] = []
    validation_observations: list[Observation] = []
    validation_ids = {d.dataset_id for d in validation_datasets}
    active_ids = {d.dataset_id for d in datasets}
    for row in _iter_nonblank_rows(ws):
        ds = str(ws.cell(row=row, column=headers["Dataset_ID"]).value or "").strip()
        if ds not in active_ids and ds not in validation_ids:
            continue
        time_s = _float(ws.cell(row=row, column=headers["Time_s"]).value, f"Experimental_Data row {row} Time_s")
        if time_s < 0:
            raise FittingValidationError(f"Experimental_Data row {row}: Time_s cannot be negative.")
        sid = str(ws.cell(row=row, column=headers["Species_ID"]).value or "").strip()
        if sid not in model.species_ids:
            raise FittingValidationError(f"Experimental_Data row {row}: unknown Species_ID '{sid}'.")
        raw_y = ws.cell(row=row, column=headers["Measured_Value_kmol_m3"]).value
        if raw_y in (None, ""):
            raise FittingValidationError(
                f"Experimental_Data row {row} has no measured value. Represent missing data by removing the row."
            )
        y = _float(raw_y, f"Experimental_Data row {row} measured value")
        sigma = None
        if "Sigma_kmol_m3" in headers:
            raw_sigma = ws.cell(row=row, column=headers["Sigma_kmol_m3"]).value
            if raw_sigma not in (None, ""):
                sigma = _float(raw_sigma, f"Experimental_Data row {row} Sigma_kmol_m3")
                if sigma <= 0:
                    raise FittingValidationError(f"Experimental_Data row {row}: Sigma_kmol_m3 must be positive.")
        (observations if ds in active_ids else validation_observations).append(Observation(ds, time_s, sid, y, sigma))

    if not observations:
        raise FittingValidationError("No observations remain for the selected datasets.")
    if len(observations) <= len(fit_idx):
        raise FittingValidationError("The fit needs more observations than fitted parameters.")

    return KineticFittingProblem(
        model=model,
        datasets=tuple(datasets),
        observations=tuple(observations),
        fit_parameter_indices=fit_idx,
        lower_bounds=np.asarray(lower, dtype=float),
        upper_bounds=np.asarray(upper, dtype=float),
        starting_values=np.asarray(starts, dtype=float),
        validation_datasets=tuple(validation_datasets),
        validation_observations=tuple(validation_observations),
    )


def _parameter_to_unit_interval(values: np.ndarray, lower: np.ndarray, upper: np.ndarray) -> np.ndarray:
    values = np.asarray(values, dtype=float)
    out = np.empty_like(values)
    positive = (lower > 0) & (upper > 0)
    out[positive] = (np.log(values[positive]) - np.log(lower[positive])) / (np.log(upper[positive]) - np.log(lower[positive]))
    out[~positive] = (values[~positive] - lower[~positive]) / (upper[~positive] - lower[~positive])
    return np.clip(out, 0.0, 1.0)


def _unit_interval_to_parameter(z: np.ndarray, lower: np.ndarray, upper: np.ndarray) -> np.ndarray:
    z = np.asarray(z, dtype=float)
    out = np.empty_like(z)
    positive = (lower > 0) & (upper > 0)
    out[positive] = np.exp(np.log(lower[positive]) + z[positive] * (np.log(upper[positive]) - np.log(lower[positive])))
    out[~positive] = lower[~positive] + z[~positive] * (upper[~positive] - lower[~positive])
    return out


def _model_with_trial(problem: KineticFittingProblem, values: np.ndarray, dataset: FitDataset) -> GeneralBatchModel:
    params = problem.model.parameter_values.copy()
    for idx, value in zip(problem.fit_parameter_indices, values):
        params[idx] = float(value)
    return replace(
        problem.model,
        parameter_values=params,
        initial_volume_m3=float(dataset.volume_m3),
        volume_model=dataset.volume_model,
        initial_concentration_kmol_m3=np.asarray(dataset.initial_concentration_kmol_m3, dtype=float),
    )


def _simulate_dataset_at_times(problem: KineticFittingProblem, dataset: FitDataset, values: np.ndarray, times: np.ndarray, budget: FittingBudget | None = None) -> np.ndarray:
    if solve_ivp is None:
        raise RuntimeError("SciPy is required for kinetic fitting.")
    model = _model_with_trial(problem, values, dataset)
    times = np.asarray(sorted(set(float(t) for t in times)), dtype=float)
    if times.size == 0:
        return np.empty((0, len(model.species_ids)))
    if times[-1] == 0:
        return np.repeat(model.initial_concentration_kmol_m3[None, :], len(times), axis=0)
    n0 = model.initial_moles_kmol.astype(float)
    if budget is not None:
        budget.integration()

    def rhs(_t: float, n: np.ndarray) -> np.ndarray:
        if budget is not None:
            budget.rhs()
        n_rate = np.maximum(np.asarray(n, dtype=float), 0.0)
        vol = model_volume_m3(model, n_rate)
        conc = n_rate / vol
        rates = evaluate_rates(model, conc)
        return vol * (model.stoichiometric_matrix @ rates)

    t_eval = times[times > 0]
    sol = solve_ivp(rhs, (0.0, float(times[-1])), n0, t_eval=t_eval, method="RK45", rtol=1e-6, atol=1e-9)
    if not sol.success:
        raise RuntimeError(sol.message)
    conc_rows = []
    positive_map = {float(t): i for i, t in enumerate(sol.t)}
    for t in times:
        if t == 0:
            conc_rows.append(model.initial_concentration_kmol_m3.copy())
        else:
            i = positive_map[float(t)]
            n = np.maximum(sol.y[:, i], 0.0)
            vol = model_volume_m3(model, n)
            conc_rows.append(n / vol)
    return np.asarray(conc_rows, dtype=float)


def species_scales(problem: KineticFittingProblem) -> dict[str, float]:
    scales: dict[str, float] = {}
    for sid in problem.model.species_ids:
        vals = [abs(o.measured_value) for o in problem.observations if o.species_id == sid]
        if vals:
            scales[sid] = max(max(vals), 1e-6)
    return scales


def predict_observations(problem: KineticFittingProblem, values: np.ndarray, budget: FittingBudget | None = None) -> np.ndarray:
    pred = np.empty(len(problem.observations), dtype=float)
    obs_by_dataset: dict[str, list[tuple[int, Observation]]] = {}
    for i, obs in enumerate(problem.observations):
        obs_by_dataset.setdefault(obs.dataset_id, []).append((i, obs))
    dataset_map = {d.dataset_id: d for d in problem.datasets}

    for ds_id, entries in obs_by_dataset.items():
        dataset = dataset_map[ds_id]
        times = np.asarray(sorted({obs.time_s for _, obs in entries}), dtype=float)
        conc = _simulate_dataset_at_times(problem, dataset, values, times, budget=budget)
        time_index = {float(t): i for i, t in enumerate(times)}
        for out_i, obs in entries:
            sidx = problem.model.species_ids.index(obs.species_id)
            pred[out_i] = conc[time_index[float(obs.time_s)], sidx]
    return pred


def residual_vector(problem: KineticFittingProblem, values: np.ndarray, weighting: str, budget: FittingBudget | None = None) -> np.ndarray:
    prediction = predict_observations(problem, values, budget=budget)
    measured = np.asarray([o.measured_value for o in problem.observations], dtype=float)
    residual = prediction - measured
    if weighting == "unweighted":
        return residual
    if weighting == "species_scaled":
        scales = species_scales(problem)
        denom = np.asarray([scales[o.species_id] for o in problem.observations], dtype=float)
        return residual / denom
    raise ValueError(f"Unknown weighting mode '{weighting}'.")


def fit_multistart(
    problem: KineticFittingProblem,
    *,
    weighting: str = "species_scaled",
    n_starts: int = 20,
    random_seed: int = 20260904,
    max_nfev: int = 300,
    budget: FittingBudget | None = None,
) -> dict:
    if least_squares is None:
        raise RuntimeError("SciPy is required for kinetic fitting.")
    n_starts = int(n_starts)
    if not 1 <= n_starts <= 100:
        raise ValueError("n_starts must be between 1 and 100.")
    if weighting not in {"unweighted", "species_scaled"}:
        raise ValueError("weighting must be 'unweighted' or 'species_scaled'.")

    lower = problem.lower_bounds
    upper = problem.upper_bounds
    z_user = _parameter_to_unit_interval(problem.starting_values, lower, upper)
    rng = np.random.default_rng(int(random_seed))
    starts = [z_user]
    if n_starts > 1:
        # Stratified coverage of each unit interval dimension, independently permuted.
        # This gives more reliable bound coverage than simple random clustering.
        n_extra = n_starts - 1
        design = np.empty((n_extra, len(z_user)), dtype=float)
        for j in range(len(z_user)):
            bins = (np.arange(n_extra) + rng.random(n_extra)) / n_extra
            rng.shuffle(bins)
            design[:, j] = bins
        starts.extend(design)

    measured = np.asarray([o.measured_value for o in problem.observations], dtype=float)
    records = []

    def z_residual(z: np.ndarray) -> np.ndarray:
        if budget is not None:
            budget.objective_calls += 1
            budget.check()
        values = _unit_interval_to_parameter(z, lower, upper)
        try:
            res = residual_vector(problem, values, weighting, budget=budget)
            if not np.all(np.isfinite(res)):
                raise ValueError("non-finite residual")
            return res
        except FittingBudgetExceeded:
            raise
        except Exception:
            if budget is not None:
                budget.failed_forward_evaluations += 1
            return np.full(len(problem.observations), 1e6, dtype=float)

    for i, z0 in enumerate(starts, start=1):
        try:
            opt = least_squares(
                z_residual,
                np.asarray(z0, dtype=float),
                bounds=(np.zeros_like(z_user), np.ones_like(z_user)),
                method="trf",
                max_nfev=int(max_nfev),
                xtol=1e-8,
                ftol=1e-8,
                gtol=1e-8,
            )
            values = _unit_interval_to_parameter(opt.x, lower, upper)
            raw_pred = predict_observations(problem, values, budget=budget)
            raw_resid = raw_pred - measured
            weighted = residual_vector(problem, values, weighting, budget=budget)
            record = {
                "start": i,
                "success": bool(opt.success),
                "message": str(opt.message),
                "nfev": int(opt.nfev),
                "sse": float(np.sum(weighted ** 2)),
                "weighted_rmse": float(np.sqrt(np.mean(weighted ** 2))),
                "raw_rmse_kmol_m3": float(np.sqrt(np.mean(raw_resid ** 2))),
                "parameter_values": np.asarray(values, dtype=float),
                "start_values": _unit_interval_to_parameter(np.asarray(z0), lower, upper),
            }
        except FittingBudgetExceeded:
            raise
        except Exception as exc:
            record = {
                "start": i,
                "success": False,
                "message": str(exc),
                "nfev": 0,
                "sse": float("inf"),
                "weighted_rmse": float("inf"),
                "raw_rmse_kmol_m3": float("inf"),
                "parameter_values": np.full(len(z_user), np.nan),
                "start_values": _unit_interval_to_parameter(np.asarray(z0), lower, upper),
            }
        records.append(record)

    finite = [r for r in records if np.isfinite(r["sse"])]
    if not finite:
        raise RuntimeError("All multistart fitting attempts failed.")
    best = min(finite, key=lambda r: r["sse"])
    best_values = np.asarray(best["parameter_values"], dtype=float)
    predictions = predict_observations(problem, best_values, budget=budget)
    measured = np.asarray([o.measured_value for o in problem.observations], dtype=float)
    raw_residuals = predictions - measured

    return {
        "weighting": weighting,
        "n_starts": n_starts,
        "random_seed": int(random_seed),
        "max_nfev": int(max_nfev),
        "records": records,
        "best": best,
        "best_values": best_values,
        "predictions": predictions,
        "raw_residuals": raw_residuals,
        "species_scales": species_scales(problem),
        "status": "converged" if bool(best["success"]) else "non_converged",
        "budget": budget.summary() if budget is not None else None,
    }



def evaluate_validation(problem: KineticFittingProblem, fit_result: dict, budget: FittingBudget | None = None) -> dict:
    """Out-of-sample predictions; only called *after* fitting. No effect on optimiser."""
    if not problem.validation_observations:
        return {"predictions": np.array([], dtype=float), "residuals": np.array([], dtype=float), "rmse_kmol_m3": None, "sse_kmol_m6": None}
    holdout = replace(problem, datasets=problem.validation_datasets, observations=problem.validation_observations)
    pred = predict_observations(holdout, fit_result["best_values"], budget=budget)
    measured = np.array([o.measured_value for o in holdout.observations], dtype=float)
    residuals = pred - measured
    return {"predictions": pred, "residuals": residuals,
            "rmse_kmol_m3": float(np.sqrt(np.mean(residuals ** 2))),
            "sse_kmol_m6": float(np.sum(residuals ** 2))}


def fitted_trajectory(problem: KineticFittingProblem, dataset: FitDataset, values: np.ndarray, points: int = 500, budget: FittingBudget | None = None) -> dict:
    obs_times = [o.time_s for o in problem.observations if o.dataset_id == dataset.dataset_id]
    t_end = max(obs_times) if obs_times else problem.model.batch_duration_s
    times = np.linspace(0.0, float(t_end), int(points))
    conc = _simulate_dataset_at_times(problem, dataset, values, times, budget=budget)
    return {"time_s": times, "concentrations_kmol_m3": conc}


def synthetic_truth(workbook_bytes: bytes, fit_parameter_ids: tuple[str, ...]) -> dict[str, float]:
    """Read optional teaching truth only for post-fit comparison. Never called by fit_multistart."""
    if load_workbook is None:
        return {}
    wb = load_workbook(BytesIO(workbook_bytes), data_only=True, read_only=True)
    if "Synthetic_Data_Info" not in wb.sheetnames:
        return {}
    data = {str(r[0]): r[1] for r in wb["Synthetic_Data_Info"].iter_rows(min_row=2, values_only=True) if r and r[0]}
    truth = {}
    for pid in fit_parameter_ids:
        key = f"True_{pid}"
        if key in data and data[key] is not None:
            try:
                truth[pid] = float(data[key])
            except (TypeError, ValueError):
                pass
    return truth


def build_fitting_results_workbook(original_bytes: bytes, problem: KineticFittingProblem, fit_result: dict) -> bytes:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to generate fitting output.")
    wb = load_workbook(BytesIO(original_bytes), data_only=False, read_only=False)
    names = ["Fit_Parameters", "Fit_Starts", "Fit_Predictions", "Fit_Residuals", "Fit_Summary", "Validation_Predictions", "Validation_Residuals", "Validation_Summary"]
    for name in names:
        if name in wb.sheetnames:
            del wb[name]

    teal_fill = PatternFill("solid", fgColor="0F766E")
    dark_fill = PatternFill("solid", fgColor="111827")
    header_font = Font(bold=True, color="FFFFFF")

    def style_header(ws, dark=False):
        fill = dark_fill if dark else teal_fill
        for cell in ws[1]:
            cell.fill = fill; cell.font = header_font; cell.alignment = Alignment(horizontal="center")
        ws.freeze_panes = "A2"

    best = fit_result["best"]
    pids = problem.fit_parameter_ids
    truth = synthetic_truth(original_bytes, pids)

    ws = wb.create_sheet("Fit_Parameters")
    ws.append(["Parameter_ID", "Starting_Value", "Fitted_Value", "Lower_Bound", "Upper_Bound", "Synthetic_Truth_if_available", "Units"])
    style_header(ws)
    for j, pid in enumerate(pids):
        midx = problem.fit_parameter_indices[j]
        ws.append([
            pid,
            float(problem.starting_values[j]),
            float(fit_result["best_values"][j]),
            float(problem.lower_bounds[j]),
            float(problem.upper_bounds[j]),
            truth.get(pid),
            problem.model.parameter_units[midx],
        ])

    ws = wb.create_sheet("Fit_Starts")
    ws.append(["Start", "Success", "SSE", "Weighted_RMSE", "Raw_RMSE_kmol_m3", "Function_Evaluations", *[f"Start_{p}" for p in pids], *[f"Fit_{p}" for p in pids], "Message"])
    style_header(ws)
    for r in sorted(fit_result["records"], key=lambda x: x["sse"]):
        ws.append([
            r["start"], r["success"], r["sse"], r["weighted_rmse"], r["raw_rmse_kmol_m3"], r["nfev"],
            *[float(x) for x in r["start_values"]],
            *[float(x) if np.isfinite(x) else None for x in r["parameter_values"]],
            r["message"],
        ])

    ws = wb.create_sheet("Fit_Predictions")
    ws.append(["Dataset_ID", "Time_s", "Species_ID", "Measured_Value_kmol_m3", "Fitted_Value_kmol_m3"])
    style_header(ws)
    for obs, pred in zip(problem.observations, fit_result["predictions"]):
        ws.append([obs.dataset_id, obs.time_s, obs.species_id, obs.measured_value, float(pred)])

    ws = wb.create_sheet("Fit_Residuals")
    ws.append(["Dataset_ID", "Time_s", "Species_ID", "Residual_model_minus_measured_kmol_m3"])
    style_header(ws)
    for obs, resid in zip(problem.observations, fit_result["raw_residuals"]):
        ws.append([obs.dataset_id, obs.time_s, obs.species_id, float(resid)])

    validation = fit_result.get("validation") or {}
    if problem.validation_datasets:
        vp = list(validation.get("predictions", []))
        vr = list(validation.get("residuals", []))
        ws = wb.create_sheet("Validation_Predictions")
        ws.append(["Dataset_ID", "Time_s", "Species_ID", "Measured_Value_kmol_m3", "Predicted_Value_kmol_m3"])
        style_header(ws)
        for obs, predicted in zip(problem.validation_observations, vp):
            ws.append([obs.dataset_id, obs.time_s, obs.species_id, obs.measured_value, float(predicted)])
        ws = wb.create_sheet("Validation_Residuals")
        ws.append(["Dataset_ID", "Time_s", "Species_ID", "Residual_model_minus_measured_kmol_m3"])
        style_header(ws)
        for obs, resid in zip(problem.validation_observations, vr):
            ws.append([obs.dataset_id, obs.time_s, obs.species_id, float(resid)])
        ws = wb.create_sheet("Validation_Summary")
        ws.append(["Statistic", "Value", "Description"])
        style_header(ws)
        for row in [
            ["Validation datasets", len(problem.validation_datasets), "Not used for parameter estimation"],
            ["Validation observations", len(problem.validation_observations), "Not included in fitting objective"],
            ["Validation raw SSE", validation.get("sse_kmol_m6"), "Sum of unweighted squared prediction errors"],
            ["Validation raw RMSE (kmol/m3)", validation.get("rmse_kmol_m3"), "Validation observations only"],
        ]: ws.append(row)

    ws = wb.create_sheet("Fit_Summary")
    ws.append(["KineticFittingPot kinetic fitting summary", "Value", "Note"])
    style_header(ws, dark=True)
    rows = [
        ["Model", problem.model.model_name, ""],
        ["Datasets", len(problem.datasets), "Selected datasets"],
        ["Observations", len(problem.observations), "Only fitting observations; validation rows excluded"],
        ["Validation datasets", len(problem.validation_datasets), "Excluded from parameter fitting"],
        ["Validation observations", len(problem.validation_observations), "Held out for independent prediction"],
        ["Fitted parameters", len(pids), ", ".join(pids)],
        ["Weighting", fit_result["weighting"], "unweighted or species_scaled"],
        ["Multistart runs", fit_result["n_starts"], "Includes workbook starting estimate"],
        ["Random seed", fit_result["random_seed"], "Controls additional multistart locations"],
        ["Best start", best["start"], "Ranked by weighted SSE"],
        ["Best weighted SSE", best["sse"], "sum(weighted residual^2)"],
        ["Best weighted RMSE", best["weighted_rmse"], ""],
        ["Best raw RMSE", best["raw_rmse_kmol_m3"], "kmol m^-3"],
        ["Best optimiser success", best["success"], best["message"]],
    ]
    for row in rows:
        ws.append(row)

    for ws in wb.worksheets:
        if ws.title.startswith(("Fit_", "Validation_")):
            for col in ws.columns:
                letter = col[0].column_letter
                width = max(12, min(34, max(len(str(c.value)) if c.value is not None else 0 for c in col) + 2))
                ws.column_dimensions[letter].width = width
            for row in ws.iter_rows():
                for cell in row:
                    cell.alignment = Alignment(vertical="top", wrap_text=True)

    out = BytesIO(); wb.save(out); return out.getvalue()


def problem_with_fit_parameters(problem: KineticFittingProblem, parameter_ids: list[str] | tuple[str, ...]) -> KineticFittingProblem:
    selected = tuple(str(x) for x in parameter_ids)
    if not selected:
        raise FittingValidationError("Select at least one parameter to fit.")
    indices = []
    lower = []; upper = []; starts = []
    for pid in selected:
        if pid not in problem.model.parameter_ids:
            raise FittingValidationError(f"Unknown parameter '{pid}'.")
        idx = problem.model.parameter_ids.index(pid)
        if not problem.model.fit_later[idx]:
            raise FittingValidationError(f"Parameter '{pid}' is not marked Fit_Later=Yes in the workbook.")
        lb = problem.model.lower_bounds[idx]; ub = problem.model.upper_bounds[idx]
        if lb is None or ub is None:
            raise FittingValidationError(f"Parameter '{pid}' requires finite bounds.")
        indices.append(idx); lower.append(float(lb)); upper.append(float(ub)); starts.append(float(problem.model.parameter_values[idx]))
    return replace(
        problem,
        fit_parameter_indices=tuple(indices),
        lower_bounds=np.asarray(lower, dtype=float),
        upper_bounds=np.asarray(upper, dtype=float),
        starting_values=np.asarray(starts, dtype=float),
    )
