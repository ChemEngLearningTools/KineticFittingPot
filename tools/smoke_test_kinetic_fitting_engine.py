from __future__ import annotations

import sys
from pathlib import Path
from io import BytesIO

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from utils.kinetic_fitting_engine import (
    build_fitting_results_workbook,
    fit_multistart,
    parse_fitting_workbook,
    synthetic_truth,
)

BOOK = ROOT / "data" / "templates" / "KineticFittingPot_Kinetic_Fitting_Catalytic_Saturation_Synthetic_Data.xlsx"


def main() -> None:
    raw = BOOK.read_bytes()
    problem = parse_fitting_workbook(raw)
    assert len(problem.datasets) == 3
    assert len(problem.observations) == 64
    assert problem.fit_parameter_ids == ("k1", "KA", "k2")

    result = fit_multistart(problem, weighting="species_scaled", n_starts=16, random_seed=20260904, max_nfev=300)
    best = result["best_values"]
    truth = synthetic_truth(raw, problem.fit_parameter_ids)
    true = np.asarray([truth[p] for p in problem.fit_parameter_ids], dtype=float)
    rel = np.abs((best - true) / true)

    # Synthetic data contain measurement noise; the known truth should be
    # recovered approximately rather than exactly.
    if np.any(rel > np.asarray([0.20, 0.35, 0.20])):
        raise AssertionError(f"Fitted parameters are too far from synthetic truth: fit={best}, truth={true}, rel={rel}")

    finite = [r for r in result["records"] if np.isfinite(r["sse"])]
    if len(finite) < 12:
        raise AssertionError("Too many multistart fitting runs failed.")
    best_sse = min(r["sse"] for r in finite)
    close = [r for r in finite if r["sse"] <= best_sse * 1.01 + 1e-12]
    if len(close) < 8:
        raise AssertionError("Multistart regression did not repeatedly converge to the best basin.")

    completed = build_fitting_results_workbook(raw, problem, result)
    wb = load_workbook(BytesIO(completed), data_only=False, read_only=True)
    required = {"Fit_Parameters", "Fit_Starts", "Fit_Predictions", "Fit_Residuals", "Fit_Summary"}
    if not required.issubset(wb.sheetnames):
        raise AssertionError("Completed fit workbook is missing fitting result sheets.")
    if wb["Fit_Predictions"].max_row != len(problem.observations) + 1:
        raise AssertionError("Fit_Predictions row count does not match observations.")

    print("Kinetic fitting engine smoke test: PASS")
    print("Fit parameters:", dict(zip(problem.fit_parameter_ids, best)))
    print("Synthetic truth:", truth)
    print("Relative errors:", dict(zip(problem.fit_parameter_ids, rel)))
    print(f"Best SSE={result['best']['sse']:.9g}; near-best starts={len(close)}/{len(finite)}")


if __name__ == "__main__":
    main()
