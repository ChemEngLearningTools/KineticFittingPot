from __future__ import annotations

from pathlib import Path

import numpy as np
import pandas as pd
import plotly.graph_objects as go
import streamlit as st
from openpyxl import load_workbook

from utils.kinetic_fitting_engine import fitted_trajectory, parse_fitting_workbook
from utils.ui import chips, metric_card, navigate, teaching_table, workspace_header


ACCENT = "#f59e0b"
COLORS = ["#19D3F3", "#AB63FA", "#EF553B", "#00CC96", "#FFA15A", "#636EFA", "#FF6692"]
CASE_PATH = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "case_studies"
    / "KineticFittingPot_kinetic_fit_results_Biodiesel.xlsx"
)
DIAGNOSTICS_PATH = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "case_studies"
    / "KineticFittingPot_kinetic_fit_results_Biodiesel_model_diagnostics.xlsx"
)


def _dark(fig: go.Figure) -> go.Figure:
    fig.update_layout(
        template="plotly_white",
        paper_bgcolor="#ffffff",
        plot_bgcolor="#ffffff",
        font=dict(color="#17314f"),
        margin=dict(l=25, r=25, t=70, b=40),
        legend=dict(orientation="h", yanchor="bottom", y=1.02, xanchor="left", x=0),
        hovermode="x unified",
    )
    return fig


def _sheet_dataframe(wb, name: str) -> pd.DataFrame:
    ws = wb[name]
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return pd.DataFrame()
    return pd.DataFrame(rows[1:], columns=rows[0])


def _load_case():
    raw = CASE_PATH.read_bytes()
    problem = parse_fitting_workbook(raw)
    wb = load_workbook(CASE_PATH, data_only=True, read_only=True)
    tables = {
        name: _sheet_dataframe(wb, name)
        for name in [
            "Datasets",
            "Reactions",
            "Fit_Parameters",
            "Fit_Starts",
            "Fit_Predictions",
            "Fit_Residuals",
            "Fit_Summary",
        ]
    }
    wb.close()
    return raw, problem, tables


def _parameter_value_vector(problem, fit_parameters: pd.DataFrame) -> np.ndarray:
    lookup = {
        str(row["Parameter_ID"]): float(row["Fitted_Value"])
        for _, row in fit_parameters.iterrows()
    }
    return np.asarray([lookup[pid] for pid in problem.fit_parameter_ids], dtype=float)


def _trajectory_figure(problem, dataset, best_values: np.ndarray, predictions: pd.DataFrame) -> go.Figure:
    traj = fitted_trajectory(problem, dataset, best_values, points=500)
    measured_species = ["TG", "DG", "MG", "ME"]
    color_map = {sid: COLORS[i % len(COLORS)] for i, sid in enumerate(problem.model.species_ids)}
    fig = go.Figure()

    for sid in measured_species:
        sidx = problem.model.species_ids.index(sid)
        fig.add_trace(
            go.Scatter(
                x=traj["time_s"] / 60.0,
                y=traj["concentrations_kmol_m3"][:, sidx],
                mode="lines",
                name=f"{sid} fitted",
                line=dict(color=color_map[sid], width=2.3),
            )
        )
        subset = predictions[
            (predictions["Dataset_ID"] == dataset.dataset_id)
            & (predictions["Species_ID"] == sid)
        ]
        fig.add_trace(
            go.Scatter(
                x=subset["Time_s"].astype(float) / 60.0,
                y=subset["Measured_Value_kmol_m3"].astype(float),
                mode="markers",
                name=f"{sid} measured",
                marker=dict(
                    color=color_map[sid],
                    size=8,
                    symbol="circle-open",
                    line=dict(width=2),
                ),
            )
        )

    fig.update_layout(
        title=f"Experiment {dataset.dataset_id} — {dataset.description}",
        xaxis_title="Time (min)",
        yaxis_title="Concentration (kmol m⁻³)",
        height=500,
    )
    return _dark(fig)


def _multistart_range_table(starts: pd.DataFrame, fit_parameters: pd.DataFrame) -> list[list[str]]:
    rows = []
    fit_lookup = {
        str(row["Parameter_ID"]): float(row["Fitted_Value"])
        for _, row in fit_parameters.iterrows()
    }
    for pid in fit_lookup:
        start_values = starts[f"Start_{pid}"].astype(float)
        fit_values = starts[f"Fit_{pid}"].astype(float)
        start_factor = float(start_values.max() / start_values.min())
        final_spread_pct = 100.0 * float(fit_values.max() - fit_values.min()) / float(fit_lookup[pid])
        rows.append([
            pid,
            f"{start_values.min():.3e}",
            f"{start_values.max():.3e}",
            f"{start_factor:,.0f}×",
            f"{fit_lookup[pid]:.6e}",
            f"{final_spread_pct:.4f}%",
        ])
    return rows


def _load_diagnostics():
    raw = DIAGNOSTICS_PATH.read_bytes()
    wb = load_workbook(DIAGNOSTICS_PATH, data_only=True, read_only=True)
    names = [
        "Diag_Settings",
        "Diag_Invariants",
        "Diag_Balance_Definition",
        "Diag_Balance_Points",
        "Diag_Balance_Summary",
        "Diag_Recon_Observations",
        "Diag_Recon_Species",
        "Diag_Residual_Species",
        "Diag_Residual_Datasets",
        "Diag_Local_Summary",
        "Diag_Local_Parameters",
        "Diag_Parameter_Corr",
        "Diag_Weakest_Direction",
        "Diag_Profile_Summary",
        "Diag_Profile_Points",
        "Diag_Zero_Tests",
    ]
    tables = {name: _sheet_dataframe(wb, name) for name in names}
    wb.close()
    return raw, tables


def _balance_figure(points: pd.DataFrame) -> go.Figure:
    fig = go.Figure()
    for i, dataset_id in enumerate(["A", "B", "C"]):
        subset = points[points["Dataset"] == dataset_id].copy()
        fig.add_trace(
            go.Scatter(
                x=subset["Time_min"].astype(float),
                y=subset["Deviation_pct"].astype(float),
                mode="lines+markers",
                name=f"Dataset {dataset_id}",
                line=dict(color=COLORS[i], width=2.2),
                marker=dict(size=7),
                error_y=dict(
                    type="data",
                    array=2.0 * subset["Deviation_sigma_pct"].astype(float),
                    visible=True,
                    thickness=1.1,
                    width=3,
                    color=COLORS[i],
                ),
            )
        )
    fig.add_hline(y=0.0, line_dash="dash", line_color="#9ca3af")
    fig.update_layout(
        title="Stoichiometric-balance deviation with ±2σ measurement-uncertainty bars",
        xaxis_title="Time (min)",
        yaxis_title="Balance deviation from initial value (%)",
        height=490,
    )
    return _dark(fig)


def _reconciliation_figure(species: pd.DataFrame) -> go.Figure:
    order = ["TG", "DG", "MG", "ME"]
    data = species.set_index("Species").reindex(order).reset_index()
    fig = go.Figure()
    fig.add_trace(
        go.Bar(
            x=data["Species"],
            y=data["RMS |correction|/σ"].astype(float),
            name="RMS correction / σ",
            marker_color=COLORS[0],
        )
    )
    fig.add_trace(
        go.Bar(
            x=data["Species"],
            y=data["Maximum |correction|/σ"].astype(float),
            name="Maximum correction / σ",
            marker_color=COLORS[2],
        )
    )
    fig.add_hline(y=2.0, line_dash="dot", line_color="#fbbf24", annotation_text="2σ")
    fig.add_hline(y=3.0, line_dash="dot", line_color="#f87171", annotation_text="3σ")
    fig.update_layout(
        title="Stoichiometric data-reconciliation corrections by measured species",
        xaxis_title="Measured species",
        yaxis_title="Correction magnitude / assumed σ",
        barmode="group",
        height=470,
    )
    return _dark(fig)


def _correlation_figure(corr: pd.DataFrame) -> go.Figure:
    params = [c for c in corr.columns if c != "Parameter"]
    matrix = corr.set_index("Parameter").loc[params, params].astype(float)
    values = matrix.to_numpy()
    fig = go.Figure(
        data=go.Heatmap(
            z=values,
            x=params,
            y=params,
            zmin=-1.0,
            zmax=1.0,
            zmid=0.0,
            colorscale="RdBu_r",
            colorbar=dict(title="Correlation"),
            text=np.round(values, 2),
            texttemplate="%{text:.2f}",
            hovertemplate="%{y} / %{x}: %{z:.3f}<extra></extra>",
        )
    )
    fig.update_layout(
        title="Approximate local fitted-parameter correlation matrix",
        xaxis_title="Parameter",
        yaxis_title="Parameter",
        height=520,
    )
    return _dark(fig)


def _profile_figure(points: pd.DataFrame) -> go.Figure:
    fig = go.Figure()
    for i, pid in enumerate(["k1r", "k2r", "k3r"]):
        subset = points[points["Profile_parameter"] == pid].copy()
        fig.add_trace(
            go.Scatter(
                x=subset["Fixed_over_fitted"].astype(float),
                y=subset["Relative_increase_percent"].astype(float),
                mode="lines+markers",
                name=pid,
                line=dict(color=COLORS[i], width=2.4),
                marker=dict(size=8),
            )
        )
    fig.add_vline(x=1.0, line_dash="dash", line_color="#9ca3af")
    fig.add_hline(y=0.0, line_dash="dash", line_color="#9ca3af")
    fig.update_xaxes(type="log")
    fig.update_layout(
        title="Profile-objective deterioration when each reverse rate constant is constrained",
        xaxis_title="Fixed parameter / fitted parameter",
        yaxis_title="Increase in weighted objective SSE (%)",
        height=500,
    )
    return _dark(fig)


def render() -> None:
    workspace_header(
        "Kinetic Fitting Case Study — Biodiesel",
        "A real-data demonstration of how KineticFittingPot turns an Excel-defined reversible reaction model and three experiments into one bounded multistart kinetic fit.",
        section="Kinetic Fitting",
        accent=ACCENT,
    )

    chips([
        ("Case study", "blue"),
        ("Real experimental data", "blue"),
        ("3 experiments", ""),
        ("6 shared parameters", ""),
        ("20 multistarts", "warn"),
        ("Saved public results", "warn"),
    ])

    if not CASE_PATH.exists():
        st.error("The bundled biodiesel case-study workbook is missing from this KineticFittingPot build.")
        return

    raw, problem, tables = _load_case()
    datasets = tables["Datasets"]
    reactions = tables["Reactions"]
    fit_parameters = tables["Fit_Parameters"]
    starts = tables["Fit_Starts"]
    predictions = tables["Fit_Predictions"]
    summary = tables["Fit_Summary"]
    best_values = _parameter_value_vector(problem, fit_parameters)

    st.info(
        "This public page does not run the kinetic optimiser. It displays a completed fit generated with the full local KineticFittingPot kinetic-fitting workspace, then reuses the saved fitted parameters in the same forward ODE model for the trajectory plots below."
    )

    c1, c2 = st.columns([0.7, 0.3])
    with c1:
        st.markdown(
            "The example is based on biodiesel transesterification data described by **Hong Yuan, Xingxing Li and Bolun Yang**, "
            "*Kinetics for the Synthesis of Biodiesel Based on the Calculation of Hot Spot Temperatures in the Catalyst*, "
            "Industrial & Engineering Chemistry Research (2013)."
        )
        st.markdown("Source DOI: [10.1021/ie401370n](https://doi.org/10.1021/ie401370n)")
    with c2:
        st.download_button(
            "Download completed case-study workbook",
            data=raw,
            file_name=CASE_PATH.name,
            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            use_container_width=True,
        )

    st.markdown("## Why this is a useful fitting example")
    st.markdown(
        "This is not a synthetic parameter-recovery exercise. KineticFittingPot is asked to fit **one common kinetic parameter vector** to three real experimental datasets collected at the same temperature and catalyst loading but at different methanol:oil molar ratios. The measurements are imperfect, not every model state is measured, and an exact ODE trajectory is not expected to pass through every point."
    )

    st.markdown("## Experimental system")
    conditions = []
    for _, row in datasets.iterrows():
        conditions.append([
            str(row["Dataset_ID"]),
            f"{float(row['Methanol_to_Oil_Molar_Ratio']):.0f}:1",
            f"{float(row['Temperature_K']):.0f}",
            f"{float(row['Catalyst_Load_wt_pct']):.1f}",
            f"{float(row['c0_TG']):.4f}",
            f"{float(row['c0_MeOH']):.3f}",
        ])
    teaching_table(
        ["Experiment", "MeOH:oil", "Temperature (K)", "Catalyst (wt%)", "Initial TG (kmol m⁻³)", "Initial MeOH (kmol m⁻³)"],
        conditions,
        alignments=["center", "center", "center", "center", "right", "right"],
    )
    st.caption(
        "All three datasets are at 338 K and 7.5 wt% catalyst. Their principal experimental difference is the initial methanol:oil molar ratio."
    )

    st.markdown("## Reaction network")
    st.latex(r"TG + MeOH \rightleftharpoons DG + ME")
    st.latex(r"DG + MeOH \rightleftharpoons MG + ME")
    st.latex(r"MG + MeOH \rightleftharpoons Gly + ME")
    st.markdown(
        "TG, DG and MG denote tri-, di- and monoglyceride; ME is methyl ester (biodiesel), and Gly is glycerol."
    )

    st.warning(
        "The KineticFittingPot case study deliberately uses a simpler **reversible mass-action model with six apparent rate constants**. It is not a reproduction of the paper's LHHW kinetic model. Because datasets A/B/C all use the same catalyst loading and temperature, those effects are absorbed into the fitted apparent rate constants."
    )

    st.markdown("## The kinetic model is defined in Excel")
    rate_rows = []
    for _, row in reactions.iterrows():
        rate_rows.append([
            str(row["Reaction_ID"]),
            str(row["Description"]),
            str(row["Rate_Expression"]),
        ])
    teaching_table(
        ["Reaction", "Chemical step", "Workbook rate expression"],
        rate_rows,
        alignments=["center", "left", "left"],
        monospace_columns=[2],
    )
    st.markdown(
        "The workbook also contains the stoichiometric coefficients for every species. KineticFittingPot therefore reads the chemical network, kinetic expressions, parameter bounds, dataset-specific initial conditions and experimental observations from the workbook rather than from case-specific Python code."
    )

    st.markdown("### Measured and unmeasured model states")
    teaching_table(
        ["Species", "Role in this fit"],
        [
            ["TG", "Measured and fitted observation"],
            ["DG", "Measured and fitted observation"],
            ["MG", "Measured and fitted observation"],
            ["ME", "Measured and fitted observation"],
            ["MeOH", "Dynamic model state, but not a fitted observation"],
            ["Gly", "Dynamic model state, but not a fitted observation"],
        ],
        alignments=["center", "left"],
    )
    st.markdown(
        "Glycerol is intentionally excluded from the fitted observations because the legacy workbook values were stoichiometrically reconstructed rather than independently measured. Initial methanol is reconstructed from the nominal molar ratio using `c0_MeOH = (MeOH:oil ratio) × c0_TG`."
    )
    st.markdown(
        "A normalised constant liquid volume of **1.0 m³** is used for each dataset. In the constant-volume concentration model, the absolute volume cancels from the concentration ODEs, so this convention does not alter the fitted concentration trajectories."
    )

    st.markdown("## The inverse problem solved by KineticFittingPot")
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        metric_card("Datasets", "3", "A, B and C", compact=True)
    with c2:
        metric_card("Observations", "120", "actual measured rows", compact=True)
    with c3:
        metric_card("Shared parameters", "6", "same vector for all datasets", compact=True)
    with c4:
        metric_card("Multistarts", "20", "bounded least-squares runs", compact=True)

    st.markdown(
        "All six apparent rate constants are estimated simultaneously. Each experiment retains its own initial conditions and actual observation times, but the optimiser must find **one parameter set that explains all three experiments together**."
    )
    st.latex(r"\boldsymbol\theta=[k_{1f},k_{1r},k_{2f},k_{2r},k_{3f},k_{3r}]")
    st.markdown(
        "No pointwise experimental standard deviations were supplied. The saved run therefore used KineticFittingPot's **species-scaled residuals**, with finite parameter bounds from `1×10⁻⁷` to `1×10⁻² m³ kmol⁻¹ s⁻¹`, and 20 multistart locations distributed across those bounds."
    )

    st.markdown("## Fitted kinetic parameters")
    parameter_rows = []
    for _, row in fit_parameters.iterrows():
        parameter_rows.append([
            str(row["Parameter_ID"]),
            f"{float(row['Starting_Value']):.3e}",
            f"{float(row['Fitted_Value']):.6e}",
            f"{float(row['Lower_Bound']):.1e}",
            f"{float(row['Upper_Bound']):.1e}",
            str(row["Units"]),
        ])
    teaching_table(
        ["Parameter", "Workbook start", "Best fitted value", "Lower bound", "Upper bound", "Units"],
        parameter_rows,
        alignments=["center", "right", "right", "right", "right", "center"],
        monospace_columns=[0],
    )

    summary_lookup = {str(r[0]): r[1] for r in summary.itertuples(index=False, name=None)}
    st.markdown("### Fit summary")
    c1, c2, c3, c4 = st.columns(4)
    with c1:
        metric_card("Weighted RMSE", f"{float(summary_lookup['Best weighted RMSE']):.4f}", "species-scaled residuals", compact=True)
    with c2:
        metric_card("Raw RMSE", f"{float(summary_lookup['Best raw RMSE']):.4f}", "kmol m⁻³", compact=True)
    with c3:
        metric_card("Best start", str(int(summary_lookup["Best start"])), "of 20 runs", compact=True)
    with c4:
        metric_card("Optimiser", "Success", "ftol termination", compact=True)

    st.markdown("## Measured data and fitted trajectories")
    st.markdown(
        "The lines below are generated by rerunning the **same KineticFittingPot forward ODE model** with the saved best-fit parameter vector. The open markers are the original experimental observations. No parameter optimisation is performed on this page."
    )
    tabs = st.tabs([
        "A — 18:1 MeOH:oil",
        "B — 24:1 MeOH:oil",
        "C — 30:1 MeOH:oil",
    ])
    for tab, dataset in zip(tabs, problem.datasets):
        with tab:
            st.plotly_chart(
                _trajectory_figure(problem, dataset, best_values, predictions),
                use_container_width=True,
            )

    st.markdown("## Why multistart matters")
    st.markdown(
        "Nonlinear kinetic fitting can contain local minima and parameter trade-offs. KineticFittingPot therefore does not rely on one initial guess. The saved biodiesel fit used 20 starting points spread widely across the allowed parameter bounds."
    )
    successes = int(starts["Success"].astype(bool).sum())
    best_sse = float(starts["SSE"].astype(float).min())
    worst_sse = float(starts["SSE"].astype(float).max())
    sse_spread = worst_sse - best_sse
    c1, c2, c3 = st.columns(3)
    with c1:
        metric_card("Successful starts", f"{successes}/20", "all reached an optimum", compact=True)
    with c2:
        metric_card("Best weighted SSE", f"{best_sse:.9f}", "objective value", compact=True)
    with c3:
        metric_card("SSE spread", f"{sse_spread:.2e}", "best to worst final solution", compact=True)

    teaching_table(
        ["Parameter", "Smallest start", "Largest start", "Start range", "Best fit", "Spread of final fits"],
        _multistart_range_table(starts, fit_parameters),
        alignments=["center", "right", "right", "right", "right", "right"],
        monospace_columns=[0],
    )
    st.markdown(
        "The starting estimates span roughly **four to five orders of magnitude**, yet every run converges to essentially the same fitted parameter vector and objective. That is strong evidence that this optimiser run is not simply reporting a result tied to one convenient initial guess."
    )
    st.warning(
        "Multistart agreement is a numerical optimisation check — not proof that every fitted parameter is independently identifiable. Parameter correlation and profile-objective diagnostics answer a different question."
    )

    st.markdown("## A good fit is not the end of model development")
    st.markdown(
        "Parameter estimation answers: **which parameter values give the best agreement between this model and these data?** It does not establish that the measurements are internally consistent, that the residuals resemble experimental error, that the six parameters are independently identifiable, or that every reversible kinetic term is supported by the data."
    )

    if not DIAGNOSTICS_PATH.exists():
        st.warning("The bundled completed diagnostics workbook is missing, so the diagnostic case-study section cannot be displayed in this build.")
    else:
        diagnostics_raw, diag = _load_diagnostics()
        settings = {str(r[0]): r[1] for r in diag["Diag_Settings"].itertuples(index=False, name=None)}
        local_summary = {str(r[0]): r[1] for r in diag["Diag_Local_Summary"].itertuples(index=False, name=None)}
        relative_sigma = 100.0 * float(settings["Assumed relative measurement uncertainty (fraction, 1 sigma)"])
        absolute_sigma = float(settings["Absolute measurement uncertainty floor (kmol m^-3, 1 sigma)"])

        st.markdown("## Post-fit diagnostics — what does the fitted model actually tell us?")
        c1, c2 = st.columns([0.72, 0.28])
        with c1:
            st.markdown(
                "The completed local diagnostics run now challenges the result in four stages: **the data**, **the reconciliation corrections**, **the fitted parameters**, and finally **the reversible model structure**. The public page reads the saved diagnostic record; it does not rerun Jacobian calculations or profile re-optimisation."
            )
        with c2:
            st.download_button(
                "Download complete diagnostics workbook",
                data=diagnostics_raw,
                file_name=DIAGNOSTICS_PATH.name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                use_container_width=True,
            )

        st.caption(
            f"Where experimental standard deviations were unavailable, this diagnostic run assumed {relative_sigma:.0f}% relative 1σ uncertainty plus an absolute {absolute_sigma:.3f} kmol m⁻³ floor, combined in quadrature. These assumptions support engineering diagnostics; they are not experimentally established error distributions."
        )

        st.markdown("### 1. Are the experimental measurements chemically self-consistent?")
        st.markdown(
            "Stoichiometry provides a test that does **not** depend on the fitted kinetic rate constants. For this six-species network KineticFittingPot finds three independent conserved quantities. Only one is fully observable from TG, DG, MG and ME:"
        )
        st.latex(r"3TG + 2DG + MG + ME = \mathrm{constant}")
        st.markdown(
            "If all four measured concentrations were perfectly compatible with the stated closed-batch reaction network, this measured balance would remain constant. Deviations therefore reveal internal inconsistency in the experimental observations before any kinetic-model judgement is made."
        )
        st.plotly_chart(_balance_figure(diag["Diag_Balance_Points"]), use_container_width=True)

        balance_rows = []
        for _, row in diag["Diag_Balance_Summary"].iterrows():
            balance_rows.append([
                str(row["Dataset"]),
                f"{float(row['Maximum |deviation| (%)']):.1f}%",
                f"{float(row['RMS deviation (%)']):.1f}%",
                str(int(row["Points outside ±2σ"])),
                f"{float(row['Maximum |deviation|/σ']):.2f}σ",
            ])
        teaching_table(
            ["Dataset", "Maximum deviation", "RMS deviation", "Points outside ±2σ", "Largest deviation / σ"],
            balance_rows,
            alignments=["center", "right", "right", "center", "right"],
        )
        st.markdown(
            "Dataset **A** remains within the illustrative ±2σ uncertainty envelope. Dataset **B** contains one point outside it and dataset **C** contains two. The measurements are therefore broadly compatible with the chemistry, but B and especially C contain observations that cannot all be simultaneously reconciled with the measured stoichiometric balance within the assumed uncertainty model."
        )

        st.markdown("### 2. Which measurements drive the stoichiometric inconsistency?")
        st.markdown(
            "KineticFittingPot next performs weighted stoichiometric data reconciliation. This step does **not** change the kinetic model to improve the fit. It calculates the minimum changes to the measured concentrations required to restore the measured reaction balance, weighted by the assumed measurement uncertainty."
        )
        st.plotly_chart(_reconciliation_figure(diag["Diag_Recon_Species"]), use_container_width=True)

        recon_obs = diag["Diag_Recon_Observations"].copy()
        recon_obs["abs_z"] = recon_obs["Correction_over_sigma"].astype(float).abs()
        count_rows = []
        for sid in ["TG", "DG", "MG", "ME"]:
            subset = recon_obs[recon_obs["Species"] == sid]
            rms_z = float(np.sqrt(np.mean(subset["abs_z"].to_numpy() ** 2)))
            max_z = float(subset["abs_z"].max())
            count_rows.append([
                sid,
                f"{rms_z:.2f}σ",
                f"{max_z:.2f}σ",
                str(int((subset["abs_z"] > 2.0).sum())),
                str(int((subset["abs_z"] > 3.0).sum())),
            ])
        teaching_table(
            ["Species", "RMS correction / σ", "Maximum correction / σ", "Measurements >2σ", "Measurements >3σ"],
            count_rows,
            alignments=["center", "right", "right", "center", "center"],
        )
        st.warning(
            "The reconciliation signal is strongly concentrated in **ME**: 12 of 30 ME observations require corrections larger than 2σ and 5 exceed 3σ, whereas DG and MG never exceed 2σ and TG does so once. This does **not** prove that ME alone was measured incorrectly — the allocation depends on the assumed uncertainties — but it makes the ME assay, calibration and uncertainty the first experimental issue to investigate."
        )

        st.markdown("### Does that mean the kinetic model fits ME poorly?")
        residual_species = diag["Diag_Residual_Species"]
        residual_rows = []
        for sid in ["TG", "DG", "MG", "ME"]:
            row = residual_species[residual_species["Species"] == sid].iloc[0]
            residual_rows.append([
                sid,
                f"{float(row['RMSE (kmol m⁻³)']):.4f}",
                f"{float(row['Fit-scaled RMSE']):.4f}",
                f"{float(row['Bias, model − measured (kmol m⁻³)']):+.4f}",
            ])
        teaching_table(
            ["Species", "Raw RMSE (kmol m⁻³)", "Fit-scaled RMSE", "Bias (model − measured)"],
            residual_rows,
            alignments=["center", "right", "right", "right"],
        )
        dataset_fit = diag["Diag_Residual_Datasets"]
        ds_text = ", ".join(
            f"{str(row['Dataset'])} = {float(row['Fit-scaled RMSE']):.4f}"
            for _, row in dataset_fit.iterrows()
        )
        st.markdown(
            f"The answer is **no**. ME has the largest absolute RMSE because its concentration scale is much larger, but its fit-scaled RMSE is only {float(residual_species[residual_species['Species'] == 'ME']['Fit-scaled RMSE'].iloc[0]):.4f}. The three experiments also have similar fit-scaled errors ({ds_text}). The ME concern therefore arises independently from the stoichiometric consistency check, rather than simply because the model cannot reproduce the ME trajectory."
        )
        st.caption(
            "Residuals standardised by the illustrative 3% + absolute-floor uncertainty model should not be interpreted as a formal statistical goodness-of-fit test, because those uncertainties were assumed rather than supplied with the experimental dataset."
        )

        st.markdown("### 3. Are the six fitted parameters locally determined by the data?")
        local_parameters = diag["Diag_Local_Parameters"]
        weakest = diag["Diag_Weakest_Direction"].copy()
        weakest["abs_coeff"] = weakest["Weakest_direction_coefficient"].astype(float).abs()
        weakest_pid = str(weakest.sort_values("abs_coeff", ascending=False).iloc[0]["Parameter"])
        k1r_row = local_parameters[local_parameters["Parameter"] == "k1r"].iloc[0]
        strongest_corr = float(local_parameters["Strongest correlation"].astype(float).abs().max())
        c1, c2, c3, c4 = st.columns(4)
        with c1:
            metric_card("Jacobian rank", f"{int(local_summary['Numerical rank'])}/{int(local_summary['Fitted parameters'])}", "full local rank", compact=True)
        with c2:
            metric_card("Condition number", f"{float(local_summary['Condition number']):.2f}", "scaled local Jacobian", compact=True)
        with c3:
            metric_card("Largest local relative SE", f"{float(k1r_row['Approx. relative SE (%)']):.1f}%", "k1r", compact=True)
        with c4:
            metric_card("Strongest |correlation|", f"{strongest_corr:.2f}", "forward/reverse pair", compact=True)

        st.plotly_chart(_correlation_figure(diag["Diag_Parameter_Corr"]), use_container_width=True)
        param_rows = []
        for _, row in local_parameters.iterrows():
            param_rows.append([
                str(row["Parameter"]),
                f"{float(row['Approx. relative SE (%)']):.1f}%",
                str(row["Strongest partner"]),
                f"{float(row['Strongest correlation']):.3f}",
            ])
        teaching_table(
            ["Parameter", "Approx. local relative SE", "Strongest partner", "Correlation"],
            param_rows,
            alignments=["center", "right", "center", "right"],
            monospace_columns=[0],
        )
        st.markdown(
            f"The scaled sensitivity matrix has **full numerical rank**, so the fitted parameter vector is locally identifiable at the optimum. The expected forward/reverse trade-offs are visible — particularly k2f/k2r and k3f/k3r — but the system is not rank deficient. **k1r** is the least precisely determined parameter (about {float(k1r_row['Approx. relative SE (%)']):.1f}% approximate local relative SE), and the weakest singular-vector direction is dominated by **{weakest_pid}**."
        )
        st.info(
            "Local covariance and correlation diagnostics describe the neighbourhood of the fitted solution. They are useful, but they do not by themselves tell us how the objective behaves when a parameter is moved far from its optimum or removed from the model. That is the role of profiling."
        )

        st.markdown("### 4. Do the data support the reverse kinetic terms?")
        st.markdown(
            "KineticFittingPot therefore profiles the three reverse rate constants. Each selected parameter is fixed over a multiplicative range around its fitted value while **all other fitted parameters are re-optimised**. The saved profiles have their minima at the fitted values and did not uncover a better solution elsewhere in the explored range."
        )
        st.plotly_chart(_profile_figure(diag["Diag_Profile_Points"]), use_container_width=True)

        zero_rows = []
        for _, row in diag["Diag_Zero_Tests"].iterrows():
            zero_rows.append([
                str(row["Profile_parameter"]),
                f"{float(row['Fitted_value']):.3e}",
                f"{float(row['Objective_SSE']):.4f}",
                f"{float(row['Relative_increase_percent']):.1f}%",
                f"{float(row['Weighted_RMSE']):.4f}",
            ])
        teaching_table(
            ["Reverse parameter set to zero", "Fitted value", "Re-optimised SSE", "Objective increase", "Weighted RMSE"],
            zero_rows,
            alignments=["center", "right", "right", "right", "right"],
            monospace_columns=[0],
        )
        st.markdown(
            "Removing **k2r** or **k3r** produces a large deterioration in the best achievable objective (about 159% and 178% respectively). Removing **k1r** gives a smaller but still material deterioration of about 23%. This is consistent with the local diagnostics: reverse reaction 1 is less tightly constrained, while the data provide stronger support for reversibility in reactions 2 and 3."
        )
        st.caption(
            "The objective changes are engineering model-comparison diagnostics. KineticFittingPot does not label them 'statistically significant' because a formal likelihood model based on experimentally established measurement-error distributions is not available here."
        )

        st.markdown("## Case-study conclusion")
        st.success(
            "**The simple reversible mass-action model is supported by the available experimental data as a useful kinetic representation over the conditions studied.** One six-parameter vector describes all three methanol:oil experiments, the local sensitivity matrix has full rank, the profile minima coincide with the fitted solution, and eliminating reverse rate terms worsens the best achievable fit — particularly for reactions 2 and 3."
        )
        st.warning(
            "The diagnostics also show why a visually good kinetic fit should not be the end of the analysis. The independently derived stoichiometric balance reveals inconsistencies in some measurements, and weighted reconciliation concentrates those corrections in **ME** under the assumed uncertainty model. This does not establish that ME alone is wrong, but it provides a clear experimental lead: the ME measurement method, calibration and uncertainty would be the first items to investigate."
        )
        st.markdown(
            "KineticFittingPot has therefore moved beyond curve fitting: it has used **stoichiometry to challenge the data, sensitivity and correlation to challenge the parameters, and profile re-optimisation to challenge the kinetic structure itself**."
        )

    st.latex(
        r"\boxed{\text{Excel-defined model}}\rightarrow"
        r"\boxed{\text{multi-experiment fit}}\rightarrow"
        r"\boxed{\text{challenge the data}}\rightarrow"
        r"\boxed{\text{challenge the parameters}}\rightarrow"
        r"\boxed{\text{challenge the model}}"
    )
    st.divider()
    c1, c2, c3 = st.columns(3)
    with c1:
        if st.button("← Kinetic Fitting Fundamentals", key="biodiesel_case_to_fundamentals"):
            navigate("Kinetic Fitting Fundamentals")
    with c2:
        if st.button("Open Kinetic Parameter Fitting", key="biodiesel_case_to_fit"):
            navigate("Kinetic Parameter Fitting")
    with c3:
        if st.button("Model Diagnostics Fundamentals →", key="biodiesel_case_to_diag_fund"):
            navigate("Model Diagnostics Fundamentals")
