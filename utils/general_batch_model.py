from __future__ import annotations

import ast
import math
import re
import time
from dataclasses import dataclass
from io import BytesIO
from typing import Iterable

import numpy as np
try:
    from scipy.integrate import solve_ivp
    from scipy.optimize import minimize_scalar
except ImportError:  # app shell can still start before dependencies are installed
    solve_ivp = None
    minimize_scalar = None

try:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill
except ImportError:  # template page can explain the missing dependency cleanly
    load_workbook = None


_IDENTIFIER = re.compile(r"^[A-Za-z][A-Za-z0-9_]*$")
_ALLOWED_FUNCTIONS = {
    "exp": np.exp,
    "log": np.log,
    "sqrt": np.sqrt,
    "abs": np.abs,
}
_ALLOWED_BINOPS = (ast.Add, ast.Sub, ast.Mult, ast.Div, ast.Pow)
_ALLOWED_UNARYOPS = (ast.UAdd, ast.USub)


@dataclass(frozen=True)
class GeneralBatchModel:
    model_name: str
    description: str
    temperature_c: float
    initial_volume_m3: float
    batch_duration_s: float
    output_points: int
    volume_model: str
    species_ids: tuple[str, ...]
    display_names: tuple[str, ...]
    initial_concentration_kmol_m3: np.ndarray
    molecular_weight_kg_per_kmol: np.ndarray
    density_kg_per_m3: np.ndarray
    plot_default: tuple[bool, ...]
    parameter_ids: tuple[str, ...]
    parameter_values: np.ndarray
    parameter_units: tuple[str, ...]
    fit_later: tuple[bool, ...]
    lower_bounds: tuple[float | None, ...]
    upper_bounds: tuple[float | None, ...]
    reaction_ids: tuple[str, ...]
    reaction_descriptions: tuple[str, ...]
    rate_expressions: tuple[str, ...]
    stoichiometric_matrix: np.ndarray  # species x reactions
    compiled_rates: tuple[object, ...]

    @property
    def initial_moles_kmol(self) -> np.ndarray:
        return self.initial_concentration_kmol_m3 * self.initial_volume_m3

    @property
    def parameter_map(self) -> dict[str, float]:
        return {key: float(value) for key, value in zip(self.parameter_ids, self.parameter_values)}


class ModelValidationError(ValueError):
    pass


def _literal(value: object, label: str) -> object:
    if isinstance(value, str) and value.startswith("="):
        raise ModelValidationError(f"{label} contains an Excel formula. Model-definition cells must contain literal values.")
    return value


def _as_float(value: object, label: str) -> float:
    value = _literal(value, label)
    try:
        number = float(value)
    except (TypeError, ValueError) as exc:
        raise ModelValidationError(f"{label} must be numeric.") from exc
    if not np.isfinite(number):
        raise ModelValidationError(f"{label} must be finite.")
    return number


def _as_bool(value: object, label: str) -> bool:
    value = _literal(value, label)
    if isinstance(value, bool):
        return value
    text = str(value or "").strip().lower()
    if text in {"yes", "y", "true", "1"}:
        return True
    if text in {"no", "n", "false", "0", ""}:
        return False
    raise ModelValidationError(f"{label} must be Yes/No (or True/False).")


def _check_identifier(value: object, label: str, *, forbid_c_prefix: bool = False) -> str:
    value = _literal(value, label)
    text = str(value or "").strip()
    if not _IDENTIFIER.fullmatch(text):
        raise ModelValidationError(
            f"{label}='{text}' is not a valid identifier. Use letters/numbers/underscore and start with a letter."
        )
    if forbid_c_prefix and text.startswith("c_"):
        raise ModelValidationError(f"{label}='{text}' cannot start with 'c_' because that prefix is reserved for concentrations.")
    if text in _ALLOWED_FUNCTIONS:
        raise ModelValidationError(f"{label}='{text}' conflicts with a reserved function name.")
    return text


def _header_map(ws) -> dict[str, int]:
    values = [cell.value for cell in ws[1]]
    out: dict[str, int] = {}
    for idx, value in enumerate(values, start=1):
        if value is None:
            continue
        key = str(value).strip()
        if key:
            out[key] = idx
    return out


def _iter_nonblank_rows(ws, start_row: int = 2) -> Iterable[int]:
    for row in range(start_row, ws.max_row + 1):
        values = [ws.cell(row=row, column=col).value for col in range(1, ws.max_column + 1)]
        if any(value not in (None, "") for value in values):
            yield row


def _validate_expression(expression: str, allowed_names: set[str], label: str) -> object:
    if not expression or len(expression) > 250:
        raise ModelValidationError(f"{label} must contain a rate expression no longer than 250 characters.")
    try:
        tree = ast.parse(expression, mode="eval")
    except SyntaxError as exc:
        raise ModelValidationError(f"{label} is not valid mathematical syntax: {exc.msg}.") from exc

    for node in ast.walk(tree):
        if isinstance(node, ast.Expression):
            continue
        if isinstance(node, ast.BinOp):
            if not isinstance(node.op, _ALLOWED_BINOPS):
                raise ModelValidationError(f"{label} uses an unsupported arithmetic operator.")
            continue
        if isinstance(node, _ALLOWED_BINOPS):
            continue
        if isinstance(node, ast.UnaryOp):
            if not isinstance(node.op, _ALLOWED_UNARYOPS):
                raise ModelValidationError(f"{label} uses an unsupported unary operator.")
            continue
        if isinstance(node, _ALLOWED_UNARYOPS):
            continue
        if isinstance(node, ast.Constant):
            if not isinstance(node.value, (int, float)):
                raise ModelValidationError(f"{label} contains a non-numeric constant.")
            continue
        if isinstance(node, ast.Name):
            if node.id not in allowed_names and node.id not in _ALLOWED_FUNCTIONS:
                raise ModelValidationError(f"{label} refers to unknown name '{node.id}'.")
            continue
        if isinstance(node, ast.Call):
            if not isinstance(node.func, ast.Name) or node.func.id not in _ALLOWED_FUNCTIONS:
                raise ModelValidationError(f"{label} may call only exp(), log(), sqrt() or abs().")
            if node.keywords or len(node.args) != 1:
                raise ModelValidationError(f"{label}: function calls must have exactly one positional argument.")
            continue
        if isinstance(node, ast.Load):
            continue
        raise ModelValidationError(f"{label} contains unsupported syntax ({type(node).__name__}).")

    return compile(tree, filename="<KineticFittingPot rate expression>", mode="eval")


def parse_model_workbook(workbook_bytes: bytes) -> GeneralBatchModel:
    if load_workbook is None:
        raise RuntimeError(
            "Excel model support requires openpyxl. Run update_KineticFittingPot_dependencies.bat once after upgrading KineticFittingPot."
        )
    if not workbook_bytes:
        raise ModelValidationError("The uploaded workbook is empty.")
    if len(workbook_bytes) > 10 * 1024 * 1024:
        raise ModelValidationError("The workbook is larger than the 10 MB teaching-workflow limit.")

    try:
        wb = load_workbook(BytesIO(workbook_bytes), data_only=False, read_only=False, keep_links=False)
    except Exception as exc:
        raise ModelValidationError(f"KineticFittingPot could not read this .xlsx workbook: {exc}") from exc

    required = ["Model", "Species", "Parameters", "Reactions"]
    missing = [name for name in required if name not in wb.sheetnames]
    if missing:
        raise ModelValidationError("Missing required worksheet(s): " + ", ".join(missing))

    # Model key/value sheet -------------------------------------------------
    ws = wb["Model"]
    model_values: dict[str, object] = {}
    for row in _iter_nonblank_rows(ws, 2):
        key = str(ws.cell(row=row, column=1).value or "").strip()
        if not key:
            continue
        if key in model_values:
            raise ModelValidationError(f"Model sheet contains duplicate key '{key}'.")
        model_values[key] = _literal(ws.cell(row=row, column=2).value, f"Model[{key}]")

    required_keys = [
        "Model_Name",
        "Temperature_C",
        "Initial_Volume_m3",
        "Batch_Duration_s",
        "Output_Points",
        "Volume_Model",
    ]
    missing_keys = [key for key in required_keys if key not in model_values]
    if missing_keys:
        raise ModelValidationError("Model sheet is missing key(s): " + ", ".join(missing_keys))

    model_name = str(model_values["Model_Name"] or "").strip()
    if not model_name:
        raise ModelValidationError("Model_Name cannot be blank.")
    description = str(model_values.get("Description", "") or "").strip()
    temperature_c = _as_float(model_values["Temperature_C"], "Temperature_C")
    initial_volume = _as_float(model_values["Initial_Volume_m3"], "Initial_Volume_m3")
    batch_duration = _as_float(model_values["Batch_Duration_s"], "Batch_Duration_s")
    output_points_f = _as_float(model_values["Output_Points"], "Output_Points")
    output_points = int(round(output_points_f))
    volume_model = str(model_values["Volume_Model"] or "").strip().lower()

    if initial_volume <= 0:
        raise ModelValidationError("Initial_Volume_m3 must be greater than zero.")
    if batch_duration <= 0:
        raise ModelValidationError("Batch_Duration_s must be greater than zero.")
    if output_points < 100 or output_points > 10000:
        raise ModelValidationError("Output_Points must be between 100 and 10,000.")
    if abs(output_points_f - output_points) > 1e-9:
        raise ModelValidationError("Output_Points must be an integer.")
    if volume_model not in {"constant", "composition_dependent"}:
        raise ModelValidationError("Volume_Model must be 'constant' or 'composition_dependent'.")

    # Species ---------------------------------------------------------------
    ws = wb["Species"]
    headers = _header_map(ws)
    required_headers = [
        "Species_ID",
        "Display_Name",
        "Initial_Concentration_kmol_m3",
        "Molecular_Weight_kg_kmol",
        "Density_kg_m3",
        "Plot_Default",
    ]
    missing_headers = [key for key in required_headers if key not in headers]
    if missing_headers:
        raise ModelValidationError("Species sheet is missing column(s): " + ", ".join(missing_headers))

    species_ids: list[str] = []
    display_names: list[str] = []
    initial_conc: list[float] = []
    molecular_weights: list[float] = []
    densities: list[float] = []
    plot_default: list[bool] = []
    for row in _iter_nonblank_rows(ws, 2):
        sid = _check_identifier(ws.cell(row=row, column=headers["Species_ID"]).value, f"Species row {row} ID")
        if sid in species_ids:
            raise ModelValidationError(f"Duplicate Species_ID '{sid}'.")
        display = str(_literal(ws.cell(row=row, column=headers["Display_Name"]).value, f"Species {sid} Display_Name") or sid).strip()
        c0 = _as_float(ws.cell(row=row, column=headers["Initial_Concentration_kmol_m3"]).value, f"Species {sid} initial concentration")
        if c0 < 0:
            raise ModelValidationError(f"Species {sid} initial concentration cannot be negative.")
        mw = _as_float(
            ws.cell(row=row, column=headers["Molecular_Weight_kg_kmol"]).value,
            f"Species {sid} molecular weight",
        )
        rho = _as_float(
            ws.cell(row=row, column=headers["Density_kg_m3"]).value,
            f"Species {sid} density",
        )
        if mw <= 0:
            raise ModelValidationError(f"Species {sid} molecular weight must be greater than zero.")
        if rho <= 0:
            raise ModelValidationError(f"Species {sid} density must be greater than zero.")
        pdefault = _as_bool(ws.cell(row=row, column=headers["Plot_Default"]).value, f"Species {sid} Plot_Default")
        species_ids.append(sid)
        display_names.append(display or sid)
        initial_conc.append(c0)
        molecular_weights.append(mw)
        densities.append(rho)
        plot_default.append(pdefault)

    if not species_ids:
        raise ModelValidationError("Species sheet does not contain any species.")
    if len(species_ids) > 30:
        raise ModelValidationError("The teaching simulator currently supports at most 30 species.")

    if volume_model == "composition_dependent":
        initial_moles_check = np.asarray(initial_conc, dtype=float) * initial_volume
        implied_volume = float(
            np.sum(
                initial_moles_check
                * np.asarray(molecular_weights, dtype=float)
                / np.asarray(densities, dtype=float)
            )
        )
        relative_mismatch = abs(implied_volume - initial_volume) / initial_volume
        if relative_mismatch > 1e-3:
            raise ModelValidationError(
                "For Volume_Model='composition_dependent', the initial concentrations, Initial_Volume_m3, "
                "molecular weights and densities must describe the same initial charge. "
                f"They imply {implied_volume:.6g} m3 versus Initial_Volume_m3={initial_volume:.6g} m3 "
                f"({100*relative_mismatch:.3g}% mismatch)."
            )

    # Parameters ------------------------------------------------------------
    ws = wb["Parameters"]
    headers = _header_map(ws)
    required_headers = ["Parameter_ID", "Value", "Units"]
    missing_headers = [key for key in required_headers if key not in headers]
    if missing_headers:
        raise ModelValidationError("Parameters sheet is missing column(s): " + ", ".join(missing_headers))

    parameter_ids: list[str] = []
    parameter_values: list[float] = []
    parameter_units: list[str] = []
    fit_later: list[bool] = []
    lower_bounds: list[float | None] = []
    upper_bounds: list[float | None] = []

    for row in _iter_nonblank_rows(ws, 2):
        pid = _check_identifier(
            ws.cell(row=row, column=headers["Parameter_ID"]).value,
            f"Parameter row {row} ID",
            forbid_c_prefix=True,
        )
        if pid in parameter_ids:
            raise ModelValidationError(f"Duplicate Parameter_ID '{pid}'.")
        value = _as_float(ws.cell(row=row, column=headers["Value"]).value, f"Parameter {pid} value")
        units = str(_literal(ws.cell(row=row, column=headers["Units"]).value, f"Parameter {pid} units") or "").strip()
        fit = False
        if "Fit_Later" in headers:
            fit = _as_bool(ws.cell(row=row, column=headers["Fit_Later"]).value, f"Parameter {pid} Fit_Later")

        def optional_bound(name: str) -> float | None:
            if name not in headers:
                return None
            raw = ws.cell(row=row, column=headers[name]).value
            if raw in (None, ""):
                return None
            return _as_float(raw, f"Parameter {pid} {name}")

        lb = optional_bound("Lower_Bound")
        ub = optional_bound("Upper_Bound")
        if lb is not None and ub is not None and lb >= ub:
            raise ModelValidationError(f"Parameter {pid}: Lower_Bound must be less than Upper_Bound.")
        parameter_ids.append(pid)
        parameter_values.append(value)
        parameter_units.append(units)
        fit_later.append(fit)
        lower_bounds.append(lb)
        upper_bounds.append(ub)

    if len(parameter_ids) > 100:
        raise ModelValidationError("The teaching simulator currently supports at most 100 parameters.")

    concentration_names = {f"c_{sid}" for sid in species_ids}
    overlap = concentration_names.intersection(parameter_ids)
    if overlap:
        raise ModelValidationError("Parameter names conflict with concentration variables: " + ", ".join(sorted(overlap)))
    allowed_names = set(parameter_ids) | concentration_names

    # Reactions -------------------------------------------------------------
    ws = wb["Reactions"]
    headers = _header_map(ws)
    required_headers = ["Reaction_ID", "Description", "Rate_Expression", *species_ids]
    missing_headers = [key for key in required_headers if key not in headers]
    if missing_headers:
        raise ModelValidationError(
            "Reactions sheet is missing required column(s): " + ", ".join(missing_headers)
        )

    reaction_ids: list[str] = []
    reaction_descriptions: list[str] = []
    rate_expressions: list[str] = []
    compiled_rates: list[object] = []
    stoich_rows: list[list[float]] = []

    for row in _iter_nonblank_rows(ws, 2):
        rid = _check_identifier(ws.cell(row=row, column=headers["Reaction_ID"]).value, f"Reaction row {row} ID")
        if rid in reaction_ids:
            raise ModelValidationError(f"Duplicate Reaction_ID '{rid}'.")
        description_value = _literal(ws.cell(row=row, column=headers["Description"]).value, f"Reaction {rid} description")
        reaction_description = str(description_value or rid).strip()
        expr_value = _literal(ws.cell(row=row, column=headers["Rate_Expression"]).value, f"Reaction {rid} rate expression")
        expression = str(expr_value or "").strip()
        compiled = _validate_expression(expression, allowed_names, f"Reaction {rid} rate expression")
        stoich = [
            _as_float(ws.cell(row=row, column=headers[sid]).value, f"Reaction {rid} stoichiometry for {sid}")
            for sid in species_ids
        ]
        if not any(abs(value) > 0 for value in stoich):
            raise ModelValidationError(f"Reaction {rid} has all-zero stoichiometry.")
        reaction_ids.append(rid)
        reaction_descriptions.append(reaction_description)
        rate_expressions.append(expression)
        compiled_rates.append(compiled)
        stoich_rows.append(stoich)

    if not reaction_ids:
        raise ModelValidationError("Reactions sheet does not contain any reactions.")
    if len(reaction_ids) > 30:
        raise ModelValidationError("The teaching simulator currently supports at most 30 reactions.")

    stoich_matrix = np.asarray(stoich_rows, dtype=float).T

    return GeneralBatchModel(
        model_name=model_name,
        description=description,
        temperature_c=temperature_c,
        initial_volume_m3=initial_volume,
        batch_duration_s=batch_duration,
        output_points=output_points,
        volume_model=volume_model,
        species_ids=tuple(species_ids),
        display_names=tuple(display_names),
        initial_concentration_kmol_m3=np.asarray(initial_conc, dtype=float),
        molecular_weight_kg_per_kmol=np.asarray(molecular_weights, dtype=float),
        density_kg_per_m3=np.asarray(densities, dtype=float),
        plot_default=tuple(plot_default),
        parameter_ids=tuple(parameter_ids),
        parameter_values=np.asarray(parameter_values, dtype=float),
        parameter_units=tuple(parameter_units),
        fit_later=tuple(fit_later),
        lower_bounds=tuple(lower_bounds),
        upper_bounds=tuple(upper_bounds),
        reaction_ids=tuple(reaction_ids),
        reaction_descriptions=tuple(reaction_descriptions),
        rate_expressions=tuple(rate_expressions),
        stoichiometric_matrix=stoich_matrix,
        compiled_rates=tuple(compiled_rates),
    )


def evaluate_rates(model: GeneralBatchModel, concentrations_kmol_m3: np.ndarray) -> np.ndarray:
    c = np.maximum(np.asarray(concentrations_kmol_m3, dtype=float), 0.0)
    if c.shape != (len(model.species_ids),):
        raise ValueError("Concentration vector size does not match the model species count.")
    env: dict[str, object] = {name: func for name, func in _ALLOWED_FUNCTIONS.items()}
    env.update(model.parameter_map)
    env.update({f"c_{sid}": float(value) for sid, value in zip(model.species_ids, c)})
    rates = []
    for rid, expression, code in zip(model.reaction_ids, model.rate_expressions, model.compiled_rates):
        try:
            value = float(eval(code, {"__builtins__": {}}, env))
        except Exception as exc:
            raise ValueError(f"Could not evaluate rate expression for {rid} ('{expression}'): {exc}") from exc
        if not math.isfinite(value):
            raise ValueError(f"Rate expression for {rid} returned a non-finite value.")
        rates.append(value)
    return np.asarray(rates, dtype=float)


def model_volume_m3(model: GeneralBatchModel, amounts_kmol: np.ndarray) -> float:
    if model.volume_model == "constant":
        return float(model.initial_volume_m3)
    if model.volume_model == "composition_dependent":
        n = np.maximum(np.asarray(amounts_kmol, dtype=float), 0.0)
        volume = float(np.sum(n * model.molecular_weight_kg_per_kmol / model.density_kg_per_m3))
        if not np.isfinite(volume) or volume <= 0:
            raise ValueError("Calculated composition-dependent liquid volume is not positive and finite.")
        return volume
    raise ValueError(f"Unknown volume model: {model.volume_model}")


def simulate_general_isothermal_batch(
    model: GeneralBatchModel,
    *,
    method: str = "RK45",
    rtol: float = 1e-6,
    atol: float = 1e-9,
    max_function_evaluations: int | None = None,
    max_wall_time_s: float | None = None,
) -> dict:
    if solve_ivp is None:
        raise RuntimeError("SciPy is required for the general isothermal batch simulator.")

    n0 = model.initial_moles_kmol.astype(float)
    evaluation_count = 0
    started_at = time.monotonic()

    def rhs(_t: float, n: np.ndarray) -> np.ndarray:
        nonlocal evaluation_count
        evaluation_count += 1
        if max_function_evaluations is not None and evaluation_count > int(max_function_evaluations):
            raise RuntimeError(
                "Simulation stopped because the configured ODE function-evaluation limit was reached. "
                "Simplify the model or shorten the simulated interval."
            )
        if max_wall_time_s is not None and (time.monotonic() - started_at) > float(max_wall_time_s):
            raise RuntimeError(
                "Simulation stopped because the configured solver wall-time limit was reached. "
                "Simplify the model or shorten the simulated interval."
            )
        n_for_rates = np.maximum(np.asarray(n, dtype=float), 0.0)
        volume = model_volume_m3(model, n_for_rates)
        concentrations = n_for_rates / volume
        rates = evaluate_rates(model, concentrations)
        return volume * (model.stoichiometric_matrix @ rates)

    t_eval = np.linspace(0.0, float(model.batch_duration_s), int(model.output_points))
    sol = solve_ivp(
        rhs,
        (0.0, float(model.batch_duration_s)),
        n0,
        method=method,
        rtol=rtol,
        atol=atol,
        t_eval=t_eval,
        dense_output=True,
    )
    if not sol.success:
        raise RuntimeError(f"ODE solver failed: {sol.message}")

    amounts = np.maximum(np.asarray(sol.y, dtype=float).T, 0.0)
    volumes = np.asarray([model_volume_m3(model, row) for row in amounts], dtype=float)
    concentrations = amounts / volumes[:, None]
    rates = np.vstack([evaluate_rates(model, row) for row in concentrations])
    min_amount = float(np.min(np.asarray(sol.y, dtype=float).T))

    return {
        "time_s": np.asarray(sol.t, dtype=float),
        "amounts_kmol": amounts,
        "concentrations_kmol_m3": concentrations,
        "reaction_rates_kmol_m3_s": rates,
        "volume_m3": volumes,
        "initial_volume_m3": float(volumes[0]),
        "final_volume_m3": float(volumes[-1]),
        "solver_success": bool(sol.success),
        "solver_message": str(sol.message),
        "nfev": int(sol.nfev),
        "evaluation_guard_count": int(evaluation_count),
        "solver_wall_time_s": float(time.monotonic() - started_at),
        "method": method,
        "rtol": float(rtol),
        "atol": float(atol),
        "min_amount_kmol": min_amount,
        "dense_solution": sol.sol,
    }


def species_peak(model: GeneralBatchModel, simulation: dict, species_id: str) -> dict:
    if species_id not in model.species_ids:
        raise ValueError(f"Unknown species '{species_id}'.")
    idx = model.species_ids.index(species_id)
    dense = simulation.get("dense_solution")
    t_end = float(model.batch_duration_s)

    values = simulation["amounts_kmol"][:, idx]
    times = simulation["time_s"]
    pos = int(np.argmax(values))
    t_peak = float(times[pos])
    n_peak = float(values[pos])

    # Find the global peak on the regular output grid first, then refine only
    # around that local interval with the dense solution. This avoids assuming
    # an arbitrary user-defined model has a single/unimodal peak over the full batch.
    if dense is not None and minimize_scalar is not None and 0 < pos < len(times) - 1:
        left = float(times[pos - 1])
        right = float(times[pos + 1])
        optimum = minimize_scalar(
            lambda t: -float(dense(float(t))[idx]),
            bounds=(left, right),
            method="bounded",
            options={"xatol": max(1e-8, t_end * 1e-10)},
        )
        candidate_t = float(optimum.x)
        candidate_n = float(dense(candidate_t)[idx])
        if candidate_n >= n_peak:
            t_peak = candidate_t
            n_peak = candidate_n

    if dense is not None:
        n_at_peak = np.maximum(np.asarray(dense(t_peak), dtype=float), 0.0)
        peak_volume = model_volume_m3(model, n_at_peak)
    else:
        peak_volume = float(np.interp(t_peak, simulation["time_s"], simulation["volume_m3"]))

    return {
        "species_id": species_id,
        "time_s": t_peak,
        "time_min": t_peak / 60.0,
        "amount_kmol": n_peak,
        "concentration_kmol_m3": n_peak / peak_volume,
        "volume_m3": peak_volume,
        "at_end": bool(t_peak >= 0.999 * t_end),
    }


def generated_species_equations(model: GeneralBatchModel) -> list[str]:
    equations: list[str] = []
    for sidx, sid in enumerate(model.species_ids):
        terms: list[str] = []
        for ridx, rid in enumerate(model.reaction_ids):
            coeff = float(model.stoichiometric_matrix[sidx, ridx])
            if abs(coeff) < 1e-14:
                continue
            magnitude = abs(coeff)
            if abs(magnitude - 1.0) < 1e-14:
                term = f"r_{rid}"
            else:
                term = f"{magnitude:g} r_{rid}"
            if not terms:
                terms.append(("-" if coeff < 0 else "") + term)
            else:
                terms.append((" - " if coeff < 0 else " + ") + term)
        rhs = "".join(terms) if terms else "0"
        equations.append(f"dn_{sid}/dt = V ({rhs})")
    return equations


def build_results_workbook(
    original_bytes: bytes,
    model: GeneralBatchModel,
    simulation: dict,
    target_species: str,
) -> bytes:
    if load_workbook is None:
        raise RuntimeError("openpyxl is required to generate the completed Excel workbook.")
    wb = load_workbook(BytesIO(original_bytes), data_only=False, read_only=False, keep_links=False)
    for name in ("Results_TimeSeries", "Results_Summary"):
        if name in wb.sheetnames:
            del wb[name]

    dark_fill = PatternFill("solid", fgColor="111827")
    teal_fill = PatternFill("solid", fgColor="0F766E")
    header_font = Font(bold=True, color="FFFFFF")

    ws = wb.create_sheet("Results_TimeSeries")
    headers = ["time_s", "time_min", "volume_m3"]
    headers += [f"n_{sid}_kmol" for sid in model.species_ids]
    headers += [f"c_{sid}_kmol_m3" for sid in model.species_ids]
    headers += [f"r_{rid}_kmol_m3_s" for rid in model.reaction_ids]
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = teal_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center")

    for i, t in enumerate(simulation["time_s"]):
        row = [float(t), float(t) / 60.0, float(simulation["volume_m3"][i])]
        row += [float(v) for v in simulation["amounts_kmol"][i, :]]
        row += [float(v) for v in simulation["concentrations_kmol_m3"][i, :]]
        row += [float(v) for v in simulation["reaction_rates_kmol_m3_s"][i, :]]
        ws.append(row)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for column_cells in ws.columns:
        letter = column_cells[0].column_letter
        ws.column_dimensions[letter].width = min(24, max(12, len(str(column_cells[0].value)) + 2))

    peak = species_peak(model, simulation, target_species)
    idx = model.species_ids.index(target_species)
    ws2 = wb.create_sheet("Results_Summary")
    ws2.append(["KineticFittingPot simulation summary", "Value", "Units / note"])
    for cell in ws2[1]:
        cell.fill = dark_fill
        cell.font = header_font
    rows = [
        ("Model name", model.model_name, ""),
        ("Model description", model.description, ""),
        ("Simulation type", "General isothermal batch", ""),
        ("Temperature", model.temperature_c, "°C (informational unless reflected in parameters)"),
        ("Volume model", model.volume_model, "constant or composition_dependent"),
        ("Initial volume", simulation["initial_volume_m3"], "m³"),
        ("Final volume", simulation["final_volume_m3"], "m³"),
        ("Batch duration", model.batch_duration_s, "s"),
        ("Target species", target_species, model.display_names[idx]),
        ("Target peak time", peak["time_s"], "s"),
        ("Target peak time", peak["time_min"], "min"),
        ("Target peak amount", peak["amount_kmol"], "kmol"),
        ("Target final amount", float(simulation["amounts_kmol"][-1, idx]), "kmol"),
        ("Solver", simulation["method"], "SciPy solve_ivp"),
        ("rtol", simulation["rtol"], ""),
        ("atol", simulation["atol"], ""),
        ("Function evaluations", simulation["nfev"], ""),
        ("Solver success", simulation["solver_success"], simulation["solver_message"]),
        ("Minimum simulated species amount", simulation["min_amount_kmol"], "kmol; small negative values may be numerical round-off"),
    ]
    for item in rows:
        ws2.append(list(item))
    ws2.column_dimensions["A"].width = 34
    ws2.column_dimensions["B"].width = 32
    ws2.column_dimensions["C"].width = 58
    for row in ws2.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    buffer = BytesIO()
    wb.save(buffer)
    return buffer.getvalue()
