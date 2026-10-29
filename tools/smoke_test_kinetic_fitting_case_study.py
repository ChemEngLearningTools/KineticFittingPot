from __future__ import annotations

import sys
from pathlib import Path

import numpy as np
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from utils.kinetic_fitting_engine import fitted_trajectory, parse_fitting_workbook


BOOK = ROOT / "data" / "case_studies" / "KineticFittingPot_kinetic_fit_results_Biodiesel.xlsx"
DIAG_BOOK = ROOT / "data" / "case_studies" / "KineticFittingPot_kinetic_fit_results_Biodiesel_model_diagnostics.xlsx"
MODULE = ROOT / "modules" / "kinetic_fitting_case_study.py"
APP = ROOT / "app.py"
HOME = ROOT / "modules" / "home.py"


def main() -> None:
    assert BOOK.exists(), "Bundled biodiesel case-study workbook is missing."
    assert DIAG_BOOK.exists(), "Bundled biodiesel completed diagnostics workbook is missing."
    raw = BOOK.read_bytes()
    problem = parse_fitting_workbook(raw)
    assert [d.dataset_id for d in problem.datasets] == ["A", "B", "C"]
    assert len(problem.observations) == 120
    assert tuple(problem.fit_parameter_ids) == ("k1f", "k1r", "k2f", "k2r", "k3f", "k3r")

    wb = load_workbook(BOOK, data_only=True, read_only=True)
    required = {"Datasets", "Reactions", "Fit_Parameters", "Fit_Starts", "Fit_Predictions", "Fit_Residuals", "Fit_Summary"}
    assert required.issubset(wb.sheetnames)

    fp = list(wb["Fit_Parameters"].iter_rows(min_row=2, values_only=True))
    assert len(fp) == 6
    fitted = np.asarray([float(row[2]) for row in fp], dtype=float)
    expected = np.asarray([
        4.538349784328156e-05,
        2.863653880613096e-04,
        1.088726876183076e-04,
        6.972111944217387e-04,
        9.58409421549487e-05,
        1.941979167437756e-04,
    ])
    assert np.allclose(fitted, expected, rtol=1e-10, atol=1e-14)

    starts = list(wb["Fit_Starts"].iter_rows(min_row=2, values_only=True))
    assert len(starts) == 20
    assert all(bool(row[1]) for row in starts)
    sse = np.asarray([float(row[2]) for row in starts])
    assert float(sse.max() - sse.min()) < 1e-8

    summary = {str(row[0]): row[1] for row in wb["Fit_Summary"].iter_rows(min_row=2, values_only=True) if row[0]}
    assert int(summary["Datasets"]) == 3
    assert int(summary["Observations"]) == 120
    assert int(summary["Fitted parameters"]) == 6
    assert summary["Weighting"] == "species_scaled"
    assert int(summary["Multistart runs"]) == 20

    for dataset in problem.datasets:
        traj = fitted_trajectory(problem, dataset, fitted, points=120)
        assert traj["concentrations_kmol_m3"].shape == (120, len(problem.model.species_ids))
        assert np.all(np.isfinite(traj["concentrations_kmol_m3"]))
        assert np.allclose(traj["concentrations_kmol_m3"][0], dataset.initial_concentration_kmol_m3, atol=1e-12)

    wb.close()

    dwb = load_workbook(DIAG_BOOK, data_only=True, read_only=True)
    diag_required = {
        "Diag_Settings", "Diag_Invariants", "Diag_Balance_Points", "Diag_Balance_Summary",
        "Diag_Recon_Observations", "Diag_Recon_Species", "Diag_Residual_Species",
        "Diag_Local_Summary", "Diag_Local_Parameters", "Diag_Parameter_Corr",
        "Diag_Profile_Points", "Diag_Zero_Tests"
    }
    assert diag_required.issubset(dwb.sheetnames)

    balance_summary = {
        str(r[0]): (float(r[3]), int(r[6]), float(r[7]))
        for r in dwb["Diag_Balance_Summary"].iter_rows(min_row=2, values_only=True)
        if r[0]
    }
    assert balance_summary["A"][1] == 0
    assert balance_summary["B"][1] == 1
    assert balance_summary["C"][1] == 2
    assert np.isclose(balance_summary["C"][0], 12.622002, rtol=1e-6)

    recon_rows = list(dwb["Diag_Recon_Observations"].iter_rows(min_row=2, values_only=True))
    me_z = np.asarray([abs(float(r[8])) for r in recon_rows if str(r[3]) == "ME"])
    tg_z = np.asarray([abs(float(r[8])) for r in recon_rows if str(r[3]) == "TG"])
    assert int(np.sum(me_z > 2.0)) == 12
    assert int(np.sum(me_z > 3.0)) == 5
    assert int(np.sum(tg_z > 2.0)) == 1
    assert np.isclose(float(me_z.max()), 4.700060, rtol=1e-6)

    local_summary = {
        str(r[0]): r[1]
        for r in dwb["Diag_Local_Summary"].iter_rows(min_row=2, values_only=True)
        if r[0]
    }
    assert int(local_summary["Numerical rank"]) == 6
    assert np.isclose(float(local_summary["Condition number"]), 8.527589, rtol=1e-6)

    zero_tests = {
        str(r[0]): float(r[5])
        for r in dwb["Diag_Zero_Tests"].iter_rows(min_row=2, values_only=True)
        if r[0]
    }
    assert np.isclose(zero_tests["k1r"], 23.054781, rtol=1e-6)
    assert np.isclose(zero_tests["k2r"], 159.385575, rtol=1e-6)
    assert np.isclose(zero_tests["k3r"], 177.954638, rtol=1e-6)
    dwb.close()

    module_text = MODULE.read_text(encoding="utf-8")
    assert "does not run the kinetic optimiser" in module_text
    assert "not a reproduction of the paper's LHHW kinetic model" in module_text
    assert "Download completed case-study workbook" in module_text
    assert "A good fit is not the end of model development" in module_text
    assert "Download complete diagnostics workbook" in module_text
    assert "Are the experimental measurements chemically self-consistent?" in module_text
    assert "Which measurements drive the stoichiometric inconsistency?" in module_text
    assert "Are the six fitted parameters locally determined by the data?" in module_text
    assert "Do the data support the reverse kinetic terms?" in module_text
    assert "12 of 30 ME observations" in module_text
    assert "Case-study conclusion" in module_text
    assert "diagnostics export will be expanded in the next development build" not in module_text

    app_text = APP.read_text(encoding="utf-8")
    home_text = HOME.read_text(encoding="utf-8")
    assert '"Kinetic Fitting Case Study — Biodiesel": kinetic_fitting_case_study.render' in app_text
    assert '"Kinetic Fitting Case Study — Biodiesel", "Worked case study"' in home_text
    assert "CASE_STUDIES = [" in home_text

    print("PASS: biodiesel fitting + completed diagnostics case study, saved public results and route")
    print("Datasets=3; observations=120; parameters=6; multistarts=20; diagnostics: A/B/C balance, reconciliation, local rank and profiles verified")


if __name__ == "__main__":
    main()
