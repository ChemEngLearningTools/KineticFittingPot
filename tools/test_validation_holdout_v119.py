"""Regression: held-out observations cannot affect objective, fit or parameter estimates."""
from io import BytesIO
from pathlib import Path
import numpy as np
from openpyxl import load_workbook
from utils.kinetic_fitting_engine import (parse_fitting_workbook, fit_multistart,
    evaluate_validation, build_fitting_results_workbook)
from utils.deployment import public_fitting_violations

ROOT = Path(__file__).resolve().parents[1]
TEMPLATE = ROOT / 'data/templates/KineticFittingPot_Kinetic_Fitting_Catalytic_Saturation_Synthetic_Data.xlsx'

def build(add_validation):
    wb = load_workbook(TEMPLATE)
    ds = wb['Datasets']; data = wb['Experimental_Data']
    heads = {str(c.value): i for i,c in enumerate(ds[1], start=1)}
    assert 'Use_for_Fit' in heads
    rows = list(ds.iter_rows(min_row=2, values_only=True))
    source = next(row for row in rows if row[0] is not None)
    fit_ids = {str(row[heads['Dataset_ID']-1]) for row in rows if row[heads['Dataset_ID']-1] is not None and str(row[heads['Use_for_Fit']-1]).lower().startswith('y')}
    assert len(fit_ids) >= 1
    if add_validation:
        fresh = list(source)
        new_id = 'HOLDOUT_V119'
        fresh[heads['Dataset_ID']-1] = new_id
        fresh[heads['Use_for_Fit']-1] = 'No'
        ds.append(fresh)
        dheads={str(c.value):i for i,c in enumerate(data[1],start=1)}
        source_id=str(source[heads['Dataset_ID']-1])
        copied = 0
        for row in list(data.iter_rows(min_row=2, values_only=True)):
            if str(row[dheads['Dataset_ID']-1]) == source_id:
                values=list(row)
                values[dheads['Dataset_ID']-1]=new_id
                values[dheads['Measured_Value_kmol_m3']-1]=float(values[dheads['Measured_Value_kmol_m3']-1])+0.12
                data.append(values); copied+=1
        assert copied > 0
    b=BytesIO();wb.save(b);return b.getvalue()

def main():
    fit_only = parse_fitting_workbook(build(False))
    fit_with_holdout_bytes = build(True)
    combined = parse_fitting_workbook(fit_with_holdout_bytes)
    assert len(combined.datasets)==len(fit_only.datasets)
    assert len(combined.observations)==len(fit_only.observations)
    assert len(combined.validation_datasets)==1
    assert len(combined.validation_observations)>0
    assert not public_fitting_violations(combined)
    a=fit_multistart(fit_only,weighting='species_scaled',n_starts=1,max_nfev=35,random_seed=42)
    b=fit_multistart(combined,weighting='species_scaled',n_starts=1,max_nfev=35,random_seed=42)
    assert np.allclose(a['best_values'],b['best_values'],rtol=1e-10,atol=1e-12)
    assert np.isclose(a['best']['sse'],b['best']['sse'],rtol=1e-12,atol=1e-12)
    validation=evaluate_validation(combined,b)
    assert len(validation['predictions'])==len(combined.validation_observations)
    assert validation['rmse_kmol_m3']>0
    b['validation']=validation
    output=load_workbook(BytesIO(build_fitting_results_workbook(fit_with_holdout_bytes,combined,b)))
    assert 'Validation_Predictions' in output.sheetnames
    assert 'Validation_Residuals' in output.sheetnames
    assert 'Validation_Summary' in output.sheetnames
    assert output['Fit_Predictions'].max_row == len(combined.observations)+1
    assert output['Validation_Predictions'].max_row==len(combined.validation_observations)+1
    print(f'PASS v1.19 holdout: {len(combined.datasets)} fit datasets, {len(combined.validation_datasets)} validation datasets, '
          f'{len(combined.validation_observations)} validation observations; exact matching fitting objective and parameter estimates')
    print(f'Validation RMSE: {validation["rmse_kmol_m3"]:.6g}; separate workbook exports verified')

if __name__=='__main__': main()
