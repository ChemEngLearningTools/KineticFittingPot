from __future__ import annotations

import sys
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from openpyxl import load_workbook

from utils.general_batch_model import parse_model_workbook, simulate_general_isothermal_batch


BOOK = ROOT / 'data' / 'templates' / 'KineticFittingPot_Kinetic_Fitting_Catalytic_Saturation_Synthetic_Data.xlsx'


def main() -> None:
    raw = BOOK.read_bytes()
    model = parse_model_workbook(raw)
    assert model.species_ids == ('A', 'B', 'D', 'Cat')
    assert model.reaction_ids == ('R1', 'R2')
    assert all(model.fit_later)
    assert all(x is not None for x in model.lower_bounds)
    assert all(x is not None for x in model.upper_bounds)

    # Extra fitting-data sheets must not break the established forward parser.
    result = simulate_general_isothermal_batch(model)
    assert result['solver_success']

    wb = load_workbook(BOOK, data_only=True, read_only=True)
    assert {'Datasets', 'Experimental_Data', 'Synthetic_Data_Info'}.issubset(wb.sheetnames)

    datasets = list(wb['Datasets'].iter_rows(min_row=2, values_only=True))
    assert len(datasets) == 3
    assert [row[0] for row in datasets] == ['E1', 'E2', 'E3']
    assert [float(row[8]) for row in datasets] == [0.05, 0.10, 0.20]

    obs = list(wb['Experimental_Data'].iter_rows(min_row=2, values_only=True))
    assert len(obs) == 64
    assert {row[0] for row in obs} == {'E1', 'E2', 'E3'}
    # Long format: missing combinations are absent rather than blank measured values.
    assert all(row[3] is not None for row in obs)
    assert all(row[1] is not None and row[2] is not None for row in obs)

    info = {row[0]: row[1] for row in wb['Synthetic_Data_Info'].iter_rows(min_row=2, values_only=True) if row[0]}
    assert float(info['True_k1']) == 0.020
    assert float(info['True_KA']) == 1.0
    assert float(info['True_k2']) == 0.00030
    # The model starts must intentionally differ from truth.
    start = dict(zip(model.parameter_ids, model.parameter_values))
    assert start['k1'] != float(info['True_k1'])
    assert start['KA'] != float(info['True_KA'])
    assert start['k2'] != float(info['True_k2'])

    print('Kinetic fitting fundamentals/data-architecture smoke test passed.')
    print(f'Datasets: {len(datasets)}; observations: {len(obs)}')
    print('Starting estimates:', start)


if __name__ == '__main__':
    main()
