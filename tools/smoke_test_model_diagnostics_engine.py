from __future__ import annotations

from pathlib import Path
import sys

import numpy as np
import pandas as pd
from io import BytesIO
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))

from utils.kinetic_fitting_engine import build_fitting_results_workbook, fit_multistart, parse_fitting_workbook  # noqa: E402
from utils.model_diagnostics_engine import (  # noqa: E402
    classify_invariants,
    exact_rank,
    format_invariant,
    invariant_basis,
    local_parameter_diagnostics,
    profiled_objective_diagnostics,
    measured_only_invariant_basis,
    residual_autocorrelation,
    simple_nonnegative_integer_invariants,
    stoichiometric_diagnostics,
    time_grid_regularity,
)
from utils.model_diagnostics_export import build_model_diagnostics_workbook  # noqa: E402


species = ("TG", "DG", "MG", "MeOH", "ME", "Gly")
S = np.asarray([
    [-1, +1,  0,  0,  0,  0],
    [+1, -1, -1, +1,  0,  0],
    [ 0,  0, +1, -1, -1, +1],
    [-1, +1, -1, +1, -1, +1],
    [+1, -1, +1, -1, +1, -1],
    [ 0,  0,  0,  0, +1, -1],
], dtype=float)

assert exact_rank(S) == 3
basis = invariant_basis(S)
assert len(basis) == 3
for coeffs in basis:
    assert np.allclose(np.asarray(coeffs, dtype=float) @ S, 0.0)

expected = {
    (0, 0, 0, 1, 1, 0),
    (1, 1, 1, 0, 0, 1),
    (3, 2, 1, 0, 1, 0),
}
assert set(basis) == expected, basis

measured = ("TG", "DG", "MG", "ME")
mbasis = measured_only_invariant_basis(S, species, measured)
assert mbasis == ((3, 2, 1, 0, 1, 0),), mbasis
classified = classify_invariants(basis, species, measured)
assert sum(x.fully_observed for x in classified) == 1
assert format_invariant(mbasis[0], species) == "3TG + 2DG + MG + ME = constant"

# v1.26: finite simple-integer catalogue. The non-negative primitive 0..4
# biodiesel search has 29 readable conserved quantities but still spans only
# the same 3-dimensional invariant space.
simple = simple_nonnegative_integer_invariants(S, max_coefficient=4)
assert len(simple) == 29, len(simple)
assert all(np.allclose(np.asarray(v, dtype=float) @ S, 0.0) for v in simple)
assert len({tuple(v) for v in simple}) == len(simple)
assert sum(x.fully_observed for x in classify_invariants(simple, species, measured)) == 1
assert (1, 1, 1, 1, 1, 1) in simple
assert (3, 2, 1, 0, 1, 0) in simple


# v1.27: ACF is only available for a sufficiently regular observation grid.
regular = time_grid_regularity([0, 10, 20, 30, 40, 50, 60])
assert regular["regular"] is True
irregular = time_grid_regularity([0, 15, 30, 45, 60, 75, 90, 120, 180, 240])
assert irregular["regular"] is False
acf = residual_autocorrelation([1.0, 0.5, 0.0, -0.5, -1.0, -0.5, 0.0], max_lag=3)
assert len(acf) == 4
assert abs(acf[0][1] - 1.0) < 1e-12



# v1.28: local Jacobian / parameter-correlation diagnostics on the bundled
# catalytic-saturation synthetic fitting example. Build a small completed fit
# in-memory so the smoke test does not depend on an external user workbook.
template = ROOT / "data" / "templates" / "KineticFittingPot_Kinetic_Fitting_Catalytic_Saturation_Synthetic_Data.xlsx"
raw_template = template.read_bytes()
problem = parse_fitting_workbook(raw_template)
fit = fit_multistart(problem, weighting="species_scaled", n_starts=2, random_seed=20260904, max_nfev=100)
completed = build_fitting_results_workbook(raw_template, problem, fit)
from utils.model_diagnostics_engine import parse_completed_fit_workbook  # noqa: E402
completed_info = parse_completed_fit_workbook(completed)
local = local_parameter_diagnostics(completed_info)
assert local.numerical_rank == len(problem.fit_parameter_ids)
assert np.isfinite(local.condition_number) and local.condition_number > 1.0
assert local.correlation.shape == (len(problem.fit_parameter_ids), len(problem.fit_parameter_ids))
assert np.allclose(np.diag(local.correlation), 1.0)
assert abs(local.objective_sse - fit["best"]["sse"]) < 1e-8
assert len(local.singular_values) == len(problem.fit_parameter_ids)
assert len(local.parameter_rows) == len(problem.fit_parameter_ids)


# v1.29: profiled objective. Profile one parameter on a small grid and verify
# that the completed optimum is retained as the zero-deterioration reference.
profile = profiled_objective_diagnostics(
    completed_info,
    parameter_ids=(problem.fit_parameter_ids[0],),
    grid_points=5,
    span_factor=5.0,
    max_nfev=60,
)
assert len(profile.profiles) == 1
assert abs(profile.base_objective_sse - fit["best"]["sse"]) < 1e-8
pp = profile.profiles[0]
assert any(abs(point.fixed_value - pp.fitted_value) <= max(1e-14, abs(pp.fitted_value)*1e-10) for point in pp.points)
assert min(point.delta_sse for point in pp.points if point.success) >= -1e-5

# v1.37: diagnostics export preserves the completed fitting workbook and appends
# a machine-readable diagnostic record, including local geometry and stored
# profile re-optimised parameter values.
completed_diag = stoichiometric_diagnostics(completed_info)
exported = build_model_diagnostics_workbook(
    completed,
    completed_info,
    completed_diag,
    integer_artifact={
        "max_coefficient": 4,
        "classified": classify_invariants(
            simple_nonnegative_integer_invariants(completed_info.problem.model.stoichiometric_matrix, max_coefficient=4),
            completed_info.problem.model.species_ids,
            completed_info.measured_species,
        ),
    },
    residual_artifact={
        "relative_fraction": 0.03,
        "absolute_floor": 0.002,
        "points": pd.DataFrame(),
        "species_summary": pd.DataFrame(),
        "dataset_summary": pd.DataFrame(),
        "dataset_species_summary": pd.DataFrame(),
    },
    local_artifact={
        "result": local,
        "parameter_table": pd.DataFrame([{"Parameter": r.parameter_id, "Fitted value": r.fitted_value} for r in local.parameter_rows]),
        "pair_table": pd.DataFrame(),
    },
    profile_artifact={
        "result": profile,
        "selected_parameters": (problem.fit_parameter_ids[0],),
        "zero_parameters": (),
        "grid_points": 5,
        "span_factor": 5.0,
        "max_nfev": 60,
        "controls_match": True,
    },
)
export_wb = load_workbook(BytesIO(exported), data_only=True, read_only=True)
for original_sheet in ("Fit_Parameters", "Fit_Predictions", "Fit_Residuals", "Fit_Summary"):
    assert original_sheet in export_wb.sheetnames
for diagnostic_sheet in (
    "Diag_Summary", "Diag_Settings", "Diag_Stoich_Matrix", "Diag_Invariants",
    "Diag_Local_Parameters", "Diag_Jacobian", "Diag_Parameter_Corr",
    "Diag_Profile_Summary", "Diag_Profile_Points",
):
    assert diagnostic_sheet in export_wb.sheetnames
profile_headers = [c.value for c in next(export_wb["Diag_Profile_Points"].iter_rows(min_row=1, max_row=1))]
assert any(str(h).startswith("Optimised_") for h in profile_headers if h is not None)
export_wb.close()

app_text = (ROOT / "app.py").read_text(encoding="utf-8")
home_text = (ROOT / "modules" / "home.py").read_text(encoding="utf-8")
fund_text = (ROOT / "modules" / "model_diagnostics_fundamentals.py").read_text(encoding="utf-8")
module_text = (ROOT / "modules" / "model_diagnostics.py").read_text(encoding="utf-8")
assert '"Post-fit Model Diagnostics": DIAG' in app_text
assert '("Advanced Model Diagnostics", "Local-only analysis"' in home_text
assert '("Basic Model Diagnostics", "Post-fit assessment"' in home_text
assert '"Basic Model Diagnostics": basic_model_diagnostics.render' in app_text
assert "This Fundamentals page defines the diagnostic concepts only" not in fund_text
assert "Upload KineticFittingPot kinetic-fitting results workbook" in module_text
assert "Independent closed-batch mole-balance constraints" in module_text
assert "Simple integer conserved quantities" in module_text
assert "Measurement-uncertainty model" in module_text
assert "Stoichiometric data reconciliation" in module_text
assert "Dynamic residual diagnostics" in module_text
assert "Residual versus time" in module_text
assert "Residual versus fitted value" in module_text
assert "Cross-species residual correlation" in module_text
assert "Residual autocorrelation (only when the sampling grid supports it)" in module_text
assert "Local parameter / Jacobian diagnostics" in module_text
assert "Approximate local parameter correlation" in module_text
assert "Weakest local parameter combination" in module_text
assert "Profiled objective / profile-likelihood diagnostics" in module_text
assert "Explicit zero-rate boundary tests" in module_text
assert "Run parameter profiles" in module_text
assert "Constraints available from the measured species alone" in module_text
assert "Save model diagnostics" in module_text
assert "Download completed model-diagnostics workbook" in module_text
assert "Diag_*" in module_text

print("Model diagnostics engine smoke test PASS")
print("rank(S)=3; independent invariant dimension=3; measured-only invariant dimension=1")
print("0..4 primitive non-negative conserved quantities=29")
print("residual ACF regular-grid gate: PASS")
print(f"local Jacobian diagnostics: rank={local.numerical_rank}, condition={local.condition_number:.4g}: PASS")
print("profiled-objective constrained re-optimisation: PASS")
print("completed model-diagnostics Excel export: PASS")
print("biodiesel measured invariant: 3TG + 2DG + MG + ME = constant")
