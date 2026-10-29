from __future__ import annotations

from pathlib import Path

import pandas as pd
import streamlit as st
from openpyxl import load_workbook

from utils.ui import chips, navigate, workspace_header


ACCENT = "#2dd4bf"
EXAMPLE_PATH = (
    Path(__file__).resolve().parents[1]
    / "data"
    / "templates"
    / "KineticFittingPot_Kinetic_Fitting_Catalytic_Saturation_Synthetic_Data.xlsx"
)


def _example_preview() -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    wb = load_workbook(EXAMPLE_PATH, data_only=True, read_only=True)

    ws = wb["Datasets"]
    dataset_rows = list(ws.iter_rows(min_row=1, values_only=True))
    datasets = pd.DataFrame(dataset_rows[1:], columns=dataset_rows[0])

    ws = wb["Parameters"]
    parameter_rows = list(ws.iter_rows(min_row=1, values_only=True))
    parameters = pd.DataFrame(parameter_rows[1:], columns=parameter_rows[0])

    ws = wb["Experimental_Data"]
    experimental_rows = list(ws.iter_rows(min_row=1, values_only=True))
    experimental = pd.DataFrame(experimental_rows[1:], columns=experimental_rows[0])
    return datasets, parameters, experimental


def render() -> None:
    workspace_header(
        "Kinetic Fitting Fundamentals",
        "Understand the inverse parameter-estimation problem before running an optimiser: multiple experiments, irregular observations, missing data, residual weighting, bounds, local minima and parameter identifiability.",
        section="Kinetic Fitting",
        accent=ACCENT,
    )

    chips([
        ("Learning", "blue"),
        ("Inverse problem", "blue"),
        ("Multiple datasets", "blue"),
        ("Irregular sampling", ""),
        ("Missing data", ""),
        ("Multistart fitting", "warn"),
        ("Identifiability", "warn"),
    ])

    st.info(
        "KineticFittingPot already has the forward problem: an Excel-defined kinetic model is simulated to predict concentrations. "
        "Kinetic fitting reverses that workflow. Experimental concentrations are supplied and selected model parameters are adjusted until one common parameter set explains all selected experiments as well as possible."
    )

    st.markdown("## Forward simulation and inverse fitting")
    c1, c2 = st.columns(2)
    with c1:
        st.markdown("### Forward problem")
        st.latex(r"\boldsymbol\theta\;\longrightarrow\;\text{ODE simulation}\;\longrightarrow\;\mathbf y_{model}(t)")
        st.markdown(
            "The kinetic parameters, initial conditions and reaction model are known. KineticFittingPot integrates the ODEs and predicts the concentration-time response. This is what **Simulate Your Own Model** already does."
        )
    with c2:
        st.markdown("### Inverse problem")
        st.latex(r"\mathbf y_{measured}(t)\;\longrightarrow\;\text{adjust }\boldsymbol\theta\;\longrightarrow\;\text{best agreement}")
        st.markdown(
            "The measurements are known, while one or more kinetic parameters are uncertain. The optimiser repeatedly calls the same forward ODE model with trial parameter values."
        )

    st.markdown("## One kinetic model — multiple experiments")
    st.markdown(
        "A useful kinetic fit should normally use **one shared kinetic parameter vector** across all experiments that represent the same chemistry. Each experiment may nevertheless have different initial conditions, catalyst loading, volume, duration and measurement times."
    )
    st.latex(r"\boldsymbol\theta=[k_1,K_A,k_2,\ldots]\qquad\text{shared by datasets }E_1,E_2,\ldots,E_N")
    st.markdown(
        "For every trial parameter vector, the fitter simulates each selected dataset independently and concatenate all valid model-data differences into one residual vector. The optimiser therefore has to find parameters that explain **all datasets simultaneously**, not a different parameter set for each experiment."
    )

    st.markdown("## Different time points are expected")
    st.markdown(
        "Experiments do **not** need a common measurement grid. Dataset E1 might contain measurements at 0, 180 and 420 s while E2 might use 0, 120 and 300 s. The fitter integrates each experiment and evaluates the continuous ODE solution at that dataset's actual measurement times."
    )
    st.warning(
        "Experimental data should not be resampled merely to make the spreadsheets line up. The model is evaluated at the observation time; the observation is not moved to the model grid."
    )

    st.markdown("## Missing data remain missing")
    st.markdown(
        "A species may be measured at one time and not another, and different species may have completely different sampling schedules. The fitter treats a missing observation as **no residual at that species/time combination**. Missing values are not interpolated, filled or invented."
    )
    st.markdown(
        "This is why the fitting workbook uses a **long-format Experimental_Data sheet**: one row means one real observation. If a measurement was not made, there is simply no row."
    )

    st.markdown("## Residuals and the objective function")
    st.markdown(
        "For every actual observation, the fitter compares the measured value with the model prediction at the same dataset, species and time. A simple residual is:"
    )
    st.latex(r"e_m=y_{m,model}-y_{m,measured}")
    st.markdown("A least-squares fit minimises an objective based on all residuals:")
    st.latex(r"\Phi(\boldsymbol\theta)=\sum_m e_m^2")
    st.markdown(
        "However, raw residuals can be misleading if one species is numerically much larger than another. The fitting interface provides explicit **unweighted** and **species-scaled** residual options. Weighting affects the contribution of different observations to the objective function; it does not change the underlying reactor model."
    )

    st.markdown("## Parameter estimates and finite bounds")
    st.markdown(
        "Optimisation needs a starting estimate, but a starting estimate is not a fact. The workbook therefore separates the current parameter value from a physically plausible lower and upper bound. For fitted parameters the fitting workspace requires finite bounds."
    )
    st.latex(r"\theta_{j,lower}\leq\theta_j\leq\theta_{j,upper}")
    st.markdown(
        "The bounds represent engineering knowledge as well as helping the numerical algorithm. The parameter bounds are supplied explicitly by the user, rather than being inferred from the initial guess."
    )

    st.markdown("## Why one optimisation run is not enough")
    st.markdown(
        "Nonlinear kinetic objectives can contain local minima, flat regions and strongly correlated parameter combinations. A single least-squares run may converge successfully but still finish in a poor basin—or simply reproduce the influence of a weak starting guess."
    )
    st.markdown("Parameter estimation uses **bounded multistart nonlinear least squares**:")
    st.markdown(
        "1. include the user's supplied parameter estimates as one start;  "
        "2. generate additional starting vectors inside the user-specified bounds;  "
        "3. run bounded `scipy.optimize.least_squares` from every start;  "
        "4. rank the converged solutions by objective value;  "
        "5. retain the best solution while showing the competing outcomes."
    )
    st.info(
        "Multistart does not guarantee that a mathematical global minimum has been proven. It is a practical diagnostic: if many very different starts repeatedly reach the same parameter vector and objective value, confidence in the numerical optimum is much greater."
    )

    st.markdown("## Optimisation difficulty is not the same as identifiability")
    st.markdown(
        "If different starts settle at very different parameter values with almost identical fit quality, the optimiser may not be the real problem. The available experiments may simply be unable to distinguish those parameters. That is **parameter identifiability**."
    )
    st.warning(
        "A successful optimiser message does not prove that the parameters are uniquely determined. Examine the spread of multistart solutions and residual structure alongside the best parameter estimates."
    )

    st.markdown("## What a useful fitting result must show")
    st.markdown(
        "The fitting workspace reports more than fitted values: measured points over fitted trajectories for each dataset, residuals versus time/species, the objective reached by every start, parameter estimates from competing good solutions, solver failures, and the number of actual observations contributing to the fit."
    )
    st.markdown(
        "A good fit does not by itself establish parameter uncertainty. Model adequacy, residual structure, competing optima and identifiability must be considered before interpreting fitted values as precise estimates."
    )

    st.markdown("## Worked fitting architecture — catalytic saturation example")
    st.latex(r"r_1=\frac{k_1 c_{Cat} c_A}{1+K_Ac_A},\qquad r_2=k_2c_B")
    st.markdown(
        "The worked fitting example reuses the same catalytic saturation workbook that demonstrated the general forward simulator. Three synthetic experiments use catalyst loadings of **0.05, 0.10 and 0.20 kmol m⁻³**. The experiments have different sampling schedules and deliberately incomplete species measurements."
    )

    if EXAMPLE_PATH.exists():
        try:
            datasets, parameters, experimental = _example_preview()
            st.markdown("### Dataset-specific initial conditions")
            st.dataframe(
                datasets[["Dataset_ID", "Description", "Volume_m3", "c0_A", "c0_B", "c0_D", "c0_Cat"]],
                use_container_width=True,
                hide_index=True,
            )

            parameter_view = parameters[["Parameter_ID", "Value", "Lower_Bound", "Upper_Bound"]].copy()
            parameter_view.columns = ["Parameter", "Starting estimate", "Lower bound", "Upper bound"]
            st.markdown("### Deliberately imperfect starting estimates")
            st.dataframe(parameter_view, use_container_width=True, hide_index=True)

            summary = experimental.groupby("Dataset_ID").agg(
                observations=("Measured_Value_kmol_m3", "size"),
                first_time_s=("Time_s", "min"),
                last_time_s=("Time_s", "max"),
                species_measured=("Species_ID", lambda x: ", ".join(sorted(set(x)))),
            ).reset_index()
            st.markdown("### Experimental-data structure")
            st.dataframe(summary, use_container_width=True, hide_index=True)
            st.caption(
                f"The bundled workbook contains **{len(experimental)} actual observations**. Missing measurements are represented by absent rows, not blanks that KineticFittingPot must guess how to fill."
            )

            st.download_button(
                "Download catalytic kinetic-fitting example workbook",
                data=EXAMPLE_PATH.read_bytes(),
                file_name=EXAMPLE_PATH.name,
                mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                type="primary",
            )
        except Exception as exc:
            st.error(f"The bundled fitting example could not be previewed: {exc}")
    else:
        st.error("The bundled catalytic kinetic-fitting example workbook is missing from this installation.")

    st.markdown("### Synthetic truth is kept separate")
    st.markdown(
        "The workbook contains a **Synthetic_Data_Info** sheet recording the fixed random seed, noise model and true generating values. Those values allow independent comparison of fitted estimates with the known generating parameters. They must never be read by the optimiser when estimating parameters."
    )

    st.markdown("## Fitting algorithm implemented")
    st.markdown(
        "The Kinetic Parameter Fitting workspace now reuses this workbook and the existing general ODE engine with multiple-dataset parsing, observation-time evaluation, missing-data handling, parameter selection, explicit residual weighting, finite bounds, bounded multistart `least_squares`, fitted trajectories, residual plots and multistart diagnostics. **No fitting algorithm is executed on this Fundamentals page; use the separate fitting workspace when you are ready to estimate parameters.**"
    )

    st.divider()
    c1, c2 = st.columns(2)
    with c1:
        if st.button("← Open Simulate Your Own Model", key="fit_fund_to_forward"):
            navigate("Simulate Your Own Model — Isothermal Batch Reactor")
    with c2:
        if st.button("Open Kinetic Parameter Fitting →", key="fit_fund_to_planned"):
            navigate("Kinetic Parameter Fitting")
